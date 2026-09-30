# -*- coding: utf-8 -*-
"""hespokestyle.com 批量商品采集工具（WordPress + WooCommerce 自定义主题 + DrissionPage）。

站点特征（见开发流程文档第四、十二节）：
- 商品信息在服务端渲染 DOM + Yoast JSON-LD（没有 WooCommerce 商品 schema，
  也没有 ``data-product_variations``），因此走「渲染后商品 DOM」这一层。
- 币种固定 USD，价格无需汇率换算；平铺价无划线价，``price2 = price1``。
- 定制西装只有「Jacket Try-On」一个尺码维度（``select#jacket-size``），
  领型 / 口袋 / 里衬属于定制选项，不作为变体，也不参与 styles1 笛卡尔积。
- ``onhand="0"`` 的尺码表示试穿样衣缺货，不作为可购尺码，需从变体中排除。
- 图片走 JSON-LD ``ImageObject``（干净的 wp-content 原图），NitroPack 懒加载
  图片是占位符，因此使用 ``format_image_url(..., use_query_marker=True)``
  追加 ``?_600x600=1`` 尺寸标记，不改写文件名。
"""

import json
import os
import random
import re
import tempfile
import threading
import time
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext
from urllib.parse import urlsplit

import openpyxl
from DrissionPage import Chromium, ChromiumOptions
from lxml import etree, html as lxml_html

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # 使上级目录 non_shopify_utils 可导入

from non_shopify_utils import (
    clean_body_html,
    format_image_url,
    format_price,
)


MODULE_DIR = Path(__file__).resolve().parent
SHOPIFY_ROOT = MODULE_DIR.parent.parent

DEFAULT_INPUT = str(SHOPIFY_ROOT / "01INPUT_XLSX" / "hespokestyle.com.xlsx")
PROXY_SERVER = "http://127.0.0.1:7897"

CSV_FIELDS = [
    "title", "name", "price1", "price2",
    "styles1", "styles2", "styles3",
    "src_links", "link-href", "details",
]


class ProductExtractionError(ValueError):
    """页面不是可保存的完整 hespokestyle 商品时抛出。"""


def _class_xpath(class_name):
    return (
        'contains(concat(" ", normalize-space(@class), " "), '
        f'" {class_name} ")'
    )


def _first_text(nodes):
    if not nodes:
        return ""
    if isinstance(nodes[0], str):
        return nodes[0].strip()
    return " ".join(nodes[0].itertext()).strip()


def _price_number(text):
    match = re.search(r"\d[\d,]*(?:\.\d+)?", str(text or ""))
    if not match:
        return ""
    return match.group().replace(",", "")


def _inner_html(element):
    return "".join(
        etree.tostring(child, encoding="unicode", method="html")
        for child in element
    )


def _extract_ld_images(document):
    """从 Yoast JSON-LD 图里取所有 ImageObject（当前商品主图）。

    只认 schema 图里的 ``ImageObject``，天然排除页面上的推荐商品图。
    """
    images = []
    for script in document.xpath('//script[@type="application/ld+json"]'):
        text = (script.text_content() or "").strip()
        if not text:
            continue
        try:
            payload = json.loads(text)
        except (ValueError, TypeError):
            continue
        graph = payload.get("@graph", [payload]) if isinstance(payload, dict) else []
        if not isinstance(graph, list):
            continue
        for node in graph:
            if not isinstance(node, dict):
                continue
            node_type = node.get("@type")
            types = node_type if isinstance(node_type, list) else [node_type]
            if "ImageObject" not in types:
                continue
            url = str(node.get("url") or node.get("contentUrl") or "").strip()
            if not url:
                continue
            formatted = format_image_url(url, use_query_marker=True)
            if formatted and formatted not in images:
                images.append(formatted)
    return images


def _extract_gallery_images(document):
    """JSON-LD 缺图时回退到商品 Gallery 容器，跳过 NitroPack 占位图。"""
    images = []
    image_xpath = f'//*[{_class_xpath("product-single__gallery")}]//img'
    for img in document.xpath(image_xpath):
        for attr in ("src", "data-src", "nitro-lazy-src", "data-lazy-src"):
            url = str(img.get(attr) or "").strip()
            if url and not url.startswith("data:"):
                formatted = format_image_url(url, use_query_marker=True)
                if formatted and formatted not in images:
                    images.append(formatted)
                break
    return images


def parse_product_html(source):
    """从已渲染的 hespokestyle 商品 HTML 提取主管线字段。"""
    source = str(source or "")
    lowered = source.lower()

    # 错误页 / 拦截页必须直接失败，不能当成成功商品（开发流程第十四节）
    if "attention required" in lowered or "403 forbidden" in lowered:
        raise ProductExtractionError("页面被 Cloudflare 拦截（403 Forbidden），请检查代理出口")
    if "challenges.cloudflare.com" in lowered or "just a moment" in lowered:
        raise ProductExtractionError("页面被 Cloudflare 人机验证拦截，请检查代理出口")
    if "sucuri" in lowered:
        raise ProductExtractionError("页面被 Sucuri 拦截，请检查代理出口地区")

    try:
        document = lxml_html.fromstring(source)
    except (etree.ParserError, ValueError) as exc:
        raise ProductExtractionError(f"页面 HTML 无法解析: {exc}") from exc

    # 商品名：hero 区第一个 h1.h2（页尾 fixed header 还有一份同文案）
    name = _first_text(document.xpath(f'//h1[{_class_xpath("h2")}]'))

    # 售价：紧随 h1 的 .price；USD 平铺价，无划线价
    price_xpath = (
        f'//h1[{_class_xpath("h2")}]'
        f'/following-sibling::span[{_class_xpath("price")}][1]'
    )
    price_nodes = document.xpath(price_xpath)
    current_raw = _price_number(_first_text(price_nodes[:1]))
    price1 = ""
    if current_raw:
        if "." not in current_raw:
            current_raw = f"{current_raw}.00"
        price1 = format_price(current_raw)
    price2 = price1

    # 尺码：排除 onhand="0"（试穿样衣缺货）的码
    sizes = []
    for option in document.xpath('//select[@id="jacket-size"]/option'):
        if str(option.get("onhand") or "").strip() == "0":
            continue
        value = " ".join(" ".join(option.itertext()).split())
        if value and value not in sizes:
            sizes.append(value)

    # 图片：JSON-LD 主图 -> Gallery 容器
    images = _extract_ld_images(document)
    if not images:
        images = _extract_gallery_images(document)

    # 详情：描述 + specs 选项卡（Fabric 等）+ 定制选项标签
    parts = []
    for el in document.xpath(f'//*[{_class_xpath("product-single__description")}]'):
        parts.append(_inner_html(el))
    for el in document.xpath(f'//*[{_class_xpath("tab__content")}]'):
        parts.append(_inner_html(el))
    labels = []
    for el in document.xpath(f'//*[{_class_xpath("drawer--customize")}]//label'):
        text = " ".join(" ".join(el.itertext()).split())
        if text and text not in labels:
            labels.append(text)
    if labels:
        parts.append("".join(f"<p>{label}</p>" for label in labels))
    details = clean_body_html("".join(parts))

    missing = [
        field
        for field, value in (
            ("name", name),
            ("price1", price1),
            ("src_links", images),
            ("details", details),
        )
        if not value
    ]
    if missing:
        raise ProductExtractionError(f"页面关键字段缺失: {', '.join(missing)}")

    first_image = images[0]
    segments = [
        f"{size}${price1}${price2}@{first_image}"
        for size in sizes
    ]
    styles1 = f"Size#{'#'.join(segments)}" if segments else ""

    return {
        "name": name,
        "price1": price1,
        "price2": price2,
        "styles1": styles1,
        "src_links": "#".join(images),
        "details": details,
    }


def _product_location(url):
    parsed = urlsplit(str(url or ""))
    host = (parsed.hostname or "").lower()
    if host.startswith("www."):
        host = host[4:]
    path = re.sub(r"/+", "/", parsed.path).rstrip("/") or "/"
    return host, path


def wait_for_product(
    tab,
    expected_url=None,
    timeout=30,
    poll_interval=0.25,
):
    """轮询页面渲染结果，直到 10 字段所需数据全部可解析。"""
    if expected_url and _product_location(tab.url) != _product_location(expected_url):
        raise ProductExtractionError(
            f"页面地址不匹配: 期望 {expected_url}，实际 {tab.url}"
        )
    deadline = time.monotonic() + timeout
    last_error = None
    while True:
        try:
            return parse_product_html(tab.html)
        except ProductExtractionError as exc:
            last_error = exc
            # 拦截页不会随着等待自行消失，直接失败
            if "拦截" in str(exc):
                raise

        if time.monotonic() >= deadline:
            raise ProductExtractionError(f"商品数据等待超时: {last_error}")
        time.sleep(poll_interval)


def build_record(title, url, product):
    values = {
        "title": str(title or ""),
        "link-href": str(url or ""),
        "styles2": "",
        "styles3": "",
    }
    values.update(product or {})
    return {field: values.get(field, "") for field in CSV_FIELDS}


def is_record_complete(record):
    required = ("name", "price1", "price2", "src_links", "link-href", "details")
    return all(str(record.get(field, "") or "").strip() for field in required)


def to_output_path(input_path):
    path = Path(input_path)
    parts = list(path.parts)
    try:
        input_index = parts.index("01INPUT_XLSX")
    except ValueError as exc:
        raise ValueError("输入文件必须位于 01INPUT_XLSX 目录内，避免覆盖原文件") from exc
    parts[input_index] = "02OUTPUT_XLSX"
    output = Path(*parts)
    if os.path.normcase(os.path.abspath(output)) == os.path.normcase(os.path.abspath(path)):
        raise ValueError("输出路径不能与输入路径相同")
    return str(output)


def load_tasks(input_path):
    workbook = openpyxl.load_workbook(input_path, read_only=True)
    try:
        worksheet = workbook.active
        tasks = []
        for title, url in worksheet.iter_rows(
            min_row=2, min_col=1, max_col=2, values_only=True
        ):
            clean_title = str(title).strip() if title else ""
            clean_url = str(url).strip() if url else ""
            if clean_url.startswith(("http://", "https://")):
                tasks.append((clean_title, clean_url))
        return tasks
    finally:
        workbook.close()


def load_done_links(output_path):
    path = Path(output_path)
    if not path.exists():
        return set(), []

    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        worksheet = workbook.active
        header = [
            str(cell.value or "").strip()
            for cell in worksheet[1][:len(CSV_FIELDS)]
        ]
        if header != CSV_FIELDS:
            raise ValueError("现有输出文件表头与 10 字段契约不一致")

        rows = []
        done = set()
        for values in worksheet.iter_rows(
            min_row=2, min_col=1, max_col=len(CSV_FIELDS), values_only=True
        ):
            row = {
                field: (values[index] if values[index] is not None else "")
                for index, field in enumerate(CSV_FIELDS)
            }
            if not str(row["link-href"] or "").strip():
                continue
            rows.append(row)
            if is_record_complete(row):
                done.add(str(row["link-href"]).strip())
        return done, rows
    finally:
        workbook.close()


def upsert_row(rows, record):
    url = str(record.get("link-href", "") or "").strip()
    for index, existing in enumerate(rows):
        if str(existing.get("link-href", "") or "").strip() == url:
            rows[index] = record
            return
    rows.append(record)


def save_xlsx(rows, output_path):
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(CSV_FIELDS)
    for row in rows:
        worksheet.append([row.get(field, "") for field in CSV_FIELDS])

    temp_handle = tempfile.NamedTemporaryFile(
        dir=path.parent,
        prefix=f".{path.stem}-",
        suffix=".tmp.xlsx",
        delete=False,
    )
    temp_path = Path(temp_handle.name)
    temp_handle.close()
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        workbook.close()
        if temp_path.exists():
            temp_path.unlink()


def create_browser(log):
    options = ChromiumOptions()
    options.incognito(on_off=True)
    options.set_argument(f"--proxy-server={PROXY_SERVER}")
    port = random.randint(9222, 9322)
    options.set_local_port(port)
    for attempt in range(1, 4):
        try:
            chrome = Chromium(options)
            log(f"浏览器启动成功，端口: {port}")
            return chrome
        except Exception as exc:
            log(f"端口 {port} 启动失败（{attempt}/3）: {exc}")
            port = random.randint(9222, 9322)
            options.set_local_port(port)
    raise RuntimeError("无法启动浏览器，已重试 3 次")


class HeSpokeStyleGUI:
    def __init__(self, root):
        self.root = root
        root.title("hespokestyle.com 采集工具（仅采集）")
        root.geometry("780x540")

        top = tk.Frame(root)
        top.pack(fill="x", padx=8, pady=6)
        tk.Label(top, text="输入文件:").pack(side="left")
        self.path_var = tk.StringVar(value=DEFAULT_INPUT)
        tk.Entry(top, textvariable=self.path_var).pack(
            side="left", fill="x", expand=True, padx=6
        )
        tk.Button(top, text="浏览...", command=self.pick_file).pack(side="left")

        buttons = tk.Frame(root)
        buttons.pack(fill="x", padx=8, pady=4)
        self.start_btn = tk.Button(
            buttons, text="开始采集", width=14, command=self.start
        )
        self.start_btn.pack(side="left")
        self.stop_btn = tk.Button(
            buttons, text="停止", width=10, state="disabled", command=self.stop
        )
        self.stop_btn.pack(side="left", padx=8)
        self.progress_var = tk.StringVar(value="待开始")
        tk.Label(buttons, textvariable=self.progress_var).pack(side="left", padx=12)

        self.log_box = scrolledtext.ScrolledText(root, height=25)
        self.log_box.pack(fill="both", expand=True, padx=8, pady=6)

        self.stop_event = threading.Event()
        self.worker = None

    def pick_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if path:
            self.path_var.set(path)

    def log(self, message):
        def append():
            self.log_box.insert("end", str(message) + "\n")
            self.log_box.see("end")

        self.root.after(0, append)

    def set_progress(self, message):
        self.root.after(0, self.progress_var.set, message)

    def start(self):
        input_path = self.path_var.get().strip()
        if not os.path.exists(input_path):
            messagebox.showerror("错误", f"输入文件不存在:\n{input_path}")
            return
        try:
            to_output_path(input_path)
        except ValueError as exc:
            messagebox.showerror("错误", str(exc))
            return

        self.stop_event.clear()
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="normal")
        self.worker = threading.Thread(
            target=self.run, args=(input_path,), daemon=True
        )
        self.worker.start()

    def stop(self):
        self.stop_event.set()
        self.log("→ 已请求停止，等待当前商品完成...")

    def run(self, input_path):
        try:
            self._run(input_path)
        except Exception as exc:
            self.log(f"[异常] {exc}")
        finally:
            self.root.after(0, self.start_btn.config, {"state": "normal"})
            self.root.after(0, self.stop_btn.config, {"state": "disabled"})

    def _run(self, input_path):
        output_path = to_output_path(input_path)
        tasks = load_tasks(input_path)
        done, rows = load_done_links(output_path)
        pending = sum(1 for _, url in tasks if url not in done)
        self.log(f"输入: {input_path}")
        self.log(f"输出: {output_path}")
        self.log(f"任务 {len(tasks)} 条，完整记录 {len(done)} 条，本次待采 {pending} 条")

        chrome = create_browser(self.log)
        tab = chrome.latest_tab
        processed = 0
        succeeded = 0
        try:
            for index, (title, url) in enumerate(tasks, 1):
                if self.stop_event.is_set():
                    self.log("→ 已停止")
                    break
                if url in done:
                    continue

                self.set_progress(f"{index}/{len(tasks)}")
                record = build_record(title, url, {})
                error = None
                for attempt in (1, 2):
                    try:
                        tab.get(url)
                        product = wait_for_product(
                            tab,
                            expected_url=url,
                            timeout=30,
                        )
                        record = build_record(title, url, product)
                        error = None
                        break
                    except Exception as exc:
                        error = exc
                        self.log(f"  [重试 {attempt}/2] {url}\n    {exc}")

                upsert_row(rows, record)
                save_xlsx(rows, output_path)
                processed += 1
                if error is None and is_record_complete(record):
                    done.add(url)
                    succeeded += 1
                    variant_count = max(0, len(record["styles1"].split("#")) - 1)
                    self.log(
                        f"  [{index}/{len(tasks)}] {record['name']} | "
                        f"${record['price1']} | {variant_count}个尺码"
                    )
                else:
                    self.log(f"  [失败占位，后续会重采] {url}: {error}")
        finally:
            chrome.quit()

        self.log(
            f"完成: 本次处理 {processed} 条，成功 {succeeded} 条，"
            f"输出共 {len(rows)} 条 → {output_path}"
        )


if __name__ == "__main__":
    root = tk.Tk()
    HeSpokeStyleGUI(root)
    root.mainloop()
