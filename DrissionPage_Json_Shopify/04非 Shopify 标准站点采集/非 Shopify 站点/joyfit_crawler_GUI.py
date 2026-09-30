# -*- coding: utf-8 -*-
"""joyfit.store 商品采集工具（Shopline 平台 + DrissionPage）。

Shopline 站点没有可用的 .json 接口（/products/*.json 返回 403），商品数据全部
内嵌在服务端渲染的 HTML 中：
  - ld+json ``@type=Product``：name / description / sku / offers（价格+币种）
  - ``<script name="variant-data">``：全部变体（options=[颜色,尺码] + price 分 + featured_media_id）
  - 媒体画廊 ``li[data-media-id]``：media-id 与缩略图，用于 颜色→图 映射

直接解析渲染后的 HTML 即可，无需点选交互；输出与主管线保持相同的 10 字段契约。
"""

import json
import os
import random
import tempfile
import threading
import time
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext

import openpyxl
from DrissionPage import Chromium, ChromiumOptions
from lxml import etree, html as lxml_html

MODULE_DIR = Path(__file__).resolve().parent
SHOPIFY_ROOT = MODULE_DIR.parent.parent

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # 使上级目录 non_shopify_utils 可导入

from non_shopify_utils import (
    clean_body_html,
    fetch_exchange_rates,
    format_image_url,
    format_price,
)

DEFAULT_INPUT = str(SHOPIFY_ROOT / "01INPUT_XLSX" / "joyfit.store.xlsx")
PROXY_SERVER = "http://127.0.0.1:7897"

CSV_FIELDS = [
    "title", "name", "price1", "price2",
    "styles1", "styles2", "styles3",
    "src_links", "link-href", "details",
]


class ProductExtractionError(ValueError):
    """页面不是可保存的完整 Joyfit 商品时抛出。"""


def _script_json(document, name):
    """按 script 的 name 属性取内嵌 JSON 并解析，返回对象/数组或 None。"""
    for script in document.xpath(f'//script[@name="{name}"]'):
        text = "".join(script.itertext()).strip()
        if not text:
            continue
        try:
            return json.loads(text)
        except ValueError:
            continue
    return None


def _ldjson_product(document):
    """取页面中 @type=Product 的 ld+json 结构化数据。"""
    for script in document.xpath('//script[@type="application/ld+json"]'):
        text = "".join(script.itertext()).strip()
        if not text:
            continue
        try:
            data = json.loads(text)
        except ValueError:
            continue
        if isinstance(data, dict) and data.get("@type") == "Product":
            return data
    return {}


def _inner_html(element):
    return "".join(
        etree.tostring(child, encoding="unicode", method="html")
        for child in element
    )


def parse_product_html(rendered_html, rates):
    """从渲染后的 Joyfit 商品 HTML 中提取主管线字段。"""
    source = str(rendered_html or "")
    try:
        document = lxml_html.fromstring(source)
    except (etree.ParserError, ValueError) as exc:
        raise ProductExtractionError(f"页面 HTML 无法解析: {exc}") from exc

    product = _ldjson_product(document)
    variants = _script_json(document, "variant-data") or []

    # 标题: ld+json name -> h1
    name = str(product.get("name", "") or "").strip()
    if not name:
        h1 = document.xpath('//h1[contains(@class,"product-detail__title")]')
        name = " ".join(" ".join(h1[0].itertext()).split()) if h1 else ""

    # 价格: ld+json offers[0]（主币单位）-> og:price:amount -> variant-data 分
    currency = "USD"
    price_raw = ""
    offers = product.get("offers") or []
    if offers and isinstance(offers[0], dict):
        price_raw = str(offers[0].get("price", "") or "").strip()
        currency = str(offers[0].get("priceCurrency", "") or "").strip() or "USD"
    if not price_raw:
        meta = document.xpath('//meta[@property="og:price:amount"]/@content')
        price_raw = str(meta[0]).strip() if meta else ""
    if not price_raw and variants:
        price_raw = str(variants[0].get("price", "") or "").strip()
    price1 = format_price(price_raw, rates=rates, currency=currency) if price_raw else ""
    # 页面无划线价，price2 与 price1 相同
    price2 = price1

    # 颜色/尺码/颜色->featured_media_id（顺序按变体出现先后）
    colors, sizes = [], []
    color_media = {}
    for variant in variants:
        options = variant.get("options")
        if not isinstance(options, list) or not options:
            continue
        color = str(options[0]).strip() if len(options) > 0 else ""
        size = str(options[1]).strip() if len(options) > 1 else ""
        if color and color not in colors:
            colors.append(color)
        if size and size not in sizes:
            sizes.append(size)
        media_id = str(variant.get("featured_media_id", "") or "").strip()
        if color and media_id and color not in color_media:
            color_media[color] = media_id

    # 媒体画廊: media-id -> 图片 URL（去 query，避免同图不同尺寸参数重复）
    media_src = {}
    for item in document.xpath('//li[@data-media-id]'):
        media_id = str(item.get("data-media-id", "") or "").strip()
        srcs = item.xpath('.//img[@src]/@src')
        if media_id and srcs:
            media_src[media_id] = str(srcs[0]).split("?")[0]

    # 全部商品图（按画廊顺序去重）
    images = []
    for item in document.xpath('//li[@data-media-id]'):
        for raw in item.xpath('.//img[@src]/@src'):
            formatted = format_image_url(str(raw).split("?")[0])
            if formatted and formatted not in images:
                images.append(formatted)

    # 颜色 -> 图片（featured_media_id 桥接）
    color_img = {}
    for color in colors:
        src = media_src.get(color_media.get(color, ""), "")
        if src:
            color_img[color] = format_image_url(src)

    # 详情: 描述容器 inner HTML -> ld+json description
    details = ""
    desc_nodes = document.xpath(
        '//div[contains(concat(" ", normalize-space(@class), " "), '
        '" product-expander__content ")]'
    )
    if desc_nodes:
        details = clean_body_html(_inner_html(desc_nodes[0]))
    elif product.get("description"):
        details = clean_body_html(f"<p>{product['description']}</p>")

    missing = [
        field
        for field, value in (
            ("name", name),
            ("price1", price1),
            ("src_links", images),
            ("details", details),
        )
        if not value
    ]
    if missing:
        raise ProductExtractionError(f"页面关键字段缺失: {', '.join(missing)}")

    # styles1: 颜色×尺码 笛卡尔积，每个颜色用自身缩略图
    first_img = images[0] if images else ""
    segments = []
    for color in colors:
        img = color_img.get(color, first_img)
        for size in sizes:
            seg = f"{color}&Size&{size}${price1}${price2}"
            if img:
                seg = f"{seg}@{img}"
            segments.append(seg)
    styles1 = f"Color#{'#'.join(segments)}" if colors and segments else ""

    return {
        "name": name,
        "price1": price1,
        "price2": price2,
        "styles1": styles1,
        "src_links": "#".join(images),
        "details": details,
    }


def wait_for_product(tab, rates, timeout=30, poll_interval=0.3):
    """轮询渲染结果，直到十字段所需数据全部可解析。"""
    deadline = time.monotonic() + timeout
    last_error = None
    while True:
        try:
            return parse_product_html(tab.html, rates)
        except ProductExtractionError as exc:
            last_error = exc

        if time.monotonic() >= deadline:
            raise ProductExtractionError(f"商品数据等待超时: {last_error}")
        time.sleep(poll_interval)


def build_record(title, url, product):
    values = {
        "title": str(title or ""),
        "link-href": str(url or ""),
        "styles2": "",
        "styles3": "",
    }
    values.update(product or {})
    return {field: values.get(field, "") for field in CSV_FIELDS}


def is_record_complete(record):
    required = ("name", "price1", "price2", "src_links", "link-href", "details")
    return all(str(record.get(field, "") or "").strip() for field in required)


def to_output_path(input_path):
    path = Path(input_path)
    parts = list(path.parts)
    try:
        input_index = parts.index("01INPUT_XLSX")
    except ValueError as exc:
        raise ValueError("输入文件必须位于 01INPUT_XLSX 目录内，避免覆盖原文件") from exc
    parts[input_index] = "02OUTPUT_XLSX"
    output = Path(*parts)
    if os.path.normcase(os.path.abspath(output)) == os.path.normcase(os.path.abspath(path)):
        raise ValueError("输出路径不能与输入路径相同")
    return str(output)


def load_tasks(input_path):
    workbook = openpyxl.load_workbook(input_path, read_only=True)
    try:
        worksheet = workbook.active
        tasks = []
        for title, url in worksheet.iter_rows(
            min_row=2, min_col=1, max_col=2, values_only=True
        ):
            clean_title = str(title).strip() if title else ""
            clean_url = str(url).strip() if url else ""
            if clean_url.startswith(("http://", "https://")):
                tasks.append((clean_title, clean_url))
        return tasks
    finally:
        workbook.close()


def load_done_links(output_path):
    path = Path(output_path)
    if not path.exists():
        return set(), []

    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        worksheet = workbook.active
        header = [
            str(cell.value or "").strip()
            for cell in worksheet[1][:len(CSV_FIELDS)]
        ]
        if header != CSV_FIELDS:
            raise ValueError("现有输出文件表头与 10 字段契约不一致")

        rows = []
        done = set()
        for values in worksheet.iter_rows(
            min_row=2, min_col=1, max_col=len(CSV_FIELDS), values_only=True
        ):
            row = {
                field: (values[index] if values[index] is not None else "")
                for index, field in enumerate(CSV_FIELDS)
            }
            if not str(row["link-href"] or "").strip():
                continue
            rows.append(row)
            if is_record_complete(row):
                done.add(str(row["link-href"]).strip())
        return done, rows
    finally:
        workbook.close()


def upsert_row(rows, record):
    url = str(record.get("link-href", "") or "").strip()
    for index, existing in enumerate(rows):
        if str(existing.get("link-href", "") or "").strip() == url:
            rows[index] = record
            return
    rows.append(record)


def save_xlsx(rows, output_path):
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(CSV_FIELDS)
    for row in rows:
        worksheet.append([row.get(field, "") for field in CSV_FIELDS])

    temp_handle = tempfile.NamedTemporaryFile(
        dir=path.parent,
        prefix=f".{path.stem}-",
        suffix=".tmp.xlsx",
        delete=False,
    )
    temp_path = Path(temp_handle.name)
    temp_handle.close()
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        workbook.close()
        if temp_path.exists():
            temp_path.unlink()


def create_browser(log):
    options = ChromiumOptions()
    options.incognito(on_off=True)
    options.set_argument(f"--proxy-server={PROXY_SERVER}")
    port = random.randint(9222, 9322)
    options.set_local_port(port)
    for attempt in range(1, 4):
        try:
            chrome = Chromium(options)
            log(f"浏览器启动成功，端口: {port}")
            return chrome
        except Exception as exc:
            log(f"端口 {port} 启动失败（{attempt}/3）: {exc}")
            port = random.randint(9222, 9322)
            options.set_local_port(port)
    raise RuntimeError("无法启动浏览器，已重试 3 次")


class JoyfitGUI:
    def __init__(self, root):
        self.root = root
        root.title("joyfit.store 采集工具（仅采集）")
        root.geometry("780x540")

        top = tk.Frame(root)
        top.pack(fill="x", padx=8, pady=6)
        tk.Label(top, text="输入文件:").pack(side="left")
        self.path_var = tk.StringVar(value=DEFAULT_INPUT)
        tk.Entry(top, textvariable=self.path_var).pack(
            side="left", fill="x", expand=True, padx=6
        )
        tk.Button(top, text="浏览...", command=self.pick_file).pack(side="left")

        buttons = tk.Frame(root)
        buttons.pack(fill="x", padx=8, pady=4)
        self.start_btn = tk.Button(
            buttons, text="开始采集", width=14, command=self.start
        )
        self.start_btn.pack(side="left")
        self.stop_btn = tk.Button(
            buttons, text="停止", width=10, state="disabled", command=self.stop
        )
        self.stop_btn.pack(side="left", padx=8)
        self.progress_var = tk.StringVar(value="待开始")
        tk.Label(buttons, textvariable=self.progress_var).pack(side="left", padx=12)

        self.log_box = scrolledtext.ScrolledText(root, height=25)
        self.log_box.pack(fill="both", expand=True, padx=8, pady=6)

        self.stop_event = threading.Event()
        self.worker = None

    def pick_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if path:
            self.path_var.set(path)

    def log(self, message):
        def append():
            self.log_box.insert("end", str(message) + "\n")
            self.log_box.see("end")

        self.root.after(0, append)

    def set_progress(self, message):
        self.root.after(0, self.progress_var.set, message)

    def start(self):
        input_path = self.path_var.get().strip()
        if not os.path.exists(input_path):
            messagebox.showerror("错误", f"输入文件不存在:\n{input_path}")
            return
        try:
            to_output_path(input_path)
        except ValueError as exc:
            messagebox.showerror("错误", str(exc))
            return

        self.stop_event.clear()
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="normal")
        self.worker = threading.Thread(
            target=self.run, args=(input_path,), daemon=True
        )
        self.worker.start()

    def stop(self):
        self.stop_event.set()
        self.log("→ 已请求停止，等待当前商品完成...")

    def run(self, input_path):
        try:
            self._run(input_path)
        except Exception as exc:
            self.log(f"[异常] {exc}")
        finally:
            self.root.after(0, self.start_btn.config, {"state": "normal"})
            self.root.after(0, self.stop_btn.config, {"state": "disabled"})

    def _run(self, input_path):
        output_path = to_output_path(input_path)
        tasks = load_tasks(input_path)
        done, rows = load_done_links(output_path)
        pending = sum(1 for _, url in tasks if url not in done)
        self.log(f"输入: {input_path}")
        self.log(f"输出: {output_path}")
        self.log(f"任务 {len(tasks)} 条，完整记录 {len(done)} 条，本次待采 {pending} 条")

        rates = fetch_exchange_rates()

        chrome = create_browser(self.log)
        tab = chrome.latest_tab
        processed = 0
        succeeded = 0
        try:
            for index, (title, url) in enumerate(tasks, 1):
                if self.stop_event.is_set():
                    self.log("→ 已停止")
                    break
                if url in done:
                    continue

                self.set_progress(f"{index}/{len(tasks)}")
                record = build_record(title, url, {})
                error = None
                for attempt in (1, 2):
                    try:
                        tab.get(url)
                        product = wait_for_product(tab, rates, timeout=30)
                        record = build_record(title, url, product)
                        error = None
                        break
                    except Exception as exc:
                        error = exc
                        self.log(f"  [重试 {attempt}/2] {url}\n    {exc}")

                upsert_row(rows, record)
                save_xlsx(rows, output_path)
                processed += 1
                if error is None and is_record_complete(record):
                    done.add(url)
                    succeeded += 1
                    image_count = len(record["src_links"].split("#"))
                    variant_count = max(0, len(record["styles1"].split("#")) - 1)
                    self.log(
                        f"  [{index}/{len(tasks)}] {record['name']} | "
                        f"${record['price1']} | {image_count}图 | {variant_count}个规格组合"
                    )
                else:
                    self.log(f"  [失败占位，后续会重采] {url}: {error}")
        finally:
            chrome.quit()

        self.log(
            f"完成: 本次处理 {processed} 条，成功 {succeeded} 条，"
            f"输出共 {len(rows)} 条 → {output_path}"
        )


if __name__ == "__main__":
    root = tk.Tk()
    JoyfitGUI(root)
    root.mainloop()
