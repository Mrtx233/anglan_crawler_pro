from __future__ import annotations

from pathlib import Path
from urllib.parse import urlparse
import openpyxl
import os
import tempfile
from ..resources.columns import CSV_FIELDS
from ..storage.paths import to_shopify_json_url


def read_targets_from_excel(filepath):
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    targets = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        title = str(row[0]).strip() if row and row[0] else ""
        url = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if url:
            targets.append((title, url))
    wb.close()
    return targets


def save_workbook_atomic(workbook, output_path):
    """完整写入同目录临时文件后替换，避免写入失败破坏已有工作簿。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_path = tempfile.mkstemp(suffix=".xlsx", dir=output_path.parent)
    os.close(fd)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        workbook.close()
        if os.path.exists(temporary_path):
            os.unlink(temporary_path)


def save_xlsx(rows, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(CSV_FIELDS)
    for row in rows:
        ws.append([row.get(f, "") for f in CSV_FIELDS])
    save_workbook_atomic(wb, output_path)


def load_xlsx(filepath):
    wb = openpyxl.load_workbook(filepath, read_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h) if h else "" for h in next(rows_iter, ())]
        results = [
            {h: (v if v is not None else "") for h, v in zip(headers, row)} for row in rows_iter
        ]
        return headers, results
    finally:
        wb.close()


def load_existing_link_rows(output_folder):
    rows = []
    seen = set()
    for path in sorted(Path(output_folder).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        workbook = openpyxl.load_workbook(path, read_only=True)
        try:
            values = workbook.active.iter_rows(values_only=True)
            headers = next(values, ())
            if tuple(headers[:2]) != ("title", "link"):
                raise ValueError(f"已有链接文件列格式不正确（需要 title/link）: {path}")
            for row in values:
                link = str(row[1]).strip() if len(row) > 1 and row[1] else ""
                if not link:
                    continue
                key = to_shopify_json_url(link)
                if key not in seen:
                    seen.add(key)
                    rows.append({"title": str(row[0] or ""), "link": link})
        finally:
            workbook.close()
    return rows


def write_link_workbooks(
    rows: list[dict[str, str]],
    output_folder: Path,
) -> list[Path]:
    """按商品域名写入第一阶段的 title/link 工作簿。"""
    domain_groups: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        domain = urlparse(row["link"]).netloc
        if domain:
            domain_groups.setdefault(domain, []).append(row)

    output_folder.mkdir(parents=True, exist_ok=True)
    saved_paths: list[Path] = []
    for domain, domain_rows in domain_groups.items():
        file_path = output_folder / f"{domain}.xlsx"
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.append(["title", "link"])
        for row in domain_rows:
            worksheet.append([row.get("title", ""), row.get("link", "")])
        save_workbook_atomic(workbook, file_path)
        saved_paths.append(file_path)
    return saved_paths
