# ======================== 文件读写层 ========================
#
# 职责: 输入 Excel 读取、结果 XLSX 保存/加载、安全输出路径构造、
#       恢复文件命名。
# 依赖: config（CSV_FIELDS / RECOVERY 命名）、openpyxl、pathlib。
# 禁止: 导入 tkinter、DrissionPage、converter。
#
# 安全约定（本版强化）:
#   1. build_output_file_path 保证输出路径绝不等于输入路径，
#      相同时抛 RuntimeError 阻止覆盖输入文件。
#   2. 旧 XLSX 读取失败时，调用方必须改用 build_recovery_xlsx_path
#      另存，绝不覆盖原文件。
#   3. save_xlsx 采用「同目录临时文件 + os.replace」原子写入，
#      失败清理临时文件并抛出，原文件不受影响。
# ============================================================

import os
import tempfile
import time
from pathlib import Path

import openpyxl

import config

# 恢复文件名中的标记
RECOVERY_MARK = "recovery"


def read_targets_from_excel(filepath):
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    targets = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        title = str(row[0]).strip() if row and row[0] else ""
        url = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if url:
            targets.append((title, url))
    wb.close()
    return targets


def save_xlsx(rows, output_path):
    """原子写入 XLSX：先写同目录临时文件，成功后 os.replace 覆盖目标。

    签名与原实现保持一致，所有调用方无需改动。
    失败时清理临时文件并抛出异常；原文件不会被破坏。
    """
    output_path = Path(output_path)
    target_dir = output_path.parent if str(output_path.parent) else Path(".")
    target_dir.mkdir(parents=True, exist_ok=True)

    fd, tmp_path = tempfile.mkstemp(
        prefix=f".{output_path.stem}_tmp_",
        suffix=".xlsx",
        dir=str(target_dir),
    )
    os.close(fd)
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(config.CSV_FIELDS)
        for row in rows:
            ws.append([row.get(f, "") for f in config.CSV_FIELDS])
        wb.save(tmp_path)
        wb.close()
        os.replace(tmp_path, str(output_path))
    except Exception:
        try:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
        except OSError:
            pass
        raise


def load_xlsx(filepath):
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(rows_iter)]
    results = []
    for row in rows_iter:
        results.append({h: (v or "") for h, v in zip(headers, row)})
    wb.close()
    return headers, results


def _replace_path_component(
    path,
    source_name=config.INPUT_DIRNAME,
    target_name=config.OUTPUT_DIRNAME,
):
    """仅替换完整目录名，避免字符串 replace 误改其它路径片段。"""
    path_obj = Path(path)
    parts = list(path_obj.parts)
    for index, part in enumerate(parts):
        if part.lower() == source_name.lower():
            parts[index] = target_name
            return Path(*parts), True
    return path_obj, False


def build_output_file_path(input_file):
    """构造安全输出路径，保证永远不会与输入 XLSX 为同一路径。"""
    input_path = Path(input_file)
    replaced, found = _replace_path_component(input_path)
    if found:
        output_path = replaced
    else:
        # 用户选择的目录不符合 01INPUT_XLSX 约定时，也绝不原地覆盖输入。
        output_path = input_path.parent / config.OUTPUT_DIRNAME / input_path.name

    try:
        same_path = output_path.resolve(strict=False) == input_path.resolve(strict=False)
    except OSError:
        same_path = os.path.abspath(str(output_path)) == os.path.abspath(str(input_path))
    if same_path:
        raise RuntimeError(f"输出路径与输入文件相同，已阻止覆盖: {input_path}")
    return str(output_path)


def build_output_folder_path(input_folder):
    """构造合并文件使用的安全输出目录。"""
    input_path = Path(input_folder)
    replaced, found = _replace_path_component(input_path)
    if found:
        return str(replaced)
    return str(input_path / "02OUTPUT_XLSX")


def build_recovery_xlsx_path(output_file):
    """旧 XLSX 无法读取时，为本轮数据生成绝不覆盖旧文件的恢复文件名。"""
    output_path = Path(output_file)
    timestamp = time.strftime("%Y%m%d-%H%M%S")
    candidate = output_path.with_name(
        f"{output_path.stem}.{RECOVERY_MARK}-{timestamp}{output_path.suffix}"
    )
    counter = 1
    while candidate.exists():
        candidate = output_path.with_name(
            f"{output_path.stem}.{RECOVERY_MARK}-{timestamp}-{counter}{output_path.suffix}"
        )
        counter += 1
    return str(candidate)