# ======================== 详情页链接采集 GUI ========================
#
# 用途: 从 Shopify 列表页提取产品链接
# 原理: DrissionPage 浏览器加载页面，从 <script> JSON 中提取 handle
#
# 界面:
#   - 分类列表文本框（每行: 标题, URL）
#   - 最大页数输入框（留空=不限制）
#   - 保存路径选择
#   - 开始/停止按钮 + 实时日志
#
# 输出: XLSX（title, link）
# ========================================================================

import json
import os
import random
import re
import sys
import threading
import time
from urllib.parse import parse_qsl, urlencode, urlparse, urlunparse

import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext

import openpyxl
from DrissionPage import Chromium, ChromiumOptions
from DrissionPage.errors import ContextLostError

# ======================== 样式配置 ========================
UI_COLORS = {
    "background": "#F4F7FB",
    "card": "#FFFFFF",
    "card_alt": "#F8FAFC",
    "border": "#E2E8F0",
    "text": "#0F172A",
    "text_secondary": "#475569",
    "text_muted": "#94A3B8",
    "primary": "#2563EB",
    "primary_hover": "#1D4ED8",
    "danger": "#DC2626",
    "danger_light": "#FEF2F2",
    "log_background": "#0F172A",
    "log_panel": "#111827",
    "log_text": "#D1D5DB",
}

UI_FONT = "Microsoft YaHei UI"
MONO_FONT = "Consolas"

MAX_RETRIES = 3
RETRY_DELAY = 2.0
PAGE_SETTLE_WAIT = 0.35


# ======================== 浏览器初始化 ========================
def create_browser():
    co = ChromiumOptions()
    co.incognito(on_off=True)
    # 只采集 HTML/JSON/链接，关闭图片加载可明显减少列表页等待时间。
    co.set_argument("--blink-settings=imagesEnabled=false")
    co.set_argument("--proxy-server=http://127.0.0.1:7897")
    random_port = random.randint(9222, 9322)
    co.set_local_port(random_port)

    for attempt in range(3):
        try:
            chrome = Chromium(co)
            print(f"浏览器启动成功，端口: {random_port}")
            return chrome
        except Exception as e:
            print(f"端口 {random_port} 启动失败 (尝试 {attempt + 1}/3): {e}")
            random_port = random.randint(9222, 9322)
            co.set_local_port(random_port)
    raise RuntimeError("无法启动浏览器，已重试多次")


# ======================== 分页采集核心 ========================
def extract_product_links(html_text, base_url):
    """从当前分页 HTML 提取 Shopify 商品链接并按出现顺序去重。"""
    base_parsed = urlparse(base_url)
    base = f"{base_parsed.scheme}://{base_parsed.netloc}"
    seen_handles = set()
    unique_handles = []

    def add_handle(handle):
        handle = str(handle or "").strip().strip("/")
        if handle and handle not in seen_handles:
            seen_handles.add(handle)
            unique_handles.append(handle)

    # 1) DOM 中真实存在的 /products/... 链接，优先级最高。
    href_pattern = r'href=["\'](?:https?://[^/"\']+)?(/products/[^"\'?#/]+)'
    for href in re.findall(href_pattern, html_text, re.IGNORECASE):
        add_handle(href.rstrip("/").split("/")[-1])

    # 2) Shopify 常见的 var meta.products。
    meta_match = re.search(
        r"var\s+meta\s*=\s*(\{.+?\});\s*(?:\n|</script>)",
        html_text,
        re.DOTALL,
    )
    if meta_match:
        try:
            meta_data = json.loads(meta_match.group(1))
            for product in meta_data.get("products", []):
                add_handle(product.get("handle", ""))
        except (json.JSONDecodeError, TypeError, KeyError):
            pass

    # 3) 其他 JSON 中的 handle：只接受页面中同时出现精确商品路径的 handle，
    #    避免 handle=abc 错误匹配 /products/abc-123。
    for handle in re.findall(r'"handle"\s*:\s*"([^"]+)"', html_text):
        path_pattern = rf'/products/{re.escape(handle)}(?=[/?#"\'<>\s]|$)'
        if re.search(path_pattern, html_text, re.IGNORECASE):
            add_handle(handle)

    return [f"{base}/products/{handle}" for handle in unique_handles]


def _run_js_with_retry(tab, script, default=None, retries=2, stop_event=None):
    """执行 JavaScript；重试等待可被停止按钮立即打断。"""
    for attempt in range(retries + 1):
        if stop_event is not None and stop_event.is_set():
            return default
        try:
            result = tab.run_js(script)
            return default if result is None else result
        except ContextLostError:
            if attempt >= retries:
                return default
            delay = 0.6
        except Exception as e:
            if attempt >= retries:
                print(f"  JavaScript 执行失败: {e}")
                return default
            delay = 0.3

        if stop_event is not None:
            if stop_event.wait(delay):
                return default
        else:
            time.sleep(delay)
    return default


def extract_links_by_xpath(tab, xpath, stop_event=None):
    """用 XPath 提取商品详情页链接，只保留 /products/...。"""
    result = _run_js_with_retry(
        tab,
        f"""
        return (() => {{
            const xpath = {json.dumps(xpath)};
            const results = [];
            try {{
                const snapshot = document.evaluate(
                    xpath, document, null,
                    XPathResult.ORDERED_NODE_SNAPSHOT_TYPE, null
                );
                for (let i = 0; i < snapshot.snapshotLength; i++) {{
                    const node = snapshot.snapshotItem(i);
                    const val = node.nodeType === 2 ? node.value : (node.href || node.textContent || '');
                    if (val) results.push(val.trim());
                }}
            }} catch(e) {{
                return {{__xpath_error__: String(e)}};
            }}
            return results;
        }})();
        """,
        default=[],
        stop_event=stop_event,
    )

    if isinstance(result, dict) and result.get("__xpath_error__"):
        raise ValueError(f"XPath 无效: {result['__xpath_error__']}")
    if not isinstance(result, list):
        return []

    base_parsed = urlparse(tab.url)
    base = f"{base_parsed.scheme}://{base_parsed.netloc}"
    seen = set()
    links = []

    for raw in result:
        raw = str(raw).strip()
        if not raw:
            continue
        if raw.startswith("/"):
            full = f"{base}{raw}"
        elif raw.startswith("http"):
            full = raw
        else:
            continue

        parsed = urlparse(full)
        match = re.match(r"^/products/([^/?#]+)", parsed.path, re.IGNORECASE)
        if not match:
            continue
        normalized = f"{base}/products/{match.group(1)}"
        if normalized not in seen:
            seen.add(normalized)
            links.append(normalized)

    return links


def collect_page_links(tab, stop_event=None, xpath=""):
    """采集当前分页；停止时立即放弃当前页，不返回本页结果。"""
    if stop_event is not None and stop_event.is_set():
        return []

    # tab.get() 已经完成主要加载，这里只留很短的页面稳定时间。
    if stop_event is not None:
        if stop_event.wait(PAGE_SETTLE_WAIT):
            return []
    else:
        time.sleep(PAGE_SETTLE_WAIT)

    if xpath:
        links = extract_links_by_xpath(tab, xpath, stop_event=stop_event)
    else:
        html_text = _run_js_with_retry(
            tab,
            "return document.documentElement.outerHTML;",
            default="",
            stop_event=stop_event,
        )
        if not html_text:
            return []
        links = extract_product_links(html_text, tab.url)

    # 停止按钮可能在本页解析过程中被按下：明确丢弃本页结果。
    if stop_event is not None and stop_event.is_set():
        return []
    return links


def build_page_url(base_url, page_num):
    """安全设置或替换 page 参数，避免重复 ?page= 或 fragment 拼接错误。"""
    if page_num <= 1:
        return base_url

    parsed = urlparse(base_url)
    query = dict(parse_qsl(parsed.query, keep_blank_values=True))
    query["page"] = str(page_num)
    return urlunparse(parsed._replace(query=urlencode(query)))


# ======================== 文本重定向 ========================
class TextRedirector:
    def __init__(self, text_widget):
        self.text_widget = text_widget

    def write(self, string):
        if not string:
            return
        try:
            self.text_widget.after(0, self._write, string)
        except (tk.TclError, RuntimeError):
            pass

    def _write(self, string):
        try:
            self.text_widget.config(state=tk.NORMAL)
            self.text_widget.insert(tk.END, string)
            self.text_widget.see(tk.END)
            self.text_widget.config(state=tk.DISABLED)
        except tk.TclError:
            pass

    def flush(self):
        pass


# ======================== GUI ========================
def load_existing_domain_rows(file_path):
    """读取已存在的 {domain}.xlsx。

    返回值：
      - 文件不存在：{}，允许直接新建；
      - 读取成功：{link: title}；
      - 文件存在但读取失败：None，调用方必须保护原文件，禁止覆盖。
    """
    existing = {}
    if not os.path.exists(file_path):
        return existing
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            title = str(row[0] or "").strip()
            link = str(row[1] or "").strip()
            if link and link not in existing:
                existing[link] = title
        wb.close()
    except Exception as e:
        print(f"读取已有文件失败，已保护原文件不覆盖: {e}")
        return None
    return existing


def build_recovery_file_path(file_path):
    """为无法安全合并的结果生成不覆盖历史文件的恢复文件路径。"""
    base, ext = os.path.splitext(file_path)
    stamp = time.strftime("%Y%m%d-%H%M%S")
    candidate = f"{base}.recovery-{stamp}{ext or '.xlsx'}"
    counter = 2
    while os.path.exists(candidate):
        candidate = f"{base}.recovery-{stamp}-{counter}{ext or '.xlsx'}"
        counter += 1
    return candidate


LINK_INDEX_FILENAME = "links_index.json"


def update_link_index(output_path, pages):
    """把分类列表的 域名→分类标题→分类URL 映射记录进 {output_path}/links_index.json。

    结构: {域名: {分类标题: [分类URL, ...]}}
    已存在的文件会被读入并合并，同一分类下的 URL 按出现顺序追加去重，实现可追加。
    分类标题用列表而非单值：同域名下允许重名分类（如几个都叫 Classic 的列表页），
    URL 全部保留，不会因标题相同而丢失。
    """
    json_path = os.path.join(output_path, LINK_INDEX_FILENAME)
    index = {}
    if os.path.exists(json_path):
        try:
            with open(json_path, "r", encoding="utf-8") as f:
                loaded = json.load(f)
            index = _normalize_link_index(loaded)
        except (ValueError, TypeError, OSError):
            print(f"读取已有链接记录失败，将重建: {json_path}")

    added = 0
    title_counts = {}
    for title, url in pages:
        url = str(url or "").strip()
        if not url:
            continue
        domain = urlparse(url).netloc
        title = str(title or "").strip()
        title_counts[(domain, title)] = title_counts.get((domain, title), 0) + 1

        urls = index.setdefault(domain, {}).setdefault(title, [])
        if url not in urls:
            urls.append(url)
            added += 1

    os.makedirs(output_path, exist_ok=True)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=2)
    print(f"已更新链接记录: {json_path}（新增 {added} 个分类）")

    # 重名分类的 xlsx title 列无法区分来源列表页，索引里已按顺序保留全部 URL，
    # 这里额外提示一次，方便操作员判断是否需要改成不同的分类名。
    repeated = [
        f"{domain} / {title} × {count}"
        for (domain, title), count in title_counts.items()
        if count > 1
    ]
    if repeated:
        print("提示: 以下分类标题在同一域名下重复，索引已保留全部 URL: " + "；".join(repeated))
    return json_path


def _normalize_link_index(loaded):
    """兼容旧版 {域名: {分类标题: 分类URL}}：标量一律升级为单元素列表。"""
    index = {}
    if not isinstance(loaded, dict):
        return index
    for domain, titles in loaded.items():
        if not isinstance(titles, dict):
            continue
        bucket = {}
        for title, value in titles.items():
            if isinstance(value, dict) or value is None:
                continue
            urls = []
            for raw in (value if isinstance(value, list) else [value]):
                raw = str(raw or "").strip()
                if raw:
                    urls.append(raw)
            if urls:
                bucket[str(title)] = urls
        if bucket:
            index[str(domain)] = bucket
    return index


class LinkCollectorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("详情页链接采集")
        self.root.geometry("960x780")
        self.root.minsize(860, 680)
        self.root.configure(bg=UI_COLORS["background"])

        self.is_running = False
        self.stop_event = threading.Event()
        self.current_tab = None

        self.max_pages_var = tk.StringVar(value="")  # 留空=不限制
        self.xpath_var = tk.StringVar()
        self.output_var = tk.StringVar()

        self._build_ui()
        sys.stdout = TextRedirector(self.log_text)
        sys.stderr = sys.stdout

    def _build_ui(self):
        main = tk.Frame(self.root, bg=UI_COLORS["background"])
        main.pack(fill=tk.BOTH, expand=True, padx=20, pady=16)

        # ── 标题 ──
        header = tk.Frame(main, bg=UI_COLORS["background"])
        header.pack(fill=tk.X, pady=(0, 12))
        tk.Label(
            header, text="详情页链接采集",
            bg=UI_COLORS["background"], fg=UI_COLORS["text"],
            font=(UI_FONT, 18, "bold"),
        ).pack(side=tk.LEFT)
        tk.Label(
            header, text="从 Shopify 列表页自动提取产品链接",
            bg=UI_COLORS["background"], fg=UI_COLORS["text_secondary"],
            font=(UI_FONT, 9),
        ).pack(side=tk.LEFT, padx=(16, 0), pady=(8, 0))

        # ── 控制区 ──
        ctrl_card = tk.Frame(main, bg=UI_COLORS["card"],
                             highlightbackground=UI_COLORS["border"],
                             highlightthickness=1, bd=0)
        ctrl_card.pack(fill=tk.X, pady=(0, 10))
        ctrl_inner = tk.Frame(ctrl_card, bg=UI_COLORS["card"])
        ctrl_inner.pack(fill=tk.X, padx=16, pady=12)

        _entry_style = dict(
            bg=UI_COLORS["card_alt"], fg=UI_COLORS["text"],
            insertbackground=UI_COLORS["primary"],
            font=(UI_FONT, 10), relief=tk.FLAT, bd=0,
            highlightthickness=1, highlightbackground=UI_COLORS["border"],
            highlightcolor=UI_COLORS["primary"],
        )

        # 第一行: 保存路径 + XPath + 最大页数
        row1 = tk.Frame(ctrl_inner, bg=UI_COLORS["card"])
        row1.pack(fill=tk.X, pady=(0, 8))

        # 保存路径
        col_path = tk.Frame(row1, bg=UI_COLORS["card"])
        col_path.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))
        tk.Label(col_path, text="保存路径",
                 bg=UI_COLORS["card"], fg=UI_COLORS["text_secondary"],
                 font=(UI_FONT, 9)).pack(anchor=tk.W)
        path_row = tk.Frame(col_path, bg=UI_COLORS["card"])
        path_row.pack(fill=tk.X)
        path_entry = tk.Entry(path_row, textvariable=self.output_var, **_entry_style)
        path_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=6)
        browse_btn = tk.Button(
            path_row, text="浏览", command=self._browse_output,
            bg=UI_COLORS["card_alt"], fg=UI_COLORS["text_secondary"],
            activebackground="#EEF2F7", relief=tk.FLAT, bd=0,
            font=(UI_FONT, 9), padx=14, pady=4, cursor="hand2",
        )
        browse_btn.pack(side=tk.LEFT, padx=(6, 0))

        # XPath
        col_xpath = tk.Frame(row1, bg=UI_COLORS["card"])
        col_xpath.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))
        tk.Label(col_xpath, text="XPath",
                 bg=UI_COLORS["card"], fg=UI_COLORS["text_secondary"],
                 font=(UI_FONT, 9)).pack(anchor=tk.W)
        xpath_entry = tk.Entry(col_xpath, textvariable=self.xpath_var,
                               **{**_entry_style, "font": (MONO_FONT, 10)})
        xpath_entry.pack(fill=tk.X, ipady=6)

        # 最大页数
        col_pages = tk.Frame(row1, bg=UI_COLORS["card"])
        col_pages.pack(side=tk.LEFT, fill=tk.X, expand=False, padx=(0, 0))
        tk.Label(col_pages, text="最大页数",
                 bg=UI_COLORS["card"], fg=UI_COLORS["text_secondary"],
                 font=(UI_FONT, 9)).pack(anchor=tk.W)
        pages_entry = tk.Entry(col_pages, textvariable=self.max_pages_var,
                               width=10, **_entry_style)
        pages_entry.pack(fill=tk.X, ipady=6)

        # 第二行: 按钮
        row2 = tk.Frame(ctrl_inner, bg=UI_COLORS["card"])
        row2.pack(fill=tk.X)

        self.stop_btn = tk.Button(
            row2, text="停止", command=self._stop,
            bg=UI_COLORS["card"], fg=UI_COLORS["danger"],
            activebackground=UI_COLORS["danger_light"],
            relief=tk.FLAT, bd=1, font=(UI_FONT, 10, "bold"),
            padx=20, pady=8, state=tk.DISABLED, cursor="hand2",
            highlightthickness=1, highlightbackground="#FECACA",
            highlightcolor="#FECACA",
        )
        self.stop_btn.pack(side=tk.LEFT, padx=(0, 8))

        self.start_btn = tk.Button(
            row2, text="开始采集", command=self._start,
            bg=UI_COLORS["primary"], fg="#FFFFFF",
            activebackground=UI_COLORS["primary_hover"],
            activeforeground="#FFFFFF",
            relief=tk.FLAT, bd=0, font=(UI_FONT, 10, "bold"),
            padx=28, pady=8, cursor="hand2",
        )
        self.start_btn.pack(side=tk.LEFT)

        # ── 分类列表 ──
        pages_card = tk.Frame(main, bg=UI_COLORS["card"],
                              highlightbackground=UI_COLORS["border"],
                              highlightthickness=1, bd=0)
        pages_card.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

        pages_inner = tk.Frame(pages_card, bg=UI_COLORS["card"])
        pages_inner.pack(fill=tk.BOTH, expand=True, padx=16, pady=12)

        tk.Label(
            pages_inner, text="分类列表（每行: 标题, URL）",
            bg=UI_COLORS["card"], fg=UI_COLORS["text"],
            font=(UI_FONT, 11, "bold"),
        ).pack(anchor=tk.W)
        tk.Label(
            pages_inner,
            text='示例: Loungewear > Pajamas, https://example.com/collections/pajamas',
            bg=UI_COLORS["card"], fg=UI_COLORS["text_muted"],
            font=(UI_FONT, 9),
        ).pack(anchor=tk.W, pady=(2, 8))

        self.pages_text = scrolledtext.ScrolledText(
            pages_inner, height=8, wrap=tk.NONE,
            bg=UI_COLORS["card_alt"], fg=UI_COLORS["text"],
            insertbackground=UI_COLORS["primary"],
            font=(MONO_FONT, 9), relief=tk.FLAT, bd=0,
            highlightthickness=1, highlightbackground=UI_COLORS["border"],
            highlightcolor=UI_COLORS["primary"],
            padx=10, pady=8,
        )
        self.pages_text.pack(fill=tk.BOTH, expand=True)

        # ── 日志区 ──
        log_card = tk.Frame(main, bg=UI_COLORS["log_background"],
                            highlightbackground="#1E293B",
                            highlightthickness=1, bd=0)
        log_card.pack(fill=tk.BOTH, expand=True)

        log_header = tk.Frame(log_card, bg=UI_COLORS["log_panel"], height=36)
        log_header.pack(fill=tk.X)
        log_header.pack_propagate(False)
        tk.Label(
            log_header, text="运行日志",
            bg=UI_COLORS["log_panel"], fg="#F8FAFC",
            font=(UI_FONT, 10, "bold"),
        ).pack(side=tk.LEFT, padx=14, pady=8)
        clear_btn = tk.Button(
            log_header, text="清空", command=self._clear_log,
            bg="#1E293B", fg="#CBD5E1", activebackground="#334155",
            activeforeground="#FFFFFF", relief=tk.FLAT, bd=0,
            font=(UI_FONT, 8), padx=10, pady=4, cursor="hand2",
        )
        clear_btn.pack(side=tk.RIGHT, padx=10, pady=6)

        self.log_text = scrolledtext.ScrolledText(
            log_card, state=tk.DISABLED,
            bg=UI_COLORS["log_background"], fg=UI_COLORS["log_text"],
            insertbackground="#FFFFFF", selectbackground="#1D4ED8",
            selectforeground="#FFFFFF",
            font=(MONO_FONT, 9), relief=tk.FLAT, bd=0,
            padx=12, pady=10, wrap=tk.WORD,
        )
        self.log_text.pack(fill=tk.BOTH, expand=True)

    def _browse_output(self):
        path = filedialog.askdirectory(title="选择保存目录")
        if path:
            self.output_var.set(path)

    def _clear_log(self):
        self.log_text.config(state=tk.NORMAL)
        self.log_text.delete("1.0", tk.END)
        self.log_text.config(state=tk.DISABLED)

    def _parse_pages(self):
        """解析分类列表文本框，返回 [(title, url), ...]"""
        raw = self.pages_text.get("1.0", tk.END).strip()
        if not raw:
            return []
        pages = []
        for line in raw.splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue

            # 提取 URL（从 http 开始到下一个引号/括号/逗号截止）
            url_match = re.search(r'(https?://[^\s,"\')\]]+)', line)
            if not url_match:
                continue
            url = url_match.group(1)

            # 提取标题：URL 前面的文本，去掉括号、引号、逗号
            title = ""
            url_start = line.find(url)
            if url_start > 0:
                title_part = line[:url_start].strip(' \t,("\'')
                title = re.sub(r'[",)\]]+$', '', title_part).strip()

            if not title:
                title = urlparse(url).path.split("/")[-1]

            pages.append((title, url))
        return pages

    def _start(self):
        pages = self._parse_pages()
        if not pages:
            messagebox.showerror("错误", "请先填写分类列表")
            return

        output_path = self.output_var.get().strip()
        if not output_path:
            output_path = os.path.join(
                os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                "output",
            )
            self.output_var.set(output_path)

        max_pages_raw = self.max_pages_var.get().strip()
        if max_pages_raw:
            try:
                max_pages = int(max_pages_raw)
                if max_pages <= 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("错误", "最大页数必须为正整数，或留空表示不限制")
                return
        else:
            max_pages = 0

        xpath = self.xpath_var.get().strip()

        self.is_running = True
        self.stop_event.clear()
        self.start_btn.config(state=tk.DISABLED)
        self.stop_btn.config(state=tk.NORMAL)
        self._clear_log()

        thread = threading.Thread(
            target=self._run_task,
            args=(pages, output_path, max_pages, xpath),
            daemon=True,
        )
        thread.start()

    def _interrupt_current_page(self):
        """尽力中止正在加载的页面；不同 DrissionPage 版本没有该接口时安全忽略。"""
        tab = self.current_tab
        if tab is None:
            return
        try:
            stop_loading = getattr(tab, "stop_loading", None)
            if callable(stop_loading):
                stop_loading()
        except Exception:
            pass

    def _stop(self):
        if self.is_running:
            self.stop_event.set()
            self.stop_btn.config(state=tk.DISABLED)
            print("\n[停止] 已放弃当前页，正在保存此前已采集的链接...")
            threading.Thread(target=self._interrupt_current_page, daemon=True).start()

    def _run_task(self, pages, output_path, max_pages, xpath=""):
        all_results = []
        seen_all_links = set()
        chrome = None

        try:
            print(f"共 {len(pages)} 个分类")
            print(f"最大页数: {max_pages if max_pages else '不限制'}")
            if xpath:
                print(f"XPath: {xpath}")
            print(f"保存目录: {output_path}")
            print("采集模式: 纯分页")
            print()

            chrome = create_browser()
            tab = chrome.latest_tab
            self.current_tab = tab

            for i, (title, base_url) in enumerate(pages, 1):
                if self.stop_event.is_set():
                    break

                # 分类统计提前初始化，确保即使分类中途异常也能准确报告已保留结果。
                category_total = 0
                category_new_for_all = 0
                category_skipped = 0

                try:
                    print(f"[{i}/{len(pages)}] {title}")
                    seen_in_category = set()
                    page_num = 1

                    while not self.stop_event.is_set():
                        if max_pages and page_num > max_pages:
                            print(f"  已达到最大页数 {max_pages}")
                            break

                        url = build_page_url(base_url, page_num)
                        print(f"  第 {page_num} 页: {url}")

                        try:
                            tab.get(url)
                            if self.stop_event.is_set():
                                print("  当前页已放弃")
                                break
                            links = collect_page_links(
                                tab,
                                stop_event=self.stop_event,
                                xpath=xpath,
                            )
                        except ContextLostError:
                            if self.stop_event.is_set():
                                print("  当前页已放弃")
                                break
                            print("  页面上下文丢失，快速重试一次...")
                            if self.stop_event.wait(0.8):
                                break
                            try:
                                tab.get(url)
                                if self.stop_event.is_set():
                                    print("  当前页已放弃")
                                    break
                                links = collect_page_links(
                                    tab,
                                    stop_event=self.stop_event,
                                    xpath=xpath,
                                )
                            except ContextLostError:
                                print("  页面未恢复，结束该分类")
                                break

                        # 停止时明确不接纳当前页刚解析出来的结果。
                        if self.stop_event.is_set():
                            print("  当前页已放弃")
                            break

                        new_links = [link for link in links if link not in seen_in_category]

                        if not new_links:
                            if page_num == 1:
                                retried = False
                                for retry in range(1, MAX_RETRIES + 1):
                                    if self.stop_event.is_set():
                                        break
                                    print(f"  首页无链接，第 {retry}/{MAX_RETRIES} 次重试...")
                                    if self.stop_event.wait(RETRY_DELAY):
                                        break

                                    tab.get(url)
                                    if self.stop_event.is_set():
                                        break
                                    links = collect_page_links(
                                        tab,
                                        stop_event=self.stop_event,
                                        xpath=xpath,
                                    )
                                    if self.stop_event.is_set():
                                        break

                                    new_links = [
                                        link for link in links
                                        if link not in seen_in_category
                                    ]
                                    if new_links:
                                        print(f"  重试成功: {len(new_links)} 条")
                                        retried = True
                                        break

                                if self.stop_event.is_set():
                                    print("  当前页已放弃")
                                    break
                                if not retried:
                                    print("  跳过该分类")
                                    break
                            else:
                                print("  无新链接，该分类采集完毕")
                                break

                        # 当前页只有在完整采集结束且未收到停止信号时才提交。
                        if self.stop_event.is_set():
                            print("  当前页已放弃")
                            break

                        page_new_for_all = 0
                        page_skipped = 0
                        for link in new_links:
                            seen_in_category.add(link)
                            category_total += 1

                            # 每完成一页就立即提交到最终结果。
                            # 后续分页或当前分类发生异常时，之前完成页不会丢失。
                            if link in seen_all_links:
                                category_skipped += 1
                                page_skipped += 1
                                continue
                            seen_all_links.add(link)
                            all_results.append({"title": title, "link": link})
                            category_new_for_all += 1
                            page_new_for_all += 1

                        print(
                            f"  新增 {len(new_links)} 条（分类累计 {category_total} 条，"
                            f"全局新增 {page_new_for_all} 条，跨分类跳过 {page_skipped} 条）"
                        )
                        page_num += 1

                    print(
                        f"  共采集 {category_total} 条（新增 {category_new_for_all} 条，"
                        f"跨分类去重跳过 {category_skipped} 条）"
                    )
                    print()

                except Exception as e:
                    print(f"  分类 [{title}] 出错: {e}")
                    if category_total:
                        print(f"  已保留该分类此前完成分页的 {category_total} 条链接")
                    print()
                    if self.stop_event.is_set():
                        break
                    continue

        except Exception as e:
            print(f"发生异常: {e}")
            import traceback
            traceback.print_exc()

        finally:
            self.current_tab = None
            if chrome is not None:
                try:
                    chrome.quit()
                except Exception as e:
                    print(f"关闭浏览器时出现异常（不影响保存）: {e}")

            # 无论正常完成、用户停止还是运行异常，都保存此前已完成分页的结果。
            try:
                print(f"采集结束: 共 {len(all_results)} 条链接")
                os.makedirs(output_path, exist_ok=True)
                update_link_index(output_path, pages)

                if all_results:
                    domain_groups = {}
                    for row in all_results:
                        domain = urlparse(row["link"]).netloc
                        domain_groups.setdefault(domain, []).append(row)

                    for domain, rows in domain_groups.items():
                        file_name = f"{domain}.xlsx"
                        file_path = os.path.join(output_path, file_name)

                        existing = load_existing_domain_rows(file_path)

                        # 原 XLSX 存在但读取失败时，绝不覆盖历史文件。
                        # 本轮结果另存为 recovery 文件，方便后续人工检查/合并。
                        if existing is None:
                            recovery_path = build_recovery_file_path(file_path)
                            recovery_rows = {}
                            for row in rows:
                                link = row["link"]
                                if link not in recovery_rows:
                                    recovery_rows[link] = row.get("title", "")

                            wb = openpyxl.Workbook()
                            ws = wb.active
                            ws.append(["title", "link"])
                            for link, title in recovery_rows.items():
                                ws.append([title, link])
                            wb.save(recovery_path)
                            print(
                                f"原文件读取失败，已禁止覆盖: {file_path}"
                                f"；本轮 {len(recovery_rows)} 条另存为: {recovery_path}"
                            )
                            continue

                        merged = dict(existing)
                        added = 0
                        for row in rows:
                            link = row["link"]
                            if link not in merged:
                                merged[link] = row.get("title", "")
                                added += 1

                        wb = openpyxl.Workbook()
                        ws = wb.active
                        ws.append(["title", "link"])
                        for link, title in merged.items():
                            ws.append([title, link])
                        wb.save(file_path)
                        print(
                            f"已保存: {file_path}（已有 {len(existing)} 条 + 新增 {added} 条"
                            f" = 合并 {len(merged)} 条）"
                        )

                    print(f"共 {len(domain_groups)} 个域名文件")
                else:
                    print("本次没有新增可保存的链接")
            except Exception as e:
                print(f"保存结果失败: {e}")
                import traceback
                traceback.print_exc()

            self.is_running = False
            try:
                self.root.after(0, self._reset_buttons)
            except tk.TclError:
                pass

    def _reset_buttons(self):
        self.start_btn.config(state=tk.NORMAL)
        self.stop_btn.config(state=tk.DISABLED)


if __name__ == "__main__":
    root = tk.Tk()
    app = LinkCollectorApp(root)
    root.mainloop()

