"""
Shopify 一体化采集工作台（独立版）
已内置 detail_link_GUI.py 与 Shopify_JSON_GUI.py 核心代码，不再依赖外部模块。
包含链接采集、商品解析与转换三个独立阶段。
"""

# ======================== Shopify 链接采集核心 ========================
#
# 用途: 从 Shopify 列表页提取产品链接
# 原理: DrissionPage 浏览器加载页面，从 <script> JSON 中提取 handle
#
# 界面:
#   - 分类列表文本框（每行: 标题, URL）
#   - 最大页数输入框（留空=不限制）
#   - 保存路径选择
#   - 开始/停止按钮 + 实时日志
#
# 输出: XLSX（title, link）
# ========================================================================


# ======================== 样式配置 ========================
import csv
import hashlib
import json
import os
import random
import re
import string
import sys
import tempfile
import threading
import time
import traceback
import urllib.request
from enum import Enum
from html import escape
from html.parser import HTMLParser
from itertools import product as itertools_product
from pathlib import Path, PurePath
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse

import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk

import openpyxl
import pandas as pd
from DrissionPage import Chromium, ChromiumOptions
from DrissionPage.errors import ContextLostError

UI_COLORS = {
    "background": "#F4F7FB",
    "card": "#FFFFFF",
    "card_alt": "#F8FAFC",
    "border": "#E2E8F0",
    "border_focus": "#93C5FD",
    "text": "#0F172A",
    "text_secondary": "#475569",
    "text_muted": "#94A3B8",
    "primary": "#2563EB",
    "primary_hover": "#1D4ED8",
    "primary_light": "#EFF6FF",
    "success": "#16A34A",
    "success_light": "#F0FDF4",
    "warning": "#D97706",
    "warning_light": "#FFFBEB",
    "danger": "#DC2626",
    "danger_light": "#FEF2F2",
    "danger_border": "#FECACA",
    "log_background": "#0F172A",
    "log_panel": "#111827",
    "log_text": "#D1D5DB",
}

UI_FONT = "Microsoft YaHei UI"
MONO_FONT = "Consolas"

LINK_MAX_RETRIES = 5
RETRY_DELAY = 5
PAGE_LOAD_TIMEOUT = 30


def extract_product_links(html_text, base_url):
    """普通模式只从 JSON handle 构造商品链接，不回退到 DOM 链接。"""
    parsed = urlparse(base_url)
    base = f"{parsed.scheme}://{parsed.netloc}"
    handles = dict.fromkeys(
        handle.strip()
        for handle in re.findall(r'"handle"\s*:\s*"([^"\\]+)"', html_text)
        if handle.strip()
    )
    return [f"{base}/products/{handle}" for handle in handles]


def _run_js_with_retry(tab, script, default=None, retries=2):
    """执行 JavaScript，并处理页面上下文临时丢失。"""
    for attempt in range(retries + 1):
        try:
            result = tab.run_js(script)
            return default if result is None else result
        except ContextLostError:
            if attempt >= retries:
                return default
            time.sleep(2)
        except Exception:
            if attempt >= retries:
                return default
            time.sleep(1)
    return default


def extract_links_by_xpath(tab, xpath):
    """用 XPath 从页面提取产品链接，返回规范化后的链接列表。"""
    result = _run_js_with_retry(
        tab,
        f"""
        return (() => {{
            const xpath = {json.dumps(xpath)};
            const results = [];
            try {{
                const snapshot = document.evaluate(
                    xpath, document, null,
                    XPathResult.ORDERED_NODE_SNAPSHOT_TYPE, null
                );
                for (let i = 0; i < snapshot.snapshotLength; i++) {{
                    const node = snapshot.snapshotItem(i);
                    const val = node.nodeType === 2 ? node.value : (node.href || node.textContent || '');
                    if (val) results.push(val.trim());
                }}
            }} catch(e) {{
                console.error('XPath error:', e);
            }}
            return results;
        }})();
        """,
        default=[],
    )
    if not isinstance(result, list):
        return []

    base_parsed = urlparse(tab.url)
    base = f"{base_parsed.scheme}://{base_parsed.netloc}"
    seen = set()
    links = []
    for raw in result:
        raw = str(raw).strip()
        if not raw:
            continue
        # 相对路径或绝对路径 → 完整 URL
        if raw.startswith("/"):
            full = f"{base}{raw}"
        elif raw.startswith("http"):
            full = raw
        else:
            continue
        parsed = urlparse(full)
        match = re.match(r"^/products/([^/?#]+)", parsed.path, re.IGNORECASE)
        if match:
            norm = f"{base}/products/{match.group(1)}"
        else:
            norm = f"{parsed.scheme}://{parsed.netloc}{parsed.path.rstrip('/')}"
        if norm not in seen:
            seen.add(norm)
            links.append(norm)
    return links


def collect_page_links(tab, stop_event=None, xpath=""):
    """文档加载完成后直接提取；指定 XPath 时只使用 XPath。"""
    if stop_event is not None and stop_event.is_set():
        return []
    if not tab.wait.doc_loaded(timeout=PAGE_LOAD_TIMEOUT, raise_err=False):
        raise TimeoutError(f"页面加载超时: {tab.url}")
    if stop_event is not None and stop_event.is_set():
        return []
    if xpath:
        return extract_links_by_xpath(tab, xpath)
    html_text = tab.run_js("return document.documentElement.outerHTML;") or ""
    return extract_product_links(html_text, tab.url)


# ======================== 文本重定向 ========================


# ======================== GUI ========================


# ════════════════════════════════════════════════════════════
# 全局采集配置
# ════════════════════════════════════════════════════════════
DELAY = 0.5
PRODUCT_MAX_RETRIES = 5
RETRY_BASE_DELAY = 30
SAVE_EVERY = 20
SKIP_OPTIONS = None

CSV_FIELDS = [
    "title", "name", "price1", "price2",
    "styles1", "styles2", "styles3",
    "src_links", "link-href", "details",
]


# ════════════════════════════════════════════════════════════
# 核心业务逻辑 (保持原有逻辑不变)
# ════════════════════════════════════════════════════════════
def to_shopify_json_url(url):
    parsed = urlparse(url)
    path = parsed.path.rstrip("/")
    if not path.endswith(".json"):
        path = f"{path}.json"
    return urlunparse((parsed.scheme, parsed.netloc, path, "", "", ""))


def group_targets_by_json_url(targets):
    json_url_groups = {}
    duplicate_product_count = 0
    for title, url in targets:
        original_url = str(url).strip()
        if not original_url:
            continue
        json_url = to_shopify_json_url(original_url)
        if json_url in json_url_groups:
            duplicate_product_count += 1
            continue
        title_text = str(title).strip() if title is not None else ""
        json_url_groups[json_url] = {
            "title": title_text,
            "original_url": urlunparse(urlparse(original_url)._replace(query="", fragment="")),
        }
    return json_url_groups, duplicate_product_count


_exchange_rates_cache = {}


def fetch_exchange_rates():
    global _exchange_rates_cache
    if _exchange_rates_cache:
        return _exchange_rates_cache
    try:
        url = "https://open.er-api.com/v6/latest/USD"
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=10) as response:
            data = json.loads(response.read().decode())
        _exchange_rates_cache = data.get("rates", {})
        print(f"[汇率] 成功获取 {_exchange_rates_cache.get('USD', 1)} 相关汇率")
    except Exception as e:
        print(f"[汇率] 获取失败: {e}")
    return _exchange_rates_cache


def convert_price_to_usd(price_str, currency, rates=None):
    if not price_str or not currency:
        return price_str
    if rates is None:
        rates = _exchange_rates_cache
    if not rates:
        return price_str
    currency = str(currency).upper().strip()
    if currency == "USD":
        return price_str
    rate = rates.get(currency)
    if not rate or rate == 0:
        return price_str
    try:
        amount = float(str(price_str).strip())
    except (ValueError, TypeError):
        return price_str
    usd_amount = amount / rate
    return f"{usd_amount:.2f}"


def find_fallback_variant(product):
    variants = product.get("variants", [])
    for variant in variants:
        if variant.get("position") == 1:
            return variant
    return variants[0] if variants else {}


def format_combo_value(value):
    return "" if value is None else str(value)


def get_sorted_options(product, skip_options=None):
    options = sorted(
        product.get("options", []),
        key=lambda option: option.get("position") or 0,
    )
    if skip_options:
        skip_lower = [s.lower() for s in skip_options]
        options = [
            opt for opt in options
            if not any(s in (opt.get("name") or "").lower() for s in skip_lower)
        ]
    multi_value = [opt for opt in options if len(opt.get("values") or []) > 1]
    color_idx = next(
        (i for i, opt in enumerate(multi_value)
         if (opt.get("name") or "").lower() == "color"),
        None,
    )
    if color_idx is not None and color_idx > 0:
        color_opt = multi_value.pop(color_idx)
        multi_value.insert(0, color_opt)
    return multi_value


def build_option_combos(options):
    value_groups = [
        [format_combo_value(value) for value in option.get("values", [])]
        for option in options
    ]
    if not value_groups:
        return [()]
    return list(itertools_product(*value_groups))


def _option_field(options, combo_index):
    if combo_index < len(options):
        return f"option{options[combo_index].get('position', combo_index + 1)}"
    return f"option{combo_index + 1}"


def format_price(value, rates=None, currency=None):
    if value in (None, ""):
        return ""
    raw = str(value).strip()
    if not raw:
        return ""
    if rates and currency:
        raw = convert_price_to_usd(raw, currency, rates)
    # 当前采集的是 /products/{handle}.json，价格单位为主货币单位。
    # 整数与小数字符串使用相同单位，不能依据小数点猜测是否为“分”。
    try:
        amount = float(raw)
    except ValueError:
        return raw
    return str(int(amount)) if amount.is_integer() else f"{amount:.2f}".rstrip("0").rstrip(".")


def format_price2(variant, rates=None):
    currency = variant.get("price_currency", "")
    compare_at_price = variant.get("compare_at_price")
    if compare_at_price and str(compare_at_price).strip():
        return format_price(compare_at_price, rates=rates, currency=currency)
    return format_price(variant.get("price", ""), rates=rates, currency=currency)


def format_image_url(url, size="600x600"):
    if not url:
        return ""
    parsed = urlparse(url)
    path = parsed.path
    dot_index = path.rfind(".")
    slash_index = path.rfind("/")
    if dot_index > slash_index:
        path = f"{path[:dot_index]}_{size}{path[dot_index:]}"
    return urlunparse((
        parsed.scheme, parsed.netloc, path,
        parsed.params, parsed.query, parsed.fragment,
    ))


def build_images_by_id(product, skip_positions=None):
    images = product.get("images", [])
    if skip_positions:
        images = [img for img in images if (img.get("position") or 0) not in skip_positions]
    return {
        image.get("id"): format_image_url(image.get("src", ""))
        for image in images
    }


def build_first_image_srcs(product, limit=None, skip_positions=None):
    images = sorted(
        product.get("images", []),
        key=lambda image: image.get("position") or 0,
    )
    if skip_positions:
        images = [img for img in images if (img.get("position") or 0) not in skip_positions]
    selected_images = images if limit is None else images[:limit]
    srcs = [
        format_image_url(image.get("src", ""))
        for image in selected_images
        if image.get("src")
    ]
    return "#".join(srcs)


def build_images_by_first_option(product, options, skip_positions=None):
    images_by_id = build_images_by_id(product, skip_positions=skip_positions)
    field = _option_field(options, 0)
    result = {}
    for variant in product.get("variants", []):
        opt_value = format_combo_value(variant.get(field))
        if opt_value and opt_value not in result:
            image_id = variant.get("image_id")
            if image_id and image_id in images_by_id:
                result[opt_value] = images_by_id[image_id]
    return result


def build_segment(combo, options, variant, images_by_first_option, rates=None):
    segment_parts = []
    if combo:
        segment_parts.append(format_combo_value(combo[0]))
        for index, value in enumerate(combo[1:], start=1):
            option_name = format_combo_value(options[index].get("name", ""))
            segment_parts.extend([option_name, format_combo_value(value)])
    currency = variant.get("price_currency", "")
    price1 = format_price(variant.get("price"), rates=rates, currency=currency)
    price2 = format_price(variant.get("compare_at_price") or variant.get("price"), rates=rates, currency=currency)
    first_opt_value = format_combo_value(combo[0]) if combo else ""
    image_src = images_by_first_option.get(first_opt_value, "")
    segment = f"{'&'.join(segment_parts)}${price1}${price2}"
    if image_src:
        segment = f"{segment}@{image_src}"
    return segment


def build_variant_combo(product, skip_positions=None, skip_options=None, rates=None):
    options = get_sorted_options(product, skip_options=skip_options)
    combos = build_option_combos(options)
    fallback_variant = find_fallback_variant(product)
    images_by_first_opt = build_images_by_first_option(product, options, skip_positions=skip_positions)
    fields = [_option_field(options, i) for i in range(len(options))]
    exact_variants = {}
    first_variants = {}
    for variant in product.get("variants", []):
        key = tuple(format_combo_value(variant.get(field)) for field in fields)
        exact_variants.setdefault(key, variant)
        if key:
            first_variants.setdefault(key[0], variant)
    segments = []
    for combo in combos:
        variant = (exact_variants.get(combo)
                   or (first_variants.get(combo[0]) if combo else None)
                   or fallback_variant)
        segments.append(build_segment(combo, options, variant, images_by_first_opt, rates=rates))
    first_option_name = format_combo_value(options[0].get("name", "")) if options else ""
    if first_option_name:
        return f"{first_option_name}#{'#'.join(segments)}"
    return "#".join(segments)


class CleanBodyHtmlParser(HTMLParser):
    void_tags = {
        "area", "base", "br", "col", "embed", "hr", "img", "input",
        "link", "meta", "param", "source", "track", "wbr",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            return
        if tag == "img":
            src = dict(attrs).get("src", "")
            self.parts.append(f'<img src="{escape(src, quote=True)}">' if src else "<img>")
        else:
            self.parts.append(f"<{tag}>")

    def handle_endtag(self, tag):
        if tag == "a" or tag in self.void_tags:
            return
        self.parts.append(f"</{tag}>")

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)

    def handle_data(self, data):
        data = "".join(ch for ch in data if ord(ch) <= 0xFFFF)
        self.parts.append(escape(data, quote=False))

    def get_html(self):
        return "".join(self.parts).strip()


def clean_body_html(body_html):
    if not body_html:
        return ""
    parser = CleanBodyHtmlParser()
    parser.feed(str(body_html))
    parser.close()
    return parser.get_html()


def create_browser():
    co = ChromiumOptions()
    co.incognito(on_off=True)
    co.no_imgs(True)
    co.set_load_mode("normal")
    # 固定使用本地代理 127.0.0.1:7897
    co.set_argument("--proxy-server=http://127.0.0.1:7897")
    random_port = random.randint(9222, 9322)
    co.set_local_port(random_port)
    for attempt in range(3):
        try:
            chrome = Chromium(co)
            print(f"浏览器启动成功，端口: {random_port}")
            return chrome
        except Exception as e:
            print(f"端口 {random_port} 启动失败 (尝试 {attempt + 1}/3): {e}")
            random_port = random.randint(9222, 9322)
            co.set_local_port(random_port)
    raise RuntimeError("无法启动浏览器，已重试多次")


def read_targets_from_excel(filepath):
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    targets = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        title = str(row[0]).strip() if row and row[0] else ""
        url = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if url:
            targets.append((title, url))
    wb.close()
    return targets


def fetch_json(tab, json_url):
    tab.listen.start(json_url)
    tab.get(json_url)
    packet = tab.listen.wait(timeout=30)
    status = packet.response.status if packet else 0

    # Shopify 商品不存在时，404 不进行重试
    if status == 404:
        print("  商品不存在 (status=404)，跳过")
        return status, None

    if status == 200:
        text = tab.run_js("return document.body.innerText")
        if not text:
            text = tab.run_js("return document.body.textContent")
        try:
            return status, json.loads(text)
        except json.JSONDecodeError:
            pre_text = tab.run_js("return document.querySelector('pre')?.innerText || ''")
            if pre_text:
                try:
                    return status, json.loads(pre_text)
                except json.JSONDecodeError:
                    pass
            return status, None
    else:
        return status, None


class InvalidProductError(ValueError):
    """响应不是可解析的 Shopify 商品。"""


def validate_product_json(data):
    """验证 Shopify 商品 JSON 的基本结构；空响应和错误页不能标为成功。"""
    if not isinstance(data, dict) or not isinstance(data.get("product"), dict):
        raise InvalidProductError("JSON 缺少 product 对象")
    product = data["product"]
    if not isinstance(product.get("id"), int) or isinstance(product["id"], bool) or product["id"] <= 0:
        raise InvalidProductError("商品缺少有效 id")
    for field in ("title", "handle"):
        if not isinstance(product.get(field), str) or not product[field].strip():
            raise InvalidProductError(f"商品缺少有效 {field}")
    variants = product.get("variants")
    if not isinstance(variants, list) or not variants:
        raise InvalidProductError("商品缺少非空 variants 列表")
    for variant in variants:
        if not isinstance(variant, dict) or not isinstance(variant.get("id"), int) or isinstance(variant["id"], bool) or variant["id"] <= 0:
            raise InvalidProductError("商品变体缺少有效 id")
        try:
            price = float(variant.get("price"))
        except (ValueError, TypeError):
            raise InvalidProductError("商品变体缺少有效 price") from None
        if not 0 <= price < float("inf") or isinstance(variant.get("price"), bool):
            raise InvalidProductError("商品变体 price 无效")
    for field in ("options", "images"):
        entries = product.get(field, [])
        if not isinstance(entries, list) or any(not isinstance(entry, dict) for entry in entries):
            raise InvalidProductError(f"商品 {field} 必须是对象列表")
    return product


def crawl_config_key(raw_skip_spec, skip_options=None):
    """记录本轮有效图片/规格配置；只在相同配置下跳过已完成商品。"""
    positions = sorted({
        int(token.strip()) for token in str(raw_skip_spec or "").split(",")
        if token.strip().isdigit() or token.strip() == "-1"
    })
    return json.dumps({"version": 1, "skip_positions": positions,
                       "skip_options": skip_options or []}, sort_keys=True)


def parse_product(data, table_title, original_url, *, skip_positions=None,
                  skip_options=None, rates=None, raw_skip_spec=""):
    validate_product_json(data)
    prepared_data, skip_positions = prepare_data_for_image_skips(
        data, skip_positions, raw_skip_spec,
    )
    product = prepared_data["product"]
    variant = find_fallback_variant(product)
    return {
        "title": table_title,
        "name": product.get("title", ""),
        "price1": format_price(
            variant.get("price", ""),
            rates=rates,
            currency=variant.get("price_currency", ""),
        ),
        "price2": format_price2(variant, rates=rates),
        "styles1": build_variant_combo(
            product,
            skip_positions=skip_positions,
            skip_options=skip_options,
            rates=rates,
        ),
        "styles2": "",
        "styles3": "",
        "src_links": build_first_image_srcs(
            product, limit=5, skip_positions=skip_positions
        ),
        "link-href": original_url,
        "details": clean_body_html(product.get("body_html", "")),
    }


def save_workbook_atomic(workbook, output_path):
    """完整写入同目录临时文件后替换，避免写入失败破坏已有工作簿。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_path = tempfile.mkstemp(suffix=".xlsx", dir=output_path.parent)
    os.close(fd)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        workbook.close()
        if os.path.exists(temporary_path):
            os.unlink(temporary_path)


def save_xlsx(rows, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(CSV_FIELDS)
    for row in rows:
        ws.append([row.get(f, "") for f in CSV_FIELDS])
    configs = wb.create_sheet("_crawl_config")
    configs.sheet_state = "hidden"
    configs.append(["link-href", "config"])
    for row in rows:
        if row.get("link-href") and row.get("_crawl_config"):
            configs.append([row["link-href"], row["_crawl_config"]])
    save_workbook_atomic(wb, output_path)


def load_xlsx(filepath):
    wb = openpyxl.load_workbook(filepath, read_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h) if h else "" for h in next(rows_iter, ())]
        results = [{h: (v if v is not None else "") for h, v in zip(headers, row)}
                   for row in rows_iter]
        configs = {}
        if "_crawl_config" in wb.sheetnames:
            configs = dict(wb["_crawl_config"].iter_rows(min_row=2, max_col=2, values_only=True))
        for row in results:
            config = configs.get(row.get("link-href"))
            if config:
                row["_crawl_config"] = config
        return headers, results
    finally:
        wb.close()


# ════════════════════════════════════════════════════════════
# Shopify 转换 + 价格匹配（从 switch_gui.py 搬入）
# ════════════════════════════════════════════════════════════

PRICE_LIBRARY = [
    8.99, 9.01, 9.31, 9.33, 9.37, 9.63, 9.71, 9.72, 9.84, 9.86, 9.95, 9.99,
    10.03, 10.05, 10.11, 10.13, 10.15, 10.35, 10.36, 10.38, 10.54, 10.63, 10.64, 10.66,
    10.71, 10.78, 10.83, 10.94, 10.95, 10.99, 11.22, 11.36, 11.48, 11.65, 11.87, 11.95,
    11.99, 12.22, 12.25, 12.36, 12.41, 12.87, 12.92, 12.95, 12.99, 13.02, 13.13, 13.14,
    13.47, 13.69, 13.95, 13.99, 14.25, 14.32, 14.62, 14.83, 14.91, 14.95, 14.99, 15.31,
    15.52, 15.78, 15.95, 15.99, 16.34, 16.53, 16.66, 16.77, 16.95, 16.99, 17.11, 17.34,
    17.57, 17.95, 17.99, 18.62, 18.81, 18.84, 18.95, 18.99, 19.04, 19.25, 19.95, 19.99,
    20.14, 20.19, 20.34, 20.68, 20.74, 20.95, 20.99, 21.3, 21.47, 21.74, 21.95, 21.99,
    22.31, 22.34, 22.81, 22.95, 22.99, 23.14, 23.32, 23.95, 23.99, 24.22, 24.46, 24.63,
    24.95, 24.99, 25.37, 25.95, 25.99, 26.18, 26.34, 26.95, 26.99, 27.31, 27.54, 27.95,
    27.99, 28.11, 28.95, 28.99, 29, 29.31, 29.95, 29.99, 30.54, 30.95, 30.99, 31.26,
    31.52, 31.95, 31.99, 32.14, 32.24, 32.95, 32.99, 33.62, 33.84, 33.95, 33.99, 34.18,
    34.62, 34.95, 34.99, 35, 35.14, 35.33, 35.95, 35.99, 36.47, 36.95, 36.99, 37.21,
    37.95, 37.99, 38.14, 38.52, 38.95, 38.99, 39, 39.54, 39.95, 39.99, 40.02, 40.15,
    40.95, 40.99, 41.32, 41.95, 41.99, 42.51, 42.95, 42.99, 43.16, 43.95, 43.99, 44.82,
    44.95, 44.99, 45.17, 45.95, 45.99, 46.57, 47.82, 48.36, 49, 49.31, 50.03, 50.13,
    50.81, 55, 59, 69, 79, 85, 89, 99, 109, 119, 129, 189, 219, 259, 299, 319,
    349, 369, 399, 459, 499, 519, 599,
]

SHOPIFY_COLUMNS = [
    "Link-Href", "Handle", "Title", "Body (HTML)", "Collection",
    "Vendor", "Type", "Tags", "Published", "Option1 Name",
    "Option1 Value", "Option2 Name", "Option2 Value", "Option3 Name",
    "Option3 Value", "Variant SKU", "Variant Grams",
    "Variant Inventory Tracker", "Variant Inventory Qty",
    "Variant Inventory Policy", "Variant Fulfillment Service",
    "Variant Price", "Variant Compare At Price",
    "Variant Requires Shipping", "Variant Taxable",
    "Variant Barcode", "Image Src", "Image Position",
    "Image Alt Text", "Gift Card", "SEO Title",
    "SEO Description", "Google Shopping - Google Product Category",
    "Google Shopping - Gender", "Google Shopping - Age Group",
    "Google Shopping - MPN", "Google Shopping - AdWords Grouping",
    "Google Shopping - AdWords Labels", "Google Shopping - Condition",
    "Google Shopping - Custom Product", "Google Shopping - Custom Label 0",
    "Google Shopping - Custom Label 1", "Google Shopping - Custom Label 2",
    "Google Shopping - Custom Label 3", "Google Shopping - Custom Label 4",
    "Variant Image", "Variant Weight Unit", "Variant Tax Code",
    "Cost per item",
]


def _clean_cell(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def _detect_encoding(path):
    for enc in ("utf-8-sig", "gbk", "latin-1"):
        try:
            with open(path, encoding=enc, newline="") as f:
                f.read(2048)
            return enc
        except UnicodeDecodeError:
            continue
    return "utf-8-sig"


def _read_table(path):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(path)
    return pd.read_csv(path, encoding=_detect_encoding(path))


def _save_table(df, path):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xls"):
        df.to_excel(path, index=False)
    else:
        df.to_csv(path, index=False, encoding="utf-8-sig")


def _clean_handle_name(name):
    name = _clean_cell(name)
    cleaned = re.sub(r"[^a-zA-Z\s]", "", name).lower()
    return cleaned.replace(" ", "-").strip("-")


def _generate_unique_handle(base, used_handles):
    while True:
        digits = "".join(random.choices(string.digits, k=3))
        letters1 = "".join(random.choices(string.ascii_lowercase, k=2))
        letters2 = "".join(random.choices(string.ascii_lowercase, k=2))
        handle = f"{base}-{letters1}{digits}{letters2}" if base else f"-{letters1}{digits}{letters2}"
        if handle not in used_handles:
            used_handles.add(handle)
            return handle


def _generate_variant_sku(used_skus=None):
    while True:
        first_digit = random.choice("123456789")
        digits1 = "".join(random.choices(string.digits, k=2))
        letters1 = "".join(random.choices(string.ascii_uppercase, k=2))
        letter1 = random.choice(string.ascii_uppercase)
        digits2 = "".join(random.choices(string.digits, k=2))
        letters2 = "".join(random.choices(string.ascii_uppercase, k=2))
        digits3 = "".join(random.choices(string.digits, k=3))
        sku = f"{first_digit}{digits1}{letters1}-{letter1}{digits2}{letters2}-{digits3}{letters1}"
        if used_skus is None or sku not in used_skus:
            if used_skus is not None:
                used_skus.add(sku)
            return sku


def _parse_style_group(value):
    parts = [part.strip() for part in _clean_cell(value).split("#") if part.strip()]
    if not parts:
        return "", [""]
    if len(parts) == 1:
        return parts[0], [""]
    return parts[0], parts[1:]


def _parse_variant_segment(segment):
    text = _clean_cell(segment)
    image = ""
    if "@" in text:
        text, image = text.split("@", 1)
    price1 = ""
    price2 = ""
    price_parts = text.rsplit("$", 2)
    if len(price_parts) == 3:
        text, price1, price2 = price_parts
    option_parts = text.split("&") if text else []
    return {
        "option1_value": option_parts[0] if len(option_parts) > 0 else "",
        "option2_name": option_parts[1] if len(option_parts) > 1 else "",
        "option2_value": option_parts[2] if len(option_parts) > 2 else "",
        "option3_name": option_parts[3] if len(option_parts) > 3 else "",
        "option3_value": option_parts[4] if len(option_parts) > 4 else "",
        "price1": price1,
        "price2": price2,
        "image": image.strip(),
    }


def _styles_to_shopify(input_file, log):
    """导出 Shopify、价格匹配和 WP 文件，返回 Shopify 原价工作簿路径。"""
    input_path = Path(input_file)
    df = pd.read_excel(input_path)
    required = ["title", "name", "price1", "price2", "details", "src_links", "styles1", "styles2", "styles3"]
    missing = [col for col in required if col not in df.columns]
    if missing:
        raise ValueError(f"缺少必要列: {', '.join(missing)}")

    rows = []
    used_handles = set()
    used_skus = set()

    for _, item in df.iterrows():
        base_handle = _clean_handle_name(item.get("name", ""))
        handle = _generate_unique_handle(base_handle, used_handles)

        option1_name, option1_values = _parse_style_group(item.get("styles1", ""))
        option2_name, option2_values = _parse_style_group(item.get("styles2", ""))
        option3_name, option3_values = _parse_style_group(item.get("styles3", ""))
        if option1_name.startswith("$"):
            option1_name = ""
        if option2_name.startswith("$"):
            option2_name = ""
        if option3_name.startswith("$"):
            option3_name = ""
        images = [img.strip() for img in _clean_cell(item.get("src_links", "")).split("#") if img.strip()]
        combinations = list(itertools_product(option1_values or [""], option2_values or [""], option3_values or [""]))
        total_rows = max(len(combinations), len(images), 1)

        for index in range(total_rows):
            row = {col: "" for col in SHOPIFY_COLUMNS}
            row["Link-Href"] = _clean_cell(item.get("link-href", ""))
            row["Handle"] = handle
            row["Variant Weight Unit"] = "kg"

            if index == 0:
                row["Title"] = _clean_cell(item.get("name", ""))
                row["Body (HTML)"] = _clean_cell(item.get("details", ""))
                row["Collection"] = _clean_cell(item.get("title", ""))
                row["Vendor"] = _clean_cell(item.get("title", ""))
                row["Published"] = "TRUE"
                row["Gift Card"] = "FALSE"
                row["SEO Title"] = row["Title"]
                row["SEO Description"] = row["Title"]

            if index < len(combinations):
                combo = combinations[index]
                segment_info = _parse_variant_segment(combo[0])
                row["Option1 Name"] = option1_name
                row["Option1 Value"] = segment_info["option1_value"] or _clean_cell(combo[0])
                row["Option2 Name"] = segment_info["option2_name"] or option2_name
                row["Option2 Value"] = segment_info["option2_value"] or _clean_cell(combo[1])
                row["Option3 Name"] = segment_info["option3_name"] or option3_name
                row["Option3 Value"] = segment_info["option3_value"] or _clean_cell(combo[2])
                row["Variant SKU"] = _generate_variant_sku(used_skus)
                row["Variant Grams"] = 0
                row["Variant Inventory Tracker"] = "shopify"
                row["Variant Inventory Qty"] = 999
                row["Variant Inventory Policy"] = "deny"
                row["Variant Fulfillment Service"] = "manual"
                row["Variant Price"] = segment_info["price1"] or _clean_cell(item.get("price1", ""))
                row["Variant Compare At Price"] = segment_info["price2"] or _clean_cell(item.get("price2", ""))
                row["Variant Requires Shipping"] = "TRUE"
                row["Variant Taxable"] = "FALSE"
                row["Variant Image"] = segment_info["image"]

            if index < len(images):
                row["Image Src"] = images[index]
                row["Image Position"] = index + 1

            rows.append(row)

    xlsx_dir = input_path.parent / "xlsx"
    csv_dir = input_path.parent / "csv"
    xlsx_dir.mkdir(exist_ok=True)
    csv_dir.mkdir(exist_ok=True)

    output_path = xlsx_dir / f"{input_path.stem}_Shopify_原价.xlsx"
    result_df = pd.DataFrame(rows, columns=SHOPIFY_COLUMNS)
    result_df.to_excel(output_path, index=False)
    csv_path = csv_dir / f"{input_path.stem}_Shopify_原价.csv"
    result_df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    log(f"已生成 Shopify 文件: {output_path}")
    log(f"已生成 CSV 文件: {csv_path}")
    log(f"商品数: {len(df)}，输出行数: {len(rows)}")

    log("")
    log("开始价格匹配...")
    price_match_path, matched_df = _price_match(str(output_path), log, source_df=result_df)

    # ── Shopify 转 WP ──
    log("")
    log("开始 Shopify 转 WP...")
    try:
        _shopify_to_wp(str(csv_path), log, source_rows=result_df.to_dict("records"))
        if price_match_path:
            price_csv = Path(str(price_match_path).replace(str(price_match_path).split(".")[-1], "csv"))
            if price_csv.exists():
                _shopify_to_wp(str(price_csv), log, source_rows=matched_df.to_dict("records"))
        log("Shopify 转 WP 完成")
    except Exception as exc:
        log(f"Shopify 转 WP 出错: {exc}")
        traceback.print_exc()

    return output_path


def _price_match(input_file, log, source_df=None):
    """价格匹配并导出，返回输出路径和匹配后的 DataFrame。"""
    input_path = Path(input_file)
    df = (source_df.copy(deep=True) if source_df is not None else _read_table(input_path)).reset_index(drop=True)
    for col in ("Handle", "Variant Price"):
        if col not in df.columns:
            raise ValueError(f"缺少必要列: {col}")

    source_stem = input_path.stem
    for suffix in ("_Shopify_待匹配", "_Shopify_原价_拆分", "_Shopify_原价", "_原价"):
        if source_stem.endswith(suffix):
            source_stem = source_stem[: -len(suffix)]
            break

    # 使用 object 类型，既保留原始空值/文本，也允许写入小数价格。
    df["Variant Price"] = df["Variant Price"].astype(object)
    if "Variant Compare At Price" not in df.columns:
        df["Variant Compare At Price"] = None
    df["Variant Compare At Price"] = df["Variant Compare At Price"].astype(object)
    df["匹配价格"] = pd.Series([None] * len(df), dtype=object)

    global_used_prices = set()
    price_library_sorted = sorted(PRICE_LIBRARY)

    for handle, group in df.groupby("Handle"):
        if pd.isna(handle):
            continue
        valid_rows = []
        for idx, row in group.iterrows():
            price = pd.to_numeric(row.get("Variant Price"), errors="coerce")
            if pd.notna(price):
                valid_rows.append((idx, float(price)))
        if not valid_rows:
            continue

        available_prices = [p for p in price_library_sorted if p not in global_used_prices]
        for idx, original_price in sorted(valid_rows, key=lambda x: x[1]):
            lower_bound = original_price * 0.75
            upper_bound = original_price * 1.25
            matched = next((p for p in available_prices if lower_bound <= p <= upper_bound), None)
            if matched is not None:
                compare_at_price = round(matched * 1.2, 2)

                global_used_prices.add(matched)
                available_prices.remove(matched)

                # 保留匹配结果，方便核对价格匹配情况
                df.at[idx, "匹配价格"] = matched

                # 匹配成功后，直接更新 Shopify 售价与划线原价
                df.at[idx, "Variant Price"] = matched
                df.at[idx, "Variant Compare At Price"] = compare_at_price

    if input_path.parent.name == "xlsx":
        base_dir = input_path.parent.parent
    else:
        base_dir = input_path.parent
    xlsx_dir = base_dir / "xlsx"
    csv_dir = base_dir / "csv"
    xlsx_dir.mkdir(exist_ok=True)
    csv_dir.mkdir(exist_ok=True)

    output_path = xlsx_dir / f"{source_stem}_Shopify_价格匹配{input_path.suffix}"
    _save_table(df, output_path)
    csv_output = csv_dir / f"{source_stem}_Shopify_价格匹配.csv"
    df.to_csv(csv_output, index=False, encoding="utf-8-sig")
    log(f"已生成价格匹配 CSV: {csv_output}")

    unused_prices = [p for p in price_library_sorted if p not in global_used_prices]
    if unused_prices:
        unused_df = pd.DataFrame({"未匹配价格": unused_prices})
        unused_xlsx = xlsx_dir / f"{source_stem}_Shopify_未匹配到的价格{input_path.suffix}"
        _save_table(unused_df, unused_xlsx)
        unused_csv = csv_dir / f"{source_stem}_Shopify_未匹配到的价格.csv"
        unused_df.to_csv(unused_csv, index=False, encoding="utf-8-sig")
        log(f"未匹配价格文件: {unused_xlsx}")

    log(f"已生成价格匹配文件: {output_path}")
    log(f"已使用价格数: {len(global_used_prices)}，未使用价格数: {len(unused_prices)}")
    return output_path, df


# ════════════════════════════════════════════════════════════
# Shopify 转 WP（从 switch_gui.py 搬入）
# ════════════════════════════════════════════════════════════

WP_HEADERS = [
    "Type", "SKU", "Name", "Published", "Visibility in catalog", "Description", "In stock?", "Stock",
    "Sale price", "Regular price", "Categories", "Images", "Parent", "Position",
    "Attribute 1 name", "Attribute 1 value(s)", "Attribute 1 visible", "Attribute 1 global",
    "Attribute 2 name", "Attribute 2 value(s)", "Attribute 2 visible", "Attribute 2 global",
    "Attribute 3 name", "Attribute 3 value(s)", "Attribute 3 visible", "Attribute 3 global",
]


def _generate_parent_sku(handle="", counter=0, used_skus=None):
    """基于商品 Handle 生成确定性 parent SKU，同一 Handle 始终返回相同 SKU。"""
    while True:
        hash_input = f"{handle}_{counter}" if handle else str(counter)
        h = hashlib.md5(hash_input.encode()).hexdigest()
        digits = "".join(c for c in h if c.isdigit())
        if len(digits) < 8:
            digits = (digits + "0" * 8)[:8]
        sku = f"SKU{digits[:8]}"
        if used_skus is None or sku not in used_skus:
            if used_skus is not None:
                used_skus.add(sku)
            return sku
        counter += 1


def _wp_title(row):
    values = [row.get("Option1 Value", ""), row.get("Option2 Value", ""), row.get("Option3 Value", "")]
    values = [value for value in values if value]
    return "-" + ",".join(values) if values else ""


def _csv_join_unique(existing, value):
    value = _clean_cell(value)
    items = [item.strip() for item in _clean_cell(existing).split(",") if item.strip()]
    if value and value not in items:
        items.append(value)
    return ",".join(items)


def _read_csv_rows(path):
    encoding = _detect_encoding(path)
    with open(path, encoding=encoding, newline="") as f:
        return list(csv.DictReader(f)), encoding


def _read_shopify_rows(path):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xls"):
        df = pd.read_excel(path).fillna("")
        return [{key: _clean_cell(value) for key, value in row.items()} for row in df.to_dict("records")]
    rows, _ = _read_csv_rows(path)
    return rows


def _shopify_to_wp(input_file, log, source_rows=None):
    """将 Shopify CSV 转换为 WooCommerce (WP) 格式 CSV。"""
    input_path = Path(input_file)
    rows = (_read_shopify_rows(input_path) if source_rows is None else
            [{key: _clean_cell(value) for key, value in row.items()} for row in source_rows])
    products = {}
    used_parent_skus = set()

    for row in rows:
        handle = row.get("Handle", "").strip()
        if not handle:
            continue
        if handle not in products:
            parent = {}
            is_variable = row.get("Option1 Value", "") not in ("", "Default Title")
            parent["Type"] = "variable" if is_variable else "simple"
            parent["SKU"] = (
                _generate_parent_sku(handle, used_skus=used_parent_skus)
                if is_variable
                else (row.get("Variant SKU", "") or _generate_parent_sku(handle, used_skus=used_parent_skus))
            )
            parent["Name"] = row.get("Title", "")
            parent["Published"] = "1"
            parent["Visibility in catalog"] = "visible"
            parent["Description"] = row.get("Body (HTML)", "")
            parent["In stock?"] = "1"
            parent["Stock"] = "99999"
            parent["Sale price"] = "" if is_variable else row.get("Variant Price", "")
            parent["Regular price"] = "" if is_variable else (row.get("Variant Compare At Price", "") or row.get("Variant Price", ""))
            parent["Categories"] = row.get("Collection", "") or row.get("Type", "")
            parent["Images"] = row.get("Image Src", "")
            parent["Parent"] = ""
            parent["Position"] = "0"
            raw_opt1_name = row.get("Option1 Name", "")
            opt1_name = "" if (raw_opt1_name == "Title" or raw_opt1_name.startswith("$")) else raw_opt1_name
            parent["Attribute 1 name"] = opt1_name
            parent["Attribute 1 value(s)"] = "" if row.get("Option1 Value", "") == "Default Title" else row.get("Option1 Value", "")
            parent["Attribute 1 visible"] = "1" if parent["Attribute 1 name"] else ""
            parent["Attribute 1 global"] = "1" if parent["Attribute 1 name"] else ""
            parent["Attribute 2 name"] = row.get("Option2 Name", "")
            parent["Attribute 2 value(s)"] = row.get("Option2 Value", "")
            parent["Attribute 2 visible"] = "1" if row.get("Option2 Name", "") else ""
            parent["Attribute 2 global"] = "1" if row.get("Option2 Name", "") else ""
            parent["Attribute 3 name"] = row.get("Option3 Name", "")
            parent["Attribute 3 value(s)"] = row.get("Option3 Value", "")
            parent["Attribute 3 visible"] = "1" if row.get("Option3 Name", "") else ""
            parent["Attribute 3 global"] = "1" if row.get("Option3 Name", "") else ""
            products[handle] = [parent]

        parent = products[handle][0]
        if row.get("Image Src", ""):
            parent["Images"] = _csv_join_unique(parent["Images"], row.get("Image Src", ""))

        if row.get("Option1 Value", "") and row.get("Option1 Value", "") != "Default Title":
            parent["Attribute 1 value(s)"] = _csv_join_unique(
                parent["Attribute 1 value(s)"], row.get("Option1 Value", "")
            )
            parent["Attribute 2 value(s)"] = _csv_join_unique(
                parent["Attribute 2 value(s)"], row.get("Option2 Value", "")
            )
            parent["Attribute 3 value(s)"] = _csv_join_unique(
                parent["Attribute 3 value(s)"], row.get("Option3 Value", "")
            )
            variation = {
                "Type": "variation",
                "SKU": row.get("Variant SKU", ""),
                "Name": parent["Name"] + _wp_title(row),
                "Published": "1",
                "Visibility in catalog": "visible",
                "Description": "",
                "In stock?": "1",
                "Stock": "99999",
                "Sale price": row.get("Variant Price", ""),
                "Regular price": row.get("Variant Compare At Price", "") or row.get("Variant Price", ""),
                "Categories": parent.get("Categories", ""),
                "Images": row.get("Variant Image", ""),
                "Parent": parent["SKU"],
                "Position": str(len(products[handle])),
                "Attribute 1 name": parent["Attribute 1 name"],
                "Attribute 1 value(s)": row.get("Option1 Value", ""),
                "Attribute 1 visible": "",
                "Attribute 1 global": "1" if row.get("Option1 Value", "") else "",
                "Attribute 2 name": parent["Attribute 2 name"],
                "Attribute 2 value(s)": row.get("Option2 Value", ""),
                "Attribute 2 visible": "",
                "Attribute 2 global": "1" if row.get("Option2 Value", "") else "",
                "Attribute 3 name": parent["Attribute 3 name"],
                "Attribute 3 value(s)": row.get("Option3 Value", ""),
                "Attribute 3 visible": "",
                "Attribute 3 global": "1" if row.get("Option3 Value", "") else "",
            }
            products[handle].append(variation)

    # ── 价格校验：Sale price 不能大于 Regular price ──
    price_fix_count = 0
    for items in products.values():
        for item in items:
            try:
                sale = float(item.get("Sale price", "") or 0)
                regular = float(item.get("Regular price", "") or 0)
                if sale > 0 and regular > 0 and sale > regular:
                    item["Regular price"] = item["Sale price"]
                    price_fix_count += 1
            except (ValueError, TypeError):
                pass
    if price_fix_count:
        log(f"价格修正: {price_fix_count} 行 Sale price > Regular price，已统一为 Sale price")

    base_stem = input_path.stem
    for suffix in ("_Shopify_价格匹配", "_Shopify_原价"):
        if base_stem.endswith(suffix):
            base_stem = base_stem[: -len(suffix)]
            break
    output_dir = input_path.parent / f"wp-{base_stem}"
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / f"wp-{input_path.stem}.csv"
    with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=WP_HEADERS)
        writer.writeheader()
        for items in products.values():
            for item in items:
                writer.writerow({header: item.get(header, "") for header in WP_HEADERS})

    total_rows = sum(len(items) for items in products.values())
    log(f"已生成 WP 文件: {output_path}")
    log(f"商品组数: {len(products)}，输出行数: {total_rows}")
    return output_path


# ════════════════════════════════════════════════════════════
# 图形界面（清爽蓝白后台风）
# ════════════════════════════════════════════════════════════

class TextRedirector:
    """将 print 输出安全地转发到日志文本框。"""

    def __init__(self, text_widget):
        self.text_widget = text_widget

    def write(self, string):
        if not string:
            return
        try:
            self.text_widget.after(0, self._write, string)
        except (tk.TclError, RuntimeError):
            pass

    def _write(self, string):
        try:
            self.text_widget.config(state=tk.NORMAL)
            self.text_widget.insert(tk.END, string)
            self.text_widget.see(tk.END)
            self.text_widget.config(state=tk.DISABLED)
        except tk.TclError:
            pass

    def flush(self):
        pass


class ScraperApp:
    def __init__(self, root):
        self.root = root
        self.root.configure(bg=UI_COLORS["background"])

        self.style = ttk.Style()
        if "clam" in self.style.theme_names():
            self.style.theme_use("clam")
        self._configure_styles()

        self.is_running = False
        self.stop_event = threading.Event()
        self.file_rows = []
        self.task_files = []
        self.task_skip_options = ()
        self.task_error = ""
        self.failed_product_count = 0
        self.task_folder = ""
        self.active_file_index = None
        self.worker_thread = None
        self.shutdown_requested = False
        self.pending_saves = {}
        self.original_stdout = sys.stdout
        self.original_stderr = sys.stderr
        self.root.protocol("WM_DELETE_WINDOW", self._request_shutdown)

        self.folder_var = tk.StringVar()
        self.status_var = tk.StringVar(value="请选择分类文件夹后点击“读取文件夹”")
        self.progress_var = tk.DoubleVar(value=0)
        self.progress_percent_var = tk.StringVar(value="0%")
        self.file_count_var = tk.StringVar(value="0 个文件")

        self._build_log_window()
        sys.stdout = TextRedirector(self.log_text)
        sys.stderr = sys.stdout
        self._build_ui()
        self._center_window()

    def _can_start_task(self):
        if self.shutdown_requested or self.pending_saves:
            messagebox.showerror("尚未保存完成", "请先等待保存完成；保存失败时处理文件占用，再关闭窗口重试。")
            return False
        return True

    def _launch_task(self, stage, target, *, args=(), total=100, text="", clear_logs=False):
        self.pipeline_stage = stage
        self.is_running = True
        self.task_error = ""
        self.failed_product_count = 0
        self.stop_event.clear()
        if clear_logs:
            self._clear_logs()
        self.update_status(0, total, text)
        self.status_dot.configure(fg=UI_COLORS["primary"])
        self._apply_stage_controls()
        self.worker_thread = threading.Thread(target=target, args=args, daemon=True)
        self.worker_thread.start()

    def _queue_task_failure(self, callback, error):
        self.task_error = str(error) or type(error).__name__
        traceback.print_exc()
        self.root.after(0, callback, self.task_error)
        self.root.after(0, self._request_shutdown)

    def _fail_product_task(self, error):
        self.pipeline_stage = PipelineStage.FAILED
        if self.active_file_index is not None:
            self._set_file_status(self.active_file_index, "发生异常", "failed")
        self.update_status(0, 0, f"任务异常：{error}")
        self.status_dot.configure(fg=UI_COLORS["danger"])

    def _save_checkpoint(self, key, writer, *args):
        # 保存失败时保留数据引用，用户解除文件占用后可重试关闭。
        self.pending_saves[key] = (writer, args)
        writer(*args)
        self.pending_saves.pop(key, None)

    def _request_shutdown(self):
        if self.shutdown_requested:
            return
        self.shutdown_requested = True
        self.stop_event.set()
        detail = f"{self.task_error}；" if self.task_error else ""
        self.update_status(0, 0, f"{detail}正在保存已有数据，完成后关闭程序...")
        self.status_dot.configure(fg=self._result_color())
        self._poll_shutdown()

    def _poll_shutdown(self):
        if self.worker_thread is not None and self.worker_thread.is_alive():
            self.root.after(100, self._poll_shutdown)
            return
        try:
            for key, (writer, args) in list(self.pending_saves.items()):
                self._save_checkpoint(key, writer, *args)
        except Exception as exc:
            self.shutdown_requested = False
            self.task_error = f"保存失败：{exc}"
            self.status_dot.configure(fg=UI_COLORS["danger"])
            self._reset_buttons()
            traceback.print_exc()
            messagebox.showerror(
                "保存失败，尚未关闭",
                f"已有数据仍保留在内存中，请解除文件占用或修复保存目录后，再点击关闭窗口重试。\n{exc}",
            )
            return
        sys.stdout = self.original_stdout
        sys.stderr = self.original_stderr
        self.root.destroy()

    # ────────────────────────────────────────────────────────
    # 样式与通用组件
    # ────────────────────────────────────────────────────────
    def _configure_styles(self):
        self.style.configure(
            "App.TEntry",
            fieldbackground=UI_COLORS["card_alt"],
            foreground=UI_COLORS["text"],
            bordercolor=UI_COLORS["border"],
            lightcolor=UI_COLORS["border"],
            darkcolor=UI_COLORS["border"],
            insertcolor=UI_COLORS["primary"],
            padding=(10, 9),
            font=(UI_FONT, 10),
        )
        self.style.map(
            "App.TEntry",
            bordercolor=[("focus", UI_COLORS["border_focus"])],
            lightcolor=[("focus", UI_COLORS["border_focus"])],
            darkcolor=[("focus", UI_COLORS["border_focus"])],
        )
        self.style.configure(
            "Blue.Horizontal.TProgressbar",
            troughcolor="#E8EEF7",
            background=UI_COLORS["primary"],
            bordercolor="#E8EEF7",
            lightcolor=UI_COLORS["primary"],
            darkcolor=UI_COLORS["primary"],
            thickness=10,
        )
        self.style.configure(
            "App.Vertical.TScrollbar",
            troughcolor=UI_COLORS["card"],
            background="#CBD5E1",
            bordercolor=UI_COLORS["card"],
            arrowcolor=UI_COLORS["text_secondary"],
            gripcount=0,
        )
        self.style.map(
            "App.Vertical.TScrollbar",
            background=[("active", "#94A3B8")],
        )

    def _center_window(self):
        self.root.update_idletasks()
        width = self.root.winfo_width()
        height = self.root.winfo_height()
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        x = max((screen_width - width) // 2, 0)
        y = max((screen_height - height) // 2 - 20, 0)
        self.root.geometry(f"{width}x{height}+{x}+{y}")

    def _create_card(self, parent, **pack_kwargs):
        card = tk.Frame(
            parent,
            bg=UI_COLORS["card"],
            highlightbackground=UI_COLORS["border"],
            highlightcolor=UI_COLORS["border"],
            highlightthickness=1,
            bd=0,
        )
        card.pack(**pack_kwargs)
        return card

    def _create_button(
        self,
        parent,
        text,
        command,
        kind="secondary",
        padx=16,
        pady=9,
        font_size=10,
    ):
        styles = {
            "primary": {
                "bg": UI_COLORS["primary"],
                "fg": "#FFFFFF",
                "hover": UI_COLORS["primary_hover"],
                "disabled_bg": "#AFC7F5",
                "disabled_fg": "#F8FAFC",
                "border": 0,
            },
            "danger": {
                "bg": UI_COLORS["card"],
                "fg": UI_COLORS["danger"],
                "hover": UI_COLORS["danger_light"],
                "disabled_bg": UI_COLORS["card_alt"],
                "disabled_fg": UI_COLORS["text_muted"],
                "border": 1,
            },
            "secondary": {
                "bg": UI_COLORS["card_alt"],
                "fg": UI_COLORS["text_secondary"],
                "hover": "#EEF2F7",
                "disabled_bg": "#F1F5F9",
                "disabled_fg": UI_COLORS["text_muted"],
                "border": 1,
            },
            "ghost": {
                "bg": UI_COLORS["card"],
                "fg": UI_COLORS["primary"],
                "hover": UI_COLORS["primary_light"],
                "disabled_bg": UI_COLORS["card"],
                "disabled_fg": UI_COLORS["text_muted"],
                "border": 1,
            },
        }
        palette = styles[kind]
        button = tk.Button(
            parent,
            text=text,
            command=command,
            padx=padx,
            pady=pady,
            font=(UI_FONT, font_size, "bold" if kind == "primary" else "normal"),
            bg=palette["bg"],
            fg=palette["fg"],
            activebackground=palette["hover"],
            activeforeground=palette["fg"],
            disabledforeground=palette["disabled_fg"],
            relief=tk.FLAT,
            bd=0,
            highlightthickness=palette["border"],
            highlightbackground=(
                UI_COLORS["danger_border"]
                if kind == "danger"
                else UI_COLORS["border"]
            ),
            highlightcolor=(
                UI_COLORS["danger_border"]
                if kind == "danger"
                else UI_COLORS["border"]
            ),
            cursor="hand2",
        )
        button._normal_bg = palette["bg"]
        button._hover_bg = palette["hover"]
        button._disabled_bg = palette["disabled_bg"]
        button.bind("<Enter>", lambda _e, b=button: self._button_hover(b, True))
        button.bind("<Leave>", lambda _e, b=button: self._button_hover(b, False))
        return button

    @staticmethod
    def _button_hover(button, entering):
        if str(button.cget("state")) == str(tk.DISABLED):
            return
        button.configure(bg=button._hover_bg if entering else button._normal_bg)

    @staticmethod
    def _set_button_enabled(button, enabled):
        if enabled:
            button.configure(state=tk.NORMAL, bg=button._normal_bg, cursor="hand2")
        else:
            button.configure(state=tk.DISABLED, bg=button._disabled_bg, cursor="arrow")

    # ────────────────────────────────────────────────────────
    # 日志窗口
    # ────────────────────────────────────────────────────────
    def _build_log_window(self):
        self.log_win = tk.Toplevel(self.root)
        self.log_win.title("运行日志")
        self.log_win.geometry("820x480")
        self.log_win.minsize(650, 360)
        self.log_win.configure(bg=UI_COLORS["log_background"])
        self.log_win.protocol("WM_DELETE_WINDOW", self.log_win.withdraw)

        header = tk.Frame(self.log_win, bg=UI_COLORS["log_panel"], height=58)
        header.pack(fill=tk.X)
        header.pack_propagate(False)

        title_group = tk.Frame(header, bg=UI_COLORS["log_panel"])
        title_group.pack(side=tk.LEFT, padx=18, pady=10)
        tk.Label(
            title_group,
            text="运行日志",
            bg=UI_COLORS["log_panel"],
            fg="#F8FAFC",
            font=(UI_FONT, 12, "bold"),
        ).pack(anchor=tk.W)
        tk.Label(
            title_group,
            text="实时查看采集、保存与转换过程",
            bg=UI_COLORS["log_panel"],
            fg="#94A3B8",
            font=(UI_FONT, 9),
        ).pack(anchor=tk.W)

        log_actions = tk.Frame(header, bg=UI_COLORS["log_panel"])
        log_actions.pack(side=tk.RIGHT, padx=14)
        clear_btn = tk.Button(
            log_actions,
            text="清空",
            command=self._clear_logs,
            bg="#1E293B",
            fg="#CBD5E1",
            activebackground="#334155",
            activeforeground="#FFFFFF",
            relief=tk.FLAT,
            bd=0,
            padx=14,
            pady=6,
            font=(UI_FONT, 9),
            cursor="hand2",
        )
        clear_btn.pack(side=tk.LEFT, padx=(0, 8))
        close_btn = tk.Button(
            log_actions,
            text="关闭",
            command=self.log_win.withdraw,
            bg="#1E293B",
            fg="#CBD5E1",
            activebackground="#334155",
            activeforeground="#FFFFFF",
            relief=tk.FLAT,
            bd=0,
            padx=14,
            pady=6,
            font=(UI_FONT, 9),
            cursor="hand2",
        )
        close_btn.pack(side=tk.LEFT)

        log_body = tk.Frame(self.log_win, bg=UI_COLORS["log_background"])
        log_body.pack(fill=tk.BOTH, expand=True, padx=12, pady=12)

        self.log_text = scrolledtext.ScrolledText(
            log_body,
            state=tk.DISABLED,
            bg=UI_COLORS["log_background"],
            fg=UI_COLORS["log_text"],
            insertbackground="#FFFFFF",
            selectbackground="#1D4ED8",
            selectforeground="#FFFFFF",
            font=(MONO_FONT, 9),
            relief=tk.FLAT,
            bd=0,
            padx=12,
            pady=12,
            wrap=tk.WORD,
        )
        self.log_text.pack(fill=tk.BOTH, expand=True)
        self.log_win.withdraw()

    def _clear_logs(self):
        self.log_text.config(state=tk.NORMAL)
        self.log_text.delete("1.0", tk.END)
        self.log_text.config(state=tk.DISABLED)

    def show_logs(self):
        self.log_win.deiconify()
        self.log_win.lift()
        self.log_win.focus_force()

    # ────────────────────────────────────────────────────────
    # 主界面
    # ────────────────────────────────────────────────────────


    def _build_folder_card(self, parent):
        card = self._create_card(parent, fill=tk.X, pady=(0, 12))
        inner = tk.Frame(card, bg=UI_COLORS["card"])
        inner.pack(fill=tk.X, padx=18, pady=14)

        tk.Label(
            inner,
            text="选择分类文件夹",
            bg=UI_COLORS["card"],
            fg=UI_COLORS["text"],
            font=(UI_FONT, 11, "bold"),
        ).pack(anchor=tk.W)
        tk.Label(
            inner,
            text="文件夹中需要包含待处理的 .xlsx 文件",
            bg=UI_COLORS["card"],
            fg=UI_COLORS["text_muted"],
            font=(UI_FONT, 9),
        ).pack(anchor=tk.W, pady=(3, 10))

        input_row = tk.Frame(inner, bg=UI_COLORS["card"])
        input_row.pack(fill=tk.X)

        self.folder_entry = ttk.Entry(
            input_row,
            textvariable=self.folder_var,
            style="App.TEntry",
        )
        self.folder_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))

        self.browse_btn = self._create_button(
            input_row,
            "浏览文件夹",
            self._browse_folder,
            kind="secondary",
            padx=14,
            pady=8,
            font_size=9,
        )
        self.browse_btn.pack(side=tk.LEFT, padx=(0, 8))

        self.scan_btn = self._create_button(
            input_row,
            "读取文件夹",
            self._scan_files,
            kind="primary",
            padx=14,
            pady=8,
            font_size=9,
        )
        self.scan_btn.pack(side=tk.LEFT)

    def _build_file_card(self, parent):
        card = self._create_card(
            parent,
            fill=tk.BOTH,
            expand=True,
            pady=(0, 10),
        )
        card.configure(height=250)
        card.pack_propagate(False)
        self.file_card = card

        top = tk.Frame(card, bg=UI_COLORS["card"])
        top.pack(fill=tk.X, padx=18, pady=(10, 6))

        title_group = tk.Frame(top, bg=UI_COLORS["card"])
        title_group.pack(side=tk.LEFT)
        tk.Label(
            title_group,
            text="第二阶段 · 待处理文件",
            bg=UI_COLORS["card"],
            fg=UI_COLORS["text"],
            font=(UI_FONT, 11, "bold"),
        ).pack(anchor=tk.W)
        tk.Label(
            title_group,
            text="可设置跳过图片位置；-1 表示最后一张",
            bg=UI_COLORS["card"],
            fg=UI_COLORS["text_muted"],
            font=(UI_FONT, 9),
        ).pack(anchor=tk.W, pady=(3, 0))

        count_chip = tk.Label(
            top,
            textvariable=self.file_count_var,
            bg=UI_COLORS["primary_light"],
            fg=UI_COLORS["primary"],
            font=(UI_FONT, 9, "bold"),
            padx=10,
            pady=5,
        )
        count_chip.pack(side=tk.RIGHT)

        table_header = tk.Frame(card, bg=UI_COLORS["card_alt"], height=30)
        table_header.pack(fill=tk.X, padx=18)
        table_header.pack_propagate(False)
        table_header.grid_columnconfigure(1, weight=1)

        headers = [
            ("序号", 0, 54, tk.CENTER),
            ("文件名", 1, 0, tk.W),
            ("跳过图片位置", 2, 130, tk.W),
            ("状态", 3, 100, tk.CENTER),
        ]
        for text, column, width, anchor in headers:
            label = tk.Label(
                table_header,
                text=text,
                bg=UI_COLORS["card_alt"],
                fg=UI_COLORS["text_secondary"],
                font=(UI_FONT, 9, "bold"),
                anchor=anchor,
            )
            label.grid(
                row=0,
                column=column,
                sticky="nsew" if column == 1 else "ns",
                padx=(12 if column == 1 else 6),
            )
            if width:
                label.configure(width=max(width // 9, 1))

        list_shell = tk.Frame(card, bg=UI_COLORS["card"])
        list_shell.pack(fill=tk.BOTH, expand=True, padx=18, pady=(0, 8))

        self._canvas = tk.Canvas(
            list_shell,
            bg=UI_COLORS["card"],
            highlightthickness=0,
            bd=0,
        )
        scrollbar = ttk.Scrollbar(
            list_shell,
            orient=tk.VERTICAL,
            command=self._canvas.yview,
            style="App.Vertical.TScrollbar",
        )
        self.list_inner = tk.Frame(self._canvas, bg=UI_COLORS["card"])
        self._list_window = self._canvas.create_window(
            (0, 0),
            window=self.list_inner,
            anchor="nw",
        )

        self.list_inner.bind(
            "<Configure>",
            lambda _e: self._canvas.configure(scrollregion=self._canvas.bbox("all")),
        )
        self._canvas.bind(
            "<Configure>",
            lambda e: self._canvas.itemconfigure(self._list_window, width=e.width),
        )
        self._canvas.configure(yscrollcommand=scrollbar.set)

        self._canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self._canvas.bind("<Enter>", self._bind_mousewheel)
        self._canvas.bind("<Leave>", self._unbind_mousewheel)
        self.list_inner.bind("<Enter>", self._bind_mousewheel)
        self.list_inner.bind("<Leave>", self._unbind_mousewheel)

        self._show_empty_file_state()

    def _build_progress_card(self, parent):
        card = self._create_card(parent, fill=tk.X, pady=(0, 12))
        inner = tk.Frame(card, bg=UI_COLORS["card"])
        inner.pack(fill=tk.X, padx=18, pady=14)

        progress_header = tk.Frame(inner, bg=UI_COLORS["card"])
        progress_header.pack(fill=tk.X)
        tk.Label(
            progress_header,
            text="任务进度",
            bg=UI_COLORS["card"],
            fg=UI_COLORS["text"],
            font=(UI_FONT, 10, "bold"),
        ).pack(side=tk.LEFT)
        tk.Label(
            progress_header,
            textvariable=self.progress_percent_var,
            bg=UI_COLORS["card"],
            fg=UI_COLORS["primary"],
            font=(UI_FONT, 11, "bold"),
        ).pack(side=tk.RIGHT)

        self.progress_bar = ttk.Progressbar(
            inner,
            orient=tk.HORIZONTAL,
            mode="determinate",
            variable=self.progress_var,
            style="Blue.Horizontal.TProgressbar",
        )
        self.progress_bar.pack(fill=tk.X, pady=(10, 8), ipady=1)

        status_row = tk.Frame(inner, bg=UI_COLORS["card"])
        status_row.pack(fill=tk.X)
        self.status_dot = tk.Label(
            status_row,
            text="●",
            bg=UI_COLORS["card"],
            fg=UI_COLORS["text_muted"],
            font=(UI_FONT, 8),
        )
        self.status_dot.pack(side=tk.LEFT, padx=(0, 7))
        self.status_label = tk.Label(
            status_row,
            textvariable=self.status_var,
            bg=UI_COLORS["card"],
            fg=UI_COLORS["text_secondary"],
            font=(UI_FONT, 9),
            anchor=tk.W,
        )
        self.status_label.pack(side=tk.LEFT, fill=tk.X, expand=True)


    # ────────────────────────────────────────────────────────
    # 文件列表交互
    # ────────────────────────────────────────────────────────
    def _bind_mousewheel(self, _event=None):
        self.root.bind_all("<MouseWheel>", self._on_mousewheel)
        self.root.bind_all("<Button-4>", self._on_mousewheel_linux)
        self.root.bind_all("<Button-5>", self._on_mousewheel_linux)

    def _unbind_mousewheel(self, _event=None):
        self.root.unbind_all("<MouseWheel>")
        self.root.unbind_all("<Button-4>")
        self.root.unbind_all("<Button-5>")

    def _on_mousewheel(self, event):
        if event.delta:
            self._canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

    def _on_mousewheel_linux(self, event):
        direction = -1 if event.num == 4 else 1
        self._canvas.yview_scroll(direction, "units")

    def _show_empty_file_state(self):
        for widget in self.list_inner.winfo_children():
            widget.destroy()
        empty = tk.Frame(self.list_inner, bg=UI_COLORS["card"])
        empty.pack(fill=tk.BOTH, expand=True, pady=18)
        tk.Label(
            empty,
            text="尚未读取文件",
            bg=UI_COLORS["card"],
            fg=UI_COLORS["text_secondary"],
            font=(UI_FONT, 10, "bold"),
        ).pack()
        tk.Label(
            empty,
            text="选择分类文件夹后，文件会显示在这里",
            bg=UI_COLORS["card"],
            fg=UI_COLORS["text_muted"],
            font=(UI_FONT, 9),
        ).pack(pady=(5, 0))


    def _scan_files(self):
        folder = self.folder_var.get().strip()
        if not folder or not os.path.isdir(folder):
            messagebox.showerror("路径无效", "请先选择一个有效的分类文件夹。")
            return

        xlsx_files = sorted(
            f
            for f in os.listdir(folder)
            if f.lower().endswith(".xlsx") and not f.startswith("~")
        )

        if not xlsx_files:
            self.file_rows.clear()
            self.file_count_var.set("0 个文件")
            self._show_empty_file_state()
            self.update_status(0, 100, "该文件夹中没有找到 .xlsx 文件")
            messagebox.showinfo("未找到文件", "该文件夹下没有 .xlsx 文件。")
            return

        for widget in self.list_inner.winfo_children():
            widget.destroy()
        self.file_rows.clear()

        self.list_inner.grid_columnconfigure(0, minsize=54)
        self.list_inner.grid_columnconfigure(1, weight=1)
        self.list_inner.grid_columnconfigure(2, minsize=130)
        self.list_inner.grid_columnconfigure(3, minsize=100)

        for index, filename in enumerate(xlsx_files):
            file_path = os.path.join(folder, filename)
            skip_var = tk.StringVar()
            status_var = tk.StringVar(value="等待中")

            row_bg = UI_COLORS["card"] if index % 2 == 0 else "#FBFCFE"
            row = tk.Frame(
                self.list_inner,
                bg=row_bg,
                height=48,
                highlightbackground="#EEF2F7",
                highlightthickness=0,
            )
            row.grid(row=index, column=0, columnspan=4, sticky="ew")
            row.grid_columnconfigure(1, weight=1)
            row.grid_propagate(False)

            number_label = tk.Label(
                row,
                text=f"{index + 1:02d}",
                bg=row_bg,
                fg=UI_COLORS["text_muted"],
                font=(MONO_FONT, 9),
                anchor=tk.CENTER,
            )
            number_label.grid(row=0, column=0, sticky="nsew", padx=6)

            file_label = tk.Label(
                row,
                text=filename,
                bg=row_bg,
                fg=UI_COLORS["text"],
                font=(UI_FONT, 9),
                anchor=tk.W,
            )
            file_label.grid(row=0, column=1, sticky="nsew", padx=(12, 12))

            skip_shell = tk.Frame(
                row,
                bg=row_bg,
                width=130,
            )
            skip_shell.grid(row=0, column=2, sticky="w", padx=(0, 10))
            skip_entry = tk.Entry(
                skip_shell,
                textvariable=skip_var,
                width=11,
                bg="#FFFFFF",
                fg=UI_COLORS["text"],
                insertbackground=UI_COLORS["primary"],
                relief=tk.FLAT,
                bd=0,
                highlightthickness=1,
                highlightbackground=UI_COLORS["border"],
                highlightcolor=UI_COLORS["border_focus"],
                font=(UI_FONT, 9),
            )
            skip_entry.pack(side=tk.LEFT, ipady=5)
            tk.Label(
                skip_shell,
                text="如 1,3,-1",
                bg=row_bg,
                fg=UI_COLORS["text_muted"],
                font=(UI_FONT, 8),
            ).pack(side=tk.LEFT, padx=(6, 0))

            status_label = tk.Label(
                row,
                textvariable=status_var,
                bg="#F1F5F9",
                fg=UI_COLORS["text_secondary"],
                font=(UI_FONT, 8, "bold"),
                padx=9,
                pady=4,
                width=7,
            )
            status_label.grid(row=0, column=3, padx=(0, 8))

            separator = tk.Frame(self.list_inner, bg="#EEF2F7", height=1)
            separator.grid(row=index, column=0, columnspan=4, sticky="sew")

            self.file_rows.append(
                {
                    "path": file_path,
                    "skip_var": skip_var,
                    "skip_entry": skip_entry,
                    "status_var": status_var,
                    "status_label": status_label,
                }
            )

        self.file_count_var.set(f"{len(xlsx_files)} 个文件")
        self.update_status(0, 100, f"已加载 {len(xlsx_files)} 个文件，等待开始")

    def _set_file_status(self, index, text, state="waiting"):
        if index is None or not (0 <= index < len(self.file_rows)):
            return
        row = self.file_rows[index]
        palettes = {
            "waiting": ("#F1F5F9", UI_COLORS["text_secondary"]),
            "running": (UI_COLORS["primary_light"], UI_COLORS["primary"]),
            "success": (UI_COLORS["success_light"], UI_COLORS["success"]),
            "skipped": (UI_COLORS["warning_light"], UI_COLORS["warning"]),
            "failed": (UI_COLORS["danger_light"], UI_COLORS["danger"]),
            "partial": (UI_COLORS["warning_light"], UI_COLORS["warning"]),
            "stopped": (UI_COLORS["warning_light"], UI_COLORS["warning"]),
        }
        bg, fg = palettes.get(state, palettes["waiting"])
        row["status_var"].set(text)
        row["status_label"].configure(bg=bg, fg=fg)

    # ────────────────────────────────────────────────────────
    # 任务控制
    # ────────────────────────────────────────────────────────
    def _start(self):
        if not self._can_start_task():
            return
        if not self.file_rows:
            messagebox.showerror("无法开始", "请先选择并读取分类文件夹。")
            return

        self.task_folder = self.folder_var.get().strip()
        self.task_skip_options = tuple(SKIP_OPTIONS or ())
        self.task_files = [
            {
                "path": row["path"],
                "skip": row["skip_var"].get().strip(),
            }
            for row in self.file_rows
        ]

        self.active_file_index = None
        for index in range(len(self.file_rows)):
            self._set_file_status(index, "等待中", "waiting")
        self._launch_task(
            PipelineStage.PROCESSING_PRODUCTS, self._run_task,
            text="正在初始化任务...", clear_logs=True,
        )

    def _stop(self):
        if self.is_running:
            self.stop_event.set()
            self._set_button_enabled(self.stop_btn, False)
            self.status_dot.configure(fg=UI_COLORS["warning"])
            self.update_status(
                self.progress_var.get(),
                self.progress_bar["maximum"],
                "正在停止并保存当前进度...",
            )
            print("\n[系统] 用户已手动触发停止，保存后关闭程序...")
            self._request_shutdown()


    def update_status(self, current, total, text=""):
        safe_total = max(float(total or 0), 0)
        safe_current = max(float(current or 0), 0)

        if safe_total > 0:
            safe_current = min(safe_current, safe_total)
            percent = int(safe_current / safe_total * 100)
            self.progress_bar["maximum"] = safe_total
            self.progress_var.set(safe_current)
            self.progress_percent_var.set(f"{percent}%")
            self.status_var.set(f"{text}  ·  {int(safe_current)}/{int(safe_total)}")
        else:
            self.progress_bar["maximum"] = 100
            self.progress_var.set(0)
            self.progress_percent_var.set("0%")
            self.status_var.set(text)

    def _result_color(self):
        if self.task_error or self.pending_saves:
            return UI_COLORS["danger"]
        if self.stop_event.is_set() or self.failed_product_count:
            return UI_COLORS["warning"]
        return UI_COLORS["success"]


    # ────────────────────────────────────────────────────────
    # 采集主任务（业务逻辑保持不变，仅补充界面状态）
    # ────────────────────────────────────────────────────────

    def _run_task(self):
        try:
            folder = self.task_folder
            rates = fetch_exchange_rates()

            all_results = []
            saved_merge_count = 0
            merge_path = str(map_to_output_path(Path(folder)) / f"{Path(folder).name}_合并.xlsx")
            grand_success = 0
            grand_fail = 0
            file_count = len(self.task_files)

            for fi, file_info in enumerate(self.task_files):
                if self.stop_event.is_set():
                    break

                self.active_file_index = fi
                self.root.after(0, self._set_file_status, fi, "处理中", "running")

                fpath = file_info["path"]
                fname = os.path.basename(fpath)
                skip_str = file_info["skip"]
                file_skip = None
                if skip_str:
                    try:
                        file_skip = [
                            int(value.strip())
                            for value in skip_str.split(",")
                            if value.strip().isdigit()
                        ]
                    except Exception:
                        file_skip = None

                # 输出路径：01INPUT_XLSX → 02OUTPUT_XLSX
                output_file = str(map_to_output_path(Path(fpath)))
                os.makedirs(os.path.dirname(output_file), exist_ok=True)

                print(f"\n{'=' * 50}")
                print(f"[{fi + 1}/{file_count}] {fname}")
                print(f"  输出: {output_file}")
                if file_skip:
                    print(f"  跳过图片: {file_skip}")

                # 读取 + 去重
                targets = read_targets_from_excel(fpath)
                json_url_groups, dup_count = group_targets_by_json_url(targets)
                total = len(json_url_groups)
                print(f"  共 {total} 个唯一商品（跳过 {dup_count} 个重复）")

                if total == 0:
                    print("  无商品，跳过")
                    self.root.after(0, self._set_file_status, fi, "已跳过", "skipped")
                    continue

                # 相同配置才续传；变更配置或旧版无配置记录时重新采集并替换。
                current_config = crawl_config_key(skip_str, self.task_skip_options)
                results = []
                if os.path.exists(output_file):
                    _, results = load_xlsx(output_file)
                result_indexes = {
                    to_shopify_json_url(str(row["link-href"])): index
                    for index, row in enumerate(results) if row.get("link-href")
                }
                todo = {
                    url: group for url, group in json_url_groups.items()
                    if url not in result_indexes
                    or results[result_indexes[url]].get("_crawl_config") != current_config
                    or not results[result_indexes[url]].get("name")
                }
                print(f"  本次配置：跳过 {total - len(todo)} 个已完成商品，处理 {len(todo)} 个")
                if not todo:
                    all_results.extend(results)
                    self.root.after(0, self._set_file_status, fi, "已完成", "success")
                    continue

                # 启动浏览器
                self.root.after(
                    0,
                    self.update_status,
                    0,
                    len(todo),
                    f"[{fi + 1}/{file_count}] 正在启动浏览器 · {fname}",
                )
                chrome = create_browser()
                success_count = 0
                saved_success_count = 0
                fail_count = 0

                try:
                    tab = chrome.latest_tab
                    for index, (json_url, group) in enumerate(todo.items(), 1):
                        if self.stop_event.is_set():
                            break

                        table_title = group["title"]
                        original_url = group["original_url"]

                        self.root.after(
                            0,
                            self.update_status,
                            index,
                            len(todo),
                            f"[{fi + 1}/{file_count}] {fname} · {table_title[:24]}",
                        )
                        print(f"  [{index}/{len(todo)}] {table_title}")

                        retry_count = 0
                        item = None

                        while retry_count <= PRODUCT_MAX_RETRIES:
                            if self.stop_event.is_set():
                                break
                            status, data = fetch_json(tab, json_url)
                            if status == 200:
                                try:
                                    item = parse_product(
                                        data, table_title, original_url,
                                        skip_positions=file_skip,
                                        skip_options=self.task_skip_options,
                                        rates=rates, raw_skip_spec=skip_str,
                                    )
                                except InvalidProductError as exc:
                                    print(f"    无效商品 JSON，跳过当前商品: {exc}")
                                break
                            if status == 429:
                                retry_count += 1
                                if retry_count > PRODUCT_MAX_RETRIES:
                                    print(f"    429 超过 {PRODUCT_MAX_RETRIES} 次，放弃")
                                    break
                                delay = RETRY_BASE_DELAY * retry_count
                                print(f"    429 第{retry_count}次，{delay}秒后重试...")
                                self.stop_event.wait(delay)
                            else:
                                # 404 已确认商品不存在，不进入重试
                                if status == 404:
                                    print("    商品不存在 (status=404)，跳过当前商品")
                                    break

                                retry_count += 1
                                if retry_count > PRODUCT_MAX_RETRIES:
                                    print(
                                        f"    失败 {PRODUCT_MAX_RETRIES} 次，放弃 "
                                        f"(status={status})"
                                    )
                                    break
                                delay = 5 * retry_count
                                print(f"    失败 (status={status})，{delay}秒后重试...")
                                self.stop_event.wait(delay)
                            self.stop_event.wait(DELAY)

                        if item:
                            item["_crawl_config"] = current_config
                            if json_url in result_indexes:
                                results[result_indexes[json_url]] = item
                            else:
                                result_indexes[json_url] = len(results)
                                results.append(item)
                            success_count += 1
                            print(f"    成功: {item['name'][:40]}")
                        elif not self.stop_event.is_set():
                            fail_count += 1
                            self.failed_product_count += 1

                        if self.stop_event.is_set():
                            break

                        if success_count - saved_success_count >= SAVE_EVERY:
                            self._save_checkpoint(output_file, save_xlsx, results, output_file)
                            saved_success_count = success_count
                            print(f"    [自动保存] {len(results)} 条")

                        self.stop_event.wait(DELAY)

                finally:
                    all_results.extend(results)
                    # 先登记两份快照，即使第一份写入失败，也能在退出前重试全部数据。
                    if success_count != saved_success_count:
                        self.pending_saves[output_file] = (save_xlsx, (results, output_file))
                    if len(all_results) != saved_merge_count:
                        self.pending_saves[merge_path] = (save_xlsx, (list(all_results), merge_path))
                    try:
                        for path in (output_file, merge_path):
                            if path in self.pending_saves:
                                writer, args = self.pending_saves[path]
                                self._save_checkpoint(path, writer, *args)
                                if path == merge_path:
                                    saved_merge_count = len(all_results)
                                print(f"  数据已保存: {path}")
                    finally:
                        print("  正在关闭浏览器...")
                        try:
                            chrome.quit()
                        except Exception:
                            traceback.print_exc()

                grand_success += success_count
                grand_fail += fail_count

                if self.stop_event.is_set():
                    self.root.after(0, self._set_file_status, fi, "已停止", "stopped")
                    break

                if fail_count > 0 and success_count == 0:
                    self.root.after(0, self._set_file_status, fi, "采集失败", "failed")
                elif fail_count > 0:
                    self.root.after(0, self._set_file_status, fi, "部分失败", "partial")
                else:
                    self.root.after(0, self._set_file_status, fi, "已完成", "success")

            # 将中断后尚未处理的文件保持为等待状态
            if self.stop_event.is_set() and self.active_file_index is not None:
                for pending_index in range(self.active_file_index + 1, len(self.file_rows)):
                    self.root.after(
                        0,
                        self._set_file_status,
                        pending_index,
                        "未处理",
                        "waiting",
                    )

            # 合并
            if all_results:
                if len(all_results) != saved_merge_count:
                    self._save_checkpoint(merge_path, save_xlsx, list(all_results), merge_path)
                print(f"\n{'=' * 50}")
                print(f"合并完成: {merge_path} ({len(all_results)} 条)")

                if self.stop_event.is_set():
                    return

                self._on_merge_ready(merge_path)

            final_text = (
                f"总计：{grand_success} 成功，{grand_fail} 失败，"
                f"{len(all_results)} 条数据"
            )
            if self.stop_event.is_set():
                final_text = "任务已中断 · " + final_text
            self.root.after(
                0,
                self.update_status,
                grand_success + grand_fail,
                grand_success + grand_fail,
                final_text,
            )
            print(f"\n{final_text}")

        except Exception as exc:
            print(f"发生异常: {exc}")
            self._queue_task_failure(self._fail_product_task, exc)

        finally:
            self.is_running = False
            self.root.after(0, self._reset_buttons)


INPUT_DIR_NAME = "01INPUT_XLSX"
OUTPUT_DIR_NAME = "02OUTPUT_XLSX"


class PipelineStage(str, Enum):
    IDLE = "idle"
    COLLECTING_LINKS = "collecting_links"
    WAITING_CONTINUE = "waiting_continue"
    PROCESSING_PRODUCTS = "processing_products"
    WAITING_CONVERSION = "waiting_conversion"
    CONVERTING = "converting"
    COMPLETED = "completed"
    FAILED = "failed"
    STOPPED = "stopped"


def map_to_output_path(path: PurePath) -> PurePath:
    """把路径中精确的 01INPUT_XLSX 组件映射为 02OUTPUT_XLSX。"""
    parts = list(path.parts)
    try:
        marker_index = parts.index(INPUT_DIR_NAME)
    except ValueError as exc:
        raise ValueError(f"路径必须位于 {INPUT_DIR_NAME} 下") from exc
    parts[marker_index] = OUTPUT_DIR_NAME
    return type(path)(*parts)


def validate_input_folder(path: PurePath) -> PurePath:
    """验证路径是 01INPUT_XLSX 下的具体分类目录。"""
    parts = list(path.parts)
    try:
        marker_index = parts.index(INPUT_DIR_NAME)
    except ValueError as exc:
        raise ValueError(f"路径必须位于 {INPUT_DIR_NAME} 下") from exc
    if marker_index == len(parts) - 1:
        raise ValueError("请选择 01INPUT_XLSX 下的具体分类目录")
    return path


def parse_category_pages(raw: str) -> list[tuple[str, str]]:
    """解析每行“标题, URL”的分类列表。"""
    pages: list[tuple[str, str]] = []
    for source_line in raw.splitlines():
        line = source_line.strip()
        if not line or line.startswith("#"):
            continue
        match = re.search(r'(https?://[^\s"\')\]]+)', line)
        if not match:
            continue
        url = match.group(1)
        title = ""
        url_start = line.find(url)
        if url_start > 0:
            title_part = line[:url_start].strip(' \t,("\'')
            title = re.sub(r'[",)\]]+$', "", title_part).strip()
        if not title:
            title = urlparse(url).path.rstrip("/").split("/")[-1]
        pages.append((title, url))
    return pages


def next_stage_after_link_collection(
    link_count: int,
    stopped: bool,
) -> PipelineStage:
    if stopped:
        return PipelineStage.STOPPED
    if link_count <= 0:
        return PipelineStage.FAILED
    return PipelineStage.WAITING_CONTINUE


def can_continue_processing(stage: PipelineStage) -> bool:
    return stage == PipelineStage.WAITING_CONTINUE


def prepare_data_for_image_skips(data, configured_positions, raw_skip_spec):
    """把 -1 动态展开为当前商品最后一张图片的 position。"""
    positions = []
    for value in configured_positions or []:
        try:
            position = int(value)
        except (TypeError, ValueError):
            continue
        if position > 0 and position not in positions:
            positions.append(position)

    tokens = {
        token.strip()
        for token in str(raw_skip_spec or "").split(",")
        if token.strip()
    }
    if "-1" not in tokens:
        return data, positions

    product = data.get("product", {}) if isinstance(data, dict) else {}
    images = product.get("images", []) if isinstance(product, dict) else []
    if not images:
        return data, positions

    positioned_images = []
    for index, image in enumerate(images):
        try:
            image_position = int(image.get("position") or 0)
        except (AttributeError, TypeError, ValueError):
            image_position = 0
        positioned_images.append((image_position, index))

    positive_positions = [item for item in positioned_images if item[0] > 0]
    if positive_positions:
        last_position, _ = max(positive_positions, key=lambda item: item[0])
        if last_position not in positions:
            positions.append(last_position)
        return data, positions

    # 极少数响应不含 position；仅给列表最后一张添加临时位置，
    # 避免把所有 position 缺失的图片一起过滤掉。
    prepared_data = dict(data)
    prepared_product = dict(product)
    prepared_images = [dict(image) for image in images]
    fallback_position = 1
    prepared_images[-1]["position"] = fallback_position
    prepared_product["images"] = prepared_images
    prepared_data["product"] = prepared_product
    if fallback_position not in positions:
        positions.append(fallback_position)
    return prepared_data, positions


def load_existing_link_rows(output_folder):
    rows = []
    seen = set()
    for path in sorted(Path(output_folder).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        workbook = openpyxl.load_workbook(path, read_only=True)
        try:
            values = workbook.active.iter_rows(values_only=True)
            headers = next(values, ())
            if tuple(headers[:2]) != ("title", "link"):
                raise ValueError(f"已有链接文件列格式不正确（需要 title/link）: {path}")
            for row in values:
                link = str(row[1]).strip() if len(row) > 1 and row[1] else ""
                if not link:
                    continue
                key = to_shopify_json_url(link)
                if key not in seen:
                    seen.add(key)
                    rows.append({"title": str(row[0] or ""), "link": link})
        finally:
            workbook.close()
    return rows


def write_link_workbooks(
    rows: list[dict[str, str]],
    output_folder: Path,
) -> list[Path]:
    """按商品域名写入第一阶段的 title/link 工作簿。"""
    domain_groups: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        domain = urlparse(row["link"]).netloc
        if domain:
            domain_groups.setdefault(domain, []).append(row)

    output_folder.mkdir(parents=True, exist_ok=True)
    saved_paths: list[Path] = []
    for domain, domain_rows in domain_groups.items():
        file_path = output_folder / f"{domain}.xlsx"
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.append(["title", "link"])
        for row in domain_rows:
            worksheet.append([row.get("title", ""), row.get("link", "")])
        save_workbook_atomic(workbook, file_path)
        saved_paths.append(file_path)
    return saved_paths


class ShopifyPipelineApp(ScraperApp):
    """把链接采集和商品处理串联到同一个 Tk 窗口。"""

    def __init__(self, root: tk.Tk):
        self.pipeline_stage = PipelineStage.IDLE
        self.max_pages_var = tk.StringVar(master=root, value="")
        self.xpath_var = tk.StringVar(master=root, value="")
        self.conversion_file_var = tk.StringVar(master=root, value="")
        self.last_merge_path = ""
        root.title("Shopify 一体化采集工作台")
        root.geometry("1040x980")
        root.minsize(940, 760)
        super().__init__(root)
        self._apply_stage_controls()

    def _build_ui(self):
        page = tk.Frame(self.root, bg=UI_COLORS["background"])
        page.pack(fill=tk.BOTH, expand=True)

        content = tk.Frame(page, bg=UI_COLORS["background"])
        content.pack(fill=tk.BOTH, expand=True, padx=24, pady=(16, 14))

        footer = tk.Frame(content, bg=UI_COLORS["background"])
        footer.pack(side=tk.BOTTOM, fill=tk.X)
        self._build_progress_card(footer)
        self._build_action_bar(footer)

        body = tk.Frame(content, bg=UI_COLORS["background"])
        body.pack(side=tk.TOP, fill=tk.BOTH, expand=True)
        self._build_header(body)
        self._build_folder_card(body)
        self._build_link_card(body)
        self._build_file_card(body)
        self.file_card.configure(height=170)
        self._build_conversion_card(body)

    def _build_header(self, parent):
        colors = UI_COLORS
        header = tk.Frame(parent, bg=colors["background"])
        header.pack(fill=tk.X, pady=(0, 12))

        logo = tk.Label(
            header,
            text="S",
            width=3,
            bg=colors["primary"],
            fg="#FFFFFF",
            font=(UI_FONT, 15, "bold"),
        )
        logo.pack(side=tk.LEFT, padx=(0, 12), ipady=6)

        title_group = tk.Frame(header, bg=colors["background"])
        title_group.pack(side=tk.LEFT)
        tk.Label(
            title_group,
            text="Shopify 一体化采集工作台",
            bg=colors["background"],
            fg=colors["text"],
            font=(UI_FONT, 18, "bold"),
        ).pack(anchor=tk.W)
        tk.Label(
            title_group,
            text="链接采集 → 图片设置 → 商品采集、合并与转换",
            bg=colors["background"],
            fg=colors["text_secondary"],
            font=(UI_FONT, 9),
        ).pack(anchor=tk.W, pady=(3, 0))

        self.header_log_btn = self._create_button(
            header,
            "查看运行日志",
            self.show_logs,
            kind="ghost",
            padx=14,
            pady=8,
            font_size=9,
        )
        self.header_log_btn.pack(side=tk.RIGHT, pady=4)

    def _build_link_card(self, parent):
        colors = UI_COLORS
        card = self._create_card(parent, fill=tk.X, pady=(0, 12))
        inner = tk.Frame(card, bg=colors["card"])
        inner.pack(fill=tk.X, padx=18, pady=12)

        heading = tk.Frame(inner, bg=colors["card"])
        heading.pack(fill=tk.X, pady=(0, 8))
        tk.Label(
            heading,
            text="第一阶段 · 详情页链接采集",
            bg=colors["card"],
            fg=colors["text"],
            font=(UI_FONT, 11, "bold"),
        ).pack(side=tk.LEFT)
        tk.Label(
            heading,
            text="每行：标题, URL",
            bg=colors["card"],
            fg=colors["text_muted"],
            font=(UI_FONT, 9),
        ).pack(side=tk.RIGHT)

        fields = tk.Frame(inner, bg=colors["card"])
        fields.pack(fill=tk.X, pady=(0, 8))
        fields.grid_columnconfigure(0, weight=1)
        fields.grid_columnconfigure(1, weight=1)

        tk.Label(
            fields,
            text="XPath（可留空）",
            bg=colors["card"],
            fg=colors["text_secondary"],
            font=(UI_FONT, 9),
        ).grid(row=0, column=0, sticky="w")
        tk.Label(
            fields,
            text="最大页数（留空表示不限制）",
            bg=colors["card"],
            fg=colors["text_secondary"],
            font=(UI_FONT, 9),
        ).grid(row=0, column=1, sticky="w", padx=(10, 0))

        self.xpath_entry = ttk.Entry(
            fields,
            textvariable=self.xpath_var,
            style="App.TEntry",
        )
        self.xpath_entry.grid(row=1, column=0, sticky="ew", pady=(3, 0))
        self.max_pages_entry = ttk.Entry(
            fields,
            textvariable=self.max_pages_var,
            style="App.TEntry",
        )
        self.max_pages_entry.grid(
            row=1,
            column=1,
            sticky="ew",
            padx=(10, 0),
            pady=(3, 0),
        )

        self.pages_text = scrolledtext.ScrolledText(
            inner,
            height=5,
            wrap=tk.NONE,
            bg=colors["card_alt"],
            fg=colors["text"],
            insertbackground=colors["primary"],
            font=(MONO_FONT, 9),
            relief=tk.FLAT,
            bd=0,
            highlightthickness=1,
            highlightbackground=colors["border"],
            padx=10,
            pady=7,
        )
        self.pages_text.pack(fill=tk.X)

    def _build_conversion_card(self, parent):
        colors = UI_COLORS
        card = self._create_card(parent, fill=tk.X, pady=(0, 10))
        inner = tk.Frame(card, bg=colors["card"])
        inner.pack(fill=tk.X, padx=18, pady=12)

        tk.Label(
            inner,
            text="第三阶段 · Shopify 转换",
            bg=colors["card"],
            fg=colors["text"],
            font=(UI_FONT, 11, "bold"),
        ).pack(anchor=tk.W)
        tk.Label(
            inner,
            text="可使用第二阶段合并文件，也可单独填写或选择其他合并 XLSX",
            bg=colors["card"],
            fg=colors["text_muted"],
            font=(UI_FONT, 9),
        ).pack(anchor=tk.W, pady=(3, 8))

        file_row = tk.Frame(inner, bg=colors["card"])
        file_row.pack(fill=tk.X)
        self.conversion_entry = ttk.Entry(
            file_row,
            textvariable=self.conversion_file_var,
            style="App.TEntry",
        )
        self.conversion_entry.pack(
            side=tk.LEFT,
            fill=tk.X,
            expand=True,
            padx=(0, 8),
        )
        self.conversion_browse_btn = self._create_button(
            file_row,
            "浏览文件",
            self._browse_conversion_file,
            kind="secondary",
            padx=14,
            pady=8,
            font_size=9,
        )
        self.conversion_browse_btn.pack(side=tk.LEFT, padx=(0, 8))
        self.conversion_start_btn = self._create_button(
            file_row,
            "开始第三阶段转换",
            self._start_conversion,
            kind="primary",
            padx=16,
            pady=8,
            font_size=9,
        )
        self.conversion_start_btn.pack(side=tk.LEFT)

    def _build_action_bar(self, parent):
        colors = UI_COLORS
        actions = tk.Frame(parent, bg=colors["background"])
        actions.pack(fill=tk.X)
        tk.Label(
            actions,
            text="各阶段完成后暂停，可继续下一阶段或单独运行第三阶段",
            bg=colors["background"],
            fg=colors["text_muted"],
            font=(UI_FONT, 9),
        ).pack(side=tk.LEFT)

        buttons = tk.Frame(actions, bg=colors["background"])
        buttons.pack(side=tk.RIGHT)
        self.stop_btn = self._create_button(
            buttons, "停止", self._stop, kind="danger", padx=14, pady=9
        )
        self.stop_btn.pack(side=tk.RIGHT)
        self.continue_btn = self._create_button(
            buttons,
            "继续处理",
            self._continue_processing,
            kind="primary",
            padx=18,
            pady=9,
        )
        self.continue_btn.pack(side=tk.RIGHT, padx=(0, 8))
        self.existing_btn = self._create_button(
            buttons,
            "仅处理已有 XLSX",
            self._load_existing_files,
            kind="secondary",
            padx=14,
            pady=9,
        )
        self.existing_btn.pack(side=tk.RIGHT, padx=(0, 8))
        self.full_btn = self._create_button(
            buttons,
            "运行完整流程",
            self._start_full_pipeline,
            kind="primary",
            padx=16,
            pady=9,
        )
        self.full_btn.pack(side=tk.RIGHT, padx=(0, 8))


    def _browse_folder(self):
        path = filedialog.askdirectory(
            title="选择 01INPUT_XLSX 下的具体分类目录"
        )
        if path:
            self.folder_var.set(path)

    def _validated_task_folder(self) -> Path | None:
        raw_folder = self.folder_var.get().strip()
        if not raw_folder:
            messagebox.showerror("路径无效", "请先选择链接保存/读取目录。")
            return None
        folder = Path(raw_folder)
        try:
            validate_input_folder(folder)
        except ValueError as exc:
            messagebox.showerror("路径无效", str(exc))
            return None
        if not folder.is_dir():
            messagebox.showerror("路径无效", "所选分类目录不存在。")
            return None
        return folder

    def _scan_files(self):
        if self.pipeline_stage in {
            PipelineStage.COLLECTING_LINKS,
            PipelineStage.PROCESSING_PRODUCTS,
        }:
            return
        if self._validated_task_folder() is None:
            return
        super()._scan_files()
        if self.file_rows:
            self.pipeline_stage = PipelineStage.WAITING_CONTINUE
            self.update_status(
                0,
                100,
                "已读取链接文件；设置跳过图片位置后点击“继续处理”",
            )
            self.status_dot.configure(fg=UI_COLORS["warning"])
        else:
            self.pipeline_stage = PipelineStage.IDLE
        self._apply_stage_controls()


    def _load_existing_files(self):
        if self.pipeline_stage not in {
            PipelineStage.IDLE,
            PipelineStage.COMPLETED,
            PipelineStage.FAILED,
            PipelineStage.STOPPED,
        }:
            return
        self._scan_files()

    def _start_full_pipeline(self):
        if not self._can_start_task():
            return
        if self.pipeline_stage not in {
            PipelineStage.IDLE,
            PipelineStage.COMPLETED,
            PipelineStage.FAILED,
            PipelineStage.STOPPED,
        }:
            return
        pages = parse_category_pages(self.pages_text.get("1.0", tk.END))
        if not pages:
            messagebox.showerror("无法开始", "请先填写有效的分类列表。")
            return
        folder = self._validated_task_folder()
        if folder is None:
            return
        max_pages_raw = self.max_pages_var.get().strip()
        try:
            max_pages = int(max_pages_raw) if max_pages_raw else 0
        except ValueError:
            messagebox.showerror("页数无效", "最大页数必须是整数或留空。")
            return
        if max_pages < 0:
            messagebox.showerror("页数无效", "最大页数不能小于 0。")
            return

        self._launch_task(
            PipelineStage.COLLECTING_LINKS, self._run_link_collection,
            args=(pages, folder, max_pages, self.xpath_var.get().strip()),
            total=len(pages), text="第一阶段：正在启动链接采集...", clear_logs=True,
        )

    def _run_link_collection(
        self,
        pages: list[tuple[str, str]],
        output_folder: Path,
        max_pages: int,
        xpath: str,
    ):
        try:
            print(f"第一阶段：共 {len(pages)} 个分类")
            print(f"链接保存目录: {output_folder}")
            print(f"最大页数: {max_pages or '不限制'}")
            if xpath:
                print(f"XPath: {xpath}")

            all_results = self._collect_category_links(
                pages,
                max_pages=max_pages,
                xpath=xpath,
                output_folder=output_folder,
            )
            if all_results:
                print(f"第一阶段完成：{len(all_results)} 条链接已保存到 {output_folder}")
            else:
                print("第一阶段未采集到任何链接。")
            self.root.after(
                0,
                self._finish_link_collection,
                len(all_results),
            )
        except Exception as exc:
            print(f"链接采集发生异常: {exc}")

            self._queue_task_failure(self._fail_link_collection, exc)

    def _collect_category_links(
        self,
        pages: list[tuple[str, str]],
        max_pages: int,
        xpath: str,
        output_folder: Path,
    ) -> list[dict[str, str]]:
        all_results = load_existing_link_rows(output_folder)
        saved_links = {to_shopify_json_url(row["link"]) for row in all_results}
        domain_rows = {}
        for row in all_results:
            domain_rows.setdefault(urlparse(row["link"]).netloc, []).append(row)
        print(f"已读取已有 XLSX：{len(all_results)} 条唯一商品链接")
        chrome = create_browser()
        try:
            tab = chrome.latest_tab
            for category_index, (title, base_url) in enumerate(pages, 1):
                if self.stop_event.is_set():
                    break
                self.root.after(
                    0, self.update_status, category_index - 1, len(pages),
                    f"第一阶段 [{category_index}/{len(pages)}] {title}",
                )
                print(f"\n[{category_index}/{len(pages)}] {title}")
                seen_in_category = set()
                added_count = 0
                changed_domains = set()
                page_num = 1
                try:
                    while not self.stop_event.is_set():
                        if max_pages and page_num > max_pages:
                            break
                        parsed = urlparse(base_url)
                        query = parse_qs(parsed.query, keep_blank_values=True)
                        query["page"] = [str(page_num)]
                        url = base_url if page_num == 1 else urlunparse(
                            parsed._replace(query=urlencode(query, doseq=True), fragment="")
                        )
                        print(f"  第 {page_num} 页: {url}")
                        links = []
                        for retry in range(LINK_MAX_RETRIES + 1):
                            if self.stop_event.is_set():
                                break
                            if tab.get(url) is False:
                                raise TimeoutError(f"页面加载失败: {url}")
                            links = collect_page_links(tab, self.stop_event, xpath)
                            if links or page_num > 1 or retry == LINK_MAX_RETRIES:
                                break
                            print(f"  首页无链接，等待后重试 ({retry + 1}/{LINK_MAX_RETRIES})")
                            if self.stop_event.wait(RETRY_DELAY):
                                break
                        # 本轮分类内去重用于判断翻页终点；不能用历史数据提前结束翻页。
                        new_on_page = 0
                        for link in links:
                            key = to_shopify_json_url(link)
                            if key in seen_in_category:
                                continue
                            seen_in_category.add(key)
                            new_on_page += 1
                            if key not in saved_links:
                                saved_links.add(key)
                                row = {"title": title, "link": link}
                                all_results.append(row)
                                domain = urlparse(link).netloc
                                domain_rows.setdefault(domain, []).append(row)
                                changed_domains.add(domain)
                                added_count += 1
                        if not new_on_page:
                            print("  本页无本轮新增链接，结束该分类")
                            break
                        print(f"  分类已识别 {len(seen_in_category)} 条，去除已有数据后新增 {added_count} 条")
                        page_num += 1
                finally:
                    # 每个分类结束即保存，异常或停止也保留当前分类已获得的链接。
                    for domain in sorted(changed_domains):
                        self.pending_saves[f"links:{domain}"] = (
                            write_link_workbooks, (domain_rows[domain], output_folder),
                        )
                    for domain in sorted(changed_domains):
                        key = f"links:{domain}"
                        writer, args = self.pending_saves[key]
                        self._save_checkpoint(key, writer, *args)
                    print(f"  分类已保存：本分类新增 {added_count} 条，总计 {len(all_results)} 条")
            return all_results
        finally:
            print("第一阶段：正在关闭浏览器...")
            try:
                chrome.quit()
            except Exception:
                traceback.print_exc()


    def _finish_link_collection(self, link_count: int):
        stopped = self.stop_event.is_set()
        self.is_running = False
        self.pipeline_stage = next_stage_after_link_collection(
            link_count,
            stopped,
        )
        if self.pipeline_stage == PipelineStage.WAITING_CONTINUE:
            # 父类扫描方法负责生成可编辑的文件行。
            ScraperApp._scan_files(self)
            if self.file_rows:
                self.update_status(
                    0,
                    100,
                    "链接采集完成；设置跳过图片位置后点击“继续处理”",
                )
                self.status_dot.configure(
                    fg=UI_COLORS["warning"]
                )
            else:
                self.pipeline_stage = PipelineStage.IDLE
        elif stopped:
            self.update_status(0, 100, "链接采集已停止；不会自动进入下一阶段")
            self.status_dot.configure(fg=UI_COLORS["warning"])
        else:
            self.update_status(0, 100, "未采集到链接，请检查分类地址或 XPath")
            self.status_dot.configure(fg=UI_COLORS["danger"])
        self._apply_stage_controls()

    def _fail_link_collection(self, error: str):
        self.is_running = False
        self.task_error = error
        self.pipeline_stage = PipelineStage.FAILED
        self.update_status(0, 100, f"链接采集异常：{error}")
        self.status_dot.configure(fg=UI_COLORS["danger"])
        self._apply_stage_controls()

    def _continue_processing(self):
        if not can_continue_processing(self.pipeline_stage):
            return
        if not self.file_rows:
            messagebox.showerror("无法继续", "当前没有可处理的 XLSX 文件。")
            return
        if self.shutdown_requested or self.pending_saves:
            return
        self.last_merge_path = ""
        super()._start()

    def _on_merge_ready(self, merge_path):
        """第二阶段只提交合并结果，第三阶段由用户单独启动。"""
        self.last_merge_path = str(merge_path)
        self.root.after(0, self.conversion_file_var.set, str(merge_path))
        print("第二阶段合并完成，转换文件已送入第三阶段")

    def _browse_conversion_file(self):
        if self.pipeline_stage in {
            PipelineStage.COLLECTING_LINKS,
            PipelineStage.PROCESSING_PRODUCTS,
            PipelineStage.CONVERTING,
        }:
            return
        path = filedialog.askopenfilename(
            title="选择需要转换的合并 XLSX 文件",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")],
        )
        if path:
            self.conversion_file_var.set(path)

    def _validated_conversion_file(self):
        raw_path = self.conversion_file_var.get().strip()
        if not raw_path:
            messagebox.showerror("文件无效", "请填写或选择需要转换的 XLSX 文件。")
            return None
        input_file = Path(raw_path)
        if not input_file.is_file():
            messagebox.showerror("文件无效", "所选转换文件不存在。")
            return None
        if input_file.suffix.lower() != ".xlsx":
            messagebox.showerror("文件无效", "第三阶段只支持 .xlsx 文件。")
            return None
        return input_file

    def _start_conversion(self):
        if not self._can_start_task():
            return
        if self.pipeline_stage not in {
            PipelineStage.IDLE,
            PipelineStage.WAITING_CONVERSION,
            PipelineStage.COMPLETED,
            PipelineStage.FAILED,
            PipelineStage.STOPPED,
        }:
            return
        input_file = self._validated_conversion_file()
        if input_file is None:
            return

        self._launch_task(
            PipelineStage.CONVERTING, self._run_conversion, args=(input_file,),
            total=1, text=f"第三阶段：正在转换 {input_file.name}",
        )

    def _run_conversion(self, input_file):
        try:
            print(f"\n{'=' * 50}")
            print(f"第三阶段输入: {input_file}")
            result = _styles_to_shopify(str(input_file), print)
            print("第三阶段 Shopify 转换 + 价格匹配完成")
            self.root.after(0, self._finish_conversion, str(result or ""))
        except Exception as exc:
            print(f"第三阶段转换出错: {exc}")

            self._queue_task_failure(self._fail_conversion, exc)

    def _finish_conversion(self, result_path):
        self.is_running = False
        self.pipeline_stage = (PipelineStage.STOPPED if self.stop_event.is_set()
                               else PipelineStage.COMPLETED)
        detail = f" · {result_path}" if result_path else ""
        self.update_status(1, 1, f"第三阶段转换完成{detail}")
        self.status_dot.configure(fg=self._result_color())
        self._apply_stage_controls()

    def _fail_conversion(self, error):
        self.is_running = False
        self.task_error = error
        self.pipeline_stage = PipelineStage.FAILED
        self.update_status(0, 1, f"第三阶段转换失败：{error}")
        self.status_dot.configure(fg=UI_COLORS["danger"])
        self._apply_stage_controls()

    def _stop(self):
        if self.pipeline_stage == PipelineStage.COLLECTING_LINKS:
            if self.is_running:
                self.stop_event.set()
                self._set_button_enabled(self.stop_btn, False)
                self.update_status(
                    0,
                    100,
                    "正在停止链接采集，等待当前页面完成...",
                )
                print("\n[系统] 用户已停止完整流程，保存后关闭程序。")
                self._request_shutdown()
            return
        if self.pipeline_stage == PipelineStage.PROCESSING_PRODUCTS:
            super()._stop()


    def _apply_stage_controls(self):
        if not hasattr(self, "full_btn"):
            return
        idle = self.pipeline_stage in {
            PipelineStage.IDLE,
            PipelineStage.COMPLETED,
            PipelineStage.FAILED,
            PipelineStage.STOPPED,
        }
        waiting = self.pipeline_stage == PipelineStage.WAITING_CONTINUE
        conversion_ready = self.pipeline_stage in {
            PipelineStage.IDLE,
            PipelineStage.WAITING_CONVERSION,
            PipelineStage.COMPLETED,
            PipelineStage.FAILED,
            PipelineStage.STOPPED,
        }
        if self.shutdown_requested or self.pending_saves:
            idle = waiting = conversion_ready = False
        self._set_button_enabled(self.full_btn, idle)
        self._set_button_enabled(self.existing_btn, idle)
        self._set_button_enabled(self.continue_btn, waiting)
        self._set_button_enabled(
            self.stop_btn,
            not self.shutdown_requested and self.pipeline_stage in {
                PipelineStage.COLLECTING_LINKS,
                PipelineStage.PROCESSING_PRODUCTS,
            },
        )
        self._set_button_enabled(self.browse_btn, idle)
        self._set_button_enabled(self.scan_btn, idle)
        self._set_button_enabled(self.conversion_browse_btn, conversion_ready)
        self._set_button_enabled(self.conversion_start_btn, conversion_ready)

        normal_or_disabled = tk.NORMAL if idle else tk.DISABLED
        self.folder_entry.configure(state=normal_or_disabled)
        self.xpath_entry.configure(state=normal_or_disabled)
        self.max_pages_entry.configure(state=normal_or_disabled)
        self.pages_text.configure(state=normal_or_disabled)
        self.conversion_entry.configure(
            state=tk.NORMAL if conversion_ready else tk.DISABLED
        )
        for row in self.file_rows:
            row["skip_entry"].configure(
                state=tk.NORMAL if waiting else tk.DISABLED
            )

    def _reset_buttons(self):
        self.is_running = False
        if self.task_error or self.pending_saves:
            self.pipeline_stage = PipelineStage.FAILED
        elif self.pipeline_stage == PipelineStage.PROCESSING_PRODUCTS:
            if self.stop_event.is_set():
                self.pipeline_stage = PipelineStage.STOPPED
            elif self.last_merge_path:
                self.pipeline_stage = PipelineStage.WAITING_CONVERSION
                detail = (f"；有 {self.failed_product_count} 个商品失败"
                          if self.failed_product_count else "")
                self.update_status(
                    0, 1, f"第二阶段完成{detail}；请确认合并文件后开始第三阶段转换",
                )
            elif self.failed_product_count:
                self.pipeline_stage = PipelineStage.FAILED
            else:
                self.pipeline_stage = PipelineStage.COMPLETED
        self._apply_stage_controls()
        color = (UI_COLORS["danger"] if self.pipeline_stage == PipelineStage.FAILED
                 else self._result_color())
        self.status_dot.configure(fg=color)


def main():
    root = tk.Tk()
    ShopifyPipelineApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()