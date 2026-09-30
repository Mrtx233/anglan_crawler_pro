"""styles 数据流水线（单脚本，5 步顺序执行）：

  第一步  原价修正        D < C → D = C；styles1 内嵌价 p2 < p1 → p2 = p1
  第二步  解析 + 价格匹配  窗口 = 原价 × [LOWER_RATIO, UPPER_RATIO]，候选按解析顺序分配给变体
  第三步  导出未匹配价格
  第四步  转 Shopify      直接输出 UTF-8 CSV（逻辑复制自 03switch/01转shopify.py）
  第五步  转 WP           逻辑复制自 03switch/03shopify转WP.py

产出路径：
  <输入目录>/原价/lovedandco_styles_原价.xlsx        + lovedandco_styles_原价_Shopify.csv
  <输入目录>/价格匹配/lovedandco_styles_价格匹配.xlsx + lovedandco_styles_价格匹配_Shopify.csv
                                                     + lovedandco_styles_未匹配到的价格.csv
  WP 文件写在各子目录下，文件名为 wp-<Shopify文件名>
"""

import csv
import hashlib
import io
import random
import re
import string
from itertools import product
from pathlib import Path

import openpyxl
import pandas as pd

# ===== 配置：修改此处文件路径 =====
INPUT_FILE = r"D:\A_PythonCode\anglan_crawler_pro\DrissionPage_Json_Shopify\Data logic\styles\lovedandco_styles.xlsx"

# ===== 常量：价格库与匹配窗口 =====
price_library = [
    8.99, 9.01, 9.31, 9.33, 9.37, 9.63, 9.71, 9.72, 9.84, 9.86, 9.95, 9.99,
    10.03, 10.05, 10.11, 10.13, 10.15, 10.35, 10.36, 10.38, 10.54, 10.63, 10.64, 10.66,
    10.71, 10.78, 10.83, 10.94, 10.95, 10.99, 11.22, 11.36, 11.48, 11.65, 11.87, 11.95,
    11.99, 12.22, 12.25, 12.36, 12.41, 12.87, 12.92, 12.95, 12.99, 13.02, 13.13, 13.14,
    13.47, 13.69, 13.95, 13.99, 14.25, 14.32, 14.62, 14.83, 14.91, 14.95, 14.99, 15.31,
    15.52, 15.78, 15.95, 15.99, 16.34, 16.53, 16.66, 16.77, 16.95, 16.99, 17.11, 17.34,
    17.57, 17.95, 17.99, 18.62, 18.81, 18.84, 18.95, 18.99, 19.04, 19.25, 19.95, 19.99,
    20.14, 20.19, 20.34, 20.68, 20.74, 20.95, 20.99, 21.3, 21.47, 21.74, 21.95, 21.99,
    22.31, 22.34, 22.81, 22.95, 22.99, 23.14, 23.32, 23.95, 23.99, 24.22, 24.46, 24.63,
    24.95, 24.99, 25.37, 25.95, 25.99, 26.18, 26.34, 26.95, 26.99, 27.31, 27.54, 27.95,
    27.99, 28.11, 28.95, 28.99, 29, 29.31, 29.95, 29.99, 30.54, 30.95, 30.99, 31.26,
    31.52, 31.95, 31.99, 32.14, 32.24, 32.95, 32.99, 33.62, 33.84, 33.95, 33.99, 34.18,
    34.62, 34.95, 34.99, 35, 35.14, 35.33, 35.95, 35.99, 36.47, 36.95, 36.99, 37.21,
    37.95, 37.99, 38.14, 38.52, 38.95, 38.99, 39, 39.54, 39.95, 39.99, 40.02, 40.15,
    40.95, 40.99, 41.32, 41.95, 41.99, 42.51, 42.95, 42.99, 43.16, 43.95, 43.99, 44.82,
    44.95, 44.99, 45.17, 45.95, 45.99, 46.57, 47.82, 48.36, 49, 49.31, 50.03, 50.13,
    50.81, 55, 59, 69, 79, 85, 89, 99, 109, 119, 129, 189, 219, 259, 299, 319,
    349, 369, 399, 459, 499, 519, 599,
]
price_library_sorted = sorted(price_library)

LOWER_RATIO = 0.6
UPPER_RATIO = 1.1
MAX_LIB_PRICE = price_library_sorted[-1]  # 599


# ===== 通用工具 =====
def to_float(x):
    try:
        return float(x)
    except (ValueError, TypeError):
        return None


# ===== 第一步 / 第二步 辅助：styles1 解析与修正 =====
def parse_styles(raw):
    """解析一个 styles 单元格。

    结构: <Option1 Name>#<值1>&<Name2>&<值2>$price1$price2@图片#...
    返回 (Option1 Name, [变体dict, ...])
    """
    if pd.isna(raw) or str(raw).strip() == '':
        return None, []

    segments = str(raw).split('#')
    option1_name = segments[0]

    variants = []
    for seg in segments[1:]:
        opt_text, _, image = seg.partition('@')
        parts = opt_text.rsplit('$', 2)
        if len(parts) == 3:
            opts, price1, price2 = parts
        else:
            opts, price1, price2 = opt_text, '', ''

        p1 = to_float(price1)
        p2 = to_float(price2)
        # price2 小于 price1 时，让 price2 等于 price1
        if p1 is not None and p2 is not None and p2 < p1:
            p2 = p1

        variants.append({
            'option_values': opts.split('&'),
            'price1': p1,
            'price2': p2,
            'image': image,
            'raw': seg,
        })
    return option1_name, variants


def fix_styles(raw):
    """原价修正：styles1 内嵌价 price2 < price1 时置 price2 = price1。

    返回 (修正后的字符串, 修正的变体数)
    """
    if pd.isna(raw) or str(raw).strip() == '':
        return raw, 0

    segments = str(raw).split('#')
    fixed = [segments[0]]
    n = 0
    for seg in segments[1:]:
        opt_text, _, image = seg.partition('@')
        parts = opt_text.rsplit('$', 2)
        if len(parts) == 3:
            opts, p1s, p2s = parts
            p1 = to_float(p1s)
            p2 = to_float(p2s)
            if p1 is not None and p2 is not None and p2 < p1:
                fixed.append(f"{opts}${p1}${p1}@{image}")
                n += 1
                continue
        fixed.append(seg)
    return '#'.join(fixed), n


# ===== 路径准备 =====
input_path = Path(INPUT_FILE)

raw_dir = input_path.parent / "原价"
match_dir = input_path.parent / "价格匹配"
raw_dir.mkdir(exist_ok=True)
match_dir.mkdir(exist_ok=True)

raw_path = raw_dir / f"{input_path.stem}_原价{input_path.suffix}"
output_path = match_dir / f"{input_path.stem}_价格匹配{input_path.suffix}"

df = pd.read_excel(input_path)
df = df.reset_index(drop=True)

# ===== 第一步：原价修正（D < C → D = C；styles1 内嵌价 p2 < p1 → p2 = p1）=====
print('#' * 110)
print('第一步：原价修正（D < C → D = C；内嵌价 p2 < p1 → p2 = p1）')
print('#' * 110)

cd_fixed = 0
seg_fixed = 0
for row_idx, row in df.iterrows():
    c = to_float(row.get('price1'))
    d = to_float(row.get('price2'))

    if c is not None and d is not None and d < c:
        df.at[row_idx, 'price2'] = c
        cd_fixed += 1

    fixed_raw, n = fix_styles(row.get('styles1'))
    if n:
        df.at[row_idx, 'styles1'] = fixed_raw
        seg_fixed += n

print(f"\nC/D 修正行数: {cd_fixed}    styles1 内嵌修正变体数: {seg_fixed}")
df.to_excel(raw_path, index=False)
print(f"已保存: {raw_path}")

# ===== 第二步：解析 + 价格匹配 + 写回 =====
print('\n' + '#' * 110)
print('第二步：解析 + 价格匹配 + 写回（窗口 = 原价 × [0.6, 1.1]；候选按解析顺序从小到大分给变体）')
print('#' * 110)

used_prices = set()   # 全局统计：被任何商品用到的价格

for row_idx, row in df.iterrows():
    raw = row.get('styles1')
    option1_name, variants = parse_styles(raw)

    if option1_name is None or not variants:
        continue

    v_count = len(variants)

    # 原价：D列 price2，无效时用 C列 price1 兜底
    original = to_float(row.get('price2'))
    if original is None or original <= 0:
        original = to_float(row.get('price1'))

    # 计算每个变体的匹配价格
    if original is None or original <= 0:
        assigned = [None] * v_count
    elif original > MAX_LIB_PRICE:
        assigned = [MAX_LIB_PRICE] * v_count
    else:
        low = original * LOWER_RATIO
        high = original * UPPER_RATIO
        candidates = [p for p in price_library_sorted if low <= p <= high]
        n = len(candidates)
        if n == 0:
            assigned = [None] * v_count
        elif n >= v_count:
            assigned = candidates[:v_count]   # 候选多了，保留最小的 V 个
        else:
            assigned = [candidates[i * n // v_count] for i in range(v_count)]  # 不够按比例重复

    used = [p for p in assigned if p is not None]
    used_prices.update(used)
    if not used:
        continue

    matched_low = min(used)          # C列：最低价
    matched_original = max(used)     # D列：匹配原价

    # 写回 C列/D列
    df.at[row_idx, 'price1'] = matched_low
    df.at[row_idx, 'price2'] = matched_original

    # 重建 styles1：每个变体内嵌价 → $匹配价格$匹配原价
    segments = [option1_name]
    for v, matched in zip(variants, assigned):
        if matched is None:
            segments.append(v['raw'])          # 未匹配则保留原片段
        else:
            opts = '&'.join(v['option_values'])
            segments.append(f"{opts}${matched}${matched_original}@{v['image']}")
    df.at[row_idx, 'styles1'] = '#'.join(segments)

df.to_excel(output_path, index=False)
print(f"\n已保存: {output_path}")

# ===== 第三步：输出未匹配到的价格 =====
unused_prices = [p for p in price_library_sorted if p not in used_prices]
unused_path = match_dir / f"{input_path.stem}_未匹配到的价格.csv"
pd.DataFrame({'未匹配价格': unused_prices}).to_csv(unused_path, index=False, encoding='utf-8-sig')

print('\n' + '#' * 110)
print(f"第三步：未匹配到的价格（价格库 {len(price_library_sorted)} 个，用到 {len(used_prices)} 个，未用 {len(unused_prices)} 个）")
print('#' * 110)
print(f"已保存: {unused_path}")


# ==================================================================================
# 第四步：转 Shopify（逻辑复制自 03switch/01转shopify.py）
# ==================================================================================
def generate_unique_handle(existing_handles):
    """生成不在 existing_handles 中的唯一句柄。"""
    while True:
        digits = ''.join(random.choices(string.digits, k=3))
        letters1 = ''.join(random.choices(string.ascii_lowercase, k=2))
        letters2 = ''.join(random.choices(string.ascii_lowercase, k=2))
        handle = f"-{letters1}{digits}{letters2}"
        if handle not in existing_handles:
            existing_handles.add(handle)
            return handle


def generate_variant_sku():
    while True:
        digits = ''.join(random.choices(string.digits, k=1))
        if digits != '0':
            break
    digits1 = ''.join(random.choices(string.digits, k=2))
    letters1 = ''.join(random.choices(string.ascii_uppercase, k=2))
    letter1 = ''.join(random.choices(string.ascii_uppercase, k=1))
    digits2 = ''.join(random.choices(string.digits, k=2))
    letters2 = ''.join(random.choices(string.ascii_uppercase, k=2))
    digits3 = ''.join(random.choices(string.digits, k=3))
    return f"{digits}{digits1}{letters1}-{letter1}{digits2}{letters2}-{digits3}{letters1}"


def create_output_df(input_df):
    """创建输出数据框并填充数据。"""
    columns = [
        "Link-Href", "Handle", "Title", "Body (HTML)", "Collection",
        "Vendor", "Type", "Tags", "Published", "Option1 Name",
        "Option1 Value", "Option2 Name", "Option2 Value", "Option3 Name",
        "Option3 Value", "Variant SKU", "Variant Grams",
        "Variant Inventory Tracker", "Variant Inventory Qty",
        "Variant Inventory Policy", "Variant Fulfillment Service",
        "Variant Price", "Variant Compare At Price",
        "Variant Requires Shipping", "Variant Taxable",
        "Variant Barcode", "Image Src", "Image Position",
        "Image Alt Text", "Gift Card", "SEO Title",
        "SEO Description", "Google Shopping - Google Product Category",
        "Google Shopping - Gender", "Google Shopping - Age Group",
        "Google Shopping - MPN", "Google Shopping - AdWords Grouping",
        "Google Shopping - AdWords Labels", "Google Shopping - Condition",
        "Google Shopping - Custom Product", "Google Shopping - Custom Label 0",
        "Google Shopping - Custom Label 1", "Google Shopping - Custom Label 2",
        "Google Shopping - Custom Label 3", "Google Shopping - Custom Label 4",
        "Variant Image", "Variant Weight Unit", "Variant Tax Code",
        "Cost per item",
    ]

    output_df = pd.DataFrame(columns=columns)
    existing_handles = set()

    output_df['Link-Href'] = input_df.get('link-href', '')

    def clean_name(name):
        if pd.notna(name):
            cleaned_name = re.sub(r'[^a-zA-Z\s]', '', name).lower().replace(' ', '-').strip('-')
            return cleaned_name
        return ""

    base_handles = input_df['name'].apply(clean_name)
    handles = input_df['name'].apply(lambda name: generate_unique_handle(existing_handles) if pd.notna(name) else "")
    output_df['Handle'] = [f"{base}{handle}" for base, handle in zip(base_handles, handles)]

    output_df['Title'] = input_df['name']
    output_df['Body (HTML)'] = input_df['details']
    output_df['Collection'] = input_df['title']
    output_df['Vendor'] = input_df['title']
    output_df['Variant Compare At Price'] = input_df['price2']
    output_df['Variant Price'] = input_df['price1']
    output_df['Published'] = 'TRUE'
    output_df['Image Src'] = input_df['src_links'].str.split('#')
    output_df['Variant Weight Unit'] = 'kg'

    def safe_split(series):
        return series.fillna('').str.split('#')

    def handle_option(series):
        return safe_split(series).str[0].fillna('')

    output_df['Option1 Name'] = handle_option(input_df['styles1'])
    output_df['Option1 Value'] = safe_split(input_df['styles1']).str[1:].apply(
        lambda x: x if isinstance(x, list) else []).apply(lambda x: x if len(x) > 0 else [''])

    output_df['Option2 Name'] = handle_option(input_df['styles2'])
    output_df['Option2 Value'] = safe_split(input_df['styles2']).str[1:].apply(
        lambda x: x if isinstance(x, list) else []).apply(lambda x: x if len(x) > 0 else [''])

    output_df['Option3 Name'] = handle_option(input_df['styles3'])
    output_df['Option3 Value'] = safe_split(input_df['styles3']).str[1:].apply(
        lambda x: x if isinstance(x, list) else []).apply(lambda x: x if len(x) > 0 else [''])

    return output_df


def convert_to_shopify(input_file_name, output_file_name):
    with pd.ExcelFile(input_file_name) as xls:
        input_df = pd.read_excel(xls)

    required_columns = ['title', 'name', 'price1', 'price2', 'details', 'src_links', 'styles1', 'styles2', 'styles3']
    if not all(col in input_df.columns for col in required_columns):
        print(f"输入文件缺少必要的列，跳过: {input_file_name}")
        return

    output_df = create_output_df(input_df)

    all_rows = []
    for _, row in output_df.iterrows():
        option1_values = row['Option1 Value'] or ['']
        option2_values = row['Option2 Value'] or ['']
        option3_values = row['Option3 Value'] or ['']
        image_srcs = row['Image Src'] if isinstance(row['Image Src'], list) else [row['Image Src']]

        combinations = list(product(option1_values, option2_values, option3_values))
        combination_count = len(combinations)
        image_count = len(image_srcs)

        total_rows = max(combination_count, image_count)

        for i in range(total_rows):
            new_row = row.copy()

            if i < combination_count:
                combo = combinations[i]
                new_row['Option1 Value'] = combo[0]
                new_row['Option2 Value'] = combo[1]
                new_row['Option3 Value'] = combo[2]
            else:
                new_row['Option1 Value'] = ''
                new_row['Option2 Value'] = ''
                new_row['Option3 Value'] = ''

            if i == 0:
                new_row.update({
                    'Title': row['Title'],
                    'Body (HTML)': row['Body (HTML)'],
                    'Collection': row['Collection'],
                    'Vendor': row['Vendor'],
                    'Option1 Name': row['Option1 Name'],
                    'Option2 Name': row['Option2 Name'],
                    'Option3 Name': row['Option3 Name'],
                    'Published': row['Published'],
                    'Variant SKU': generate_variant_sku(),
                    'Variant Grams': 0,
                    'Variant Inventory Tracker': 'shopify',
                    'Variant Inventory Qty': 999,
                    'Variant Inventory Policy': 'deny',
                    'Variant Fulfillment Service': 'manual',
                    'Variant Requires Shipping': 'TRUE',
                    'Variant Taxable': 'FALSE',
                    'Gift Card': 'FALSE',
                    'SEO Title': row['Title'],
                    'SEO Description': row['Title'],
                })
            else:
                new_row.update({
                    'Title': '',
                    'Body (HTML)': '',
                    'Collection': '',
                    'Vendor': '',
                    'Published': '',
                    'Option1 Name': '',
                    'Option2 Name': '',
                    'Option3 Name': '',
                    'Variant SKU': generate_variant_sku(),
                    'Variant Grams': 0,
                    'Variant Inventory Tracker': 'shopify',
                    'Variant Inventory Qty': 999,
                    'Variant Inventory Policy': 'deny',
                    'Variant Fulfillment Service': 'manual',
                    'Variant Requires Shipping': 'TRUE',
                    'Variant Taxable': 'FALSE',
                })

            if i < image_count:
                new_row['Image Src'] = image_srcs[i]
                new_row['Image Position'] = i + 1
            else:
                new_row['Image Src'] = None
                new_row['Image Position'] = None

            if i >= combination_count:
                if 'Variant SKU' in new_row:
                    del new_row['Variant SKU']
                    del new_row['Variant Grams']
                    del new_row['Variant Inventory Tracker']
                    del new_row['Variant Inventory Qty']
                    del new_row['Variant Inventory Policy']
                    del new_row['Variant Fulfillment Service']
                    del new_row['Variant Price']
                    del new_row['Variant Compare At Price']
                    del new_row['Variant Requires Shipping']
                    del new_row['Variant Taxable']
            all_rows.append(new_row)

    final_df = pd.DataFrame(all_rows)

    option1_value_split = final_df['Option1 Value'].apply(
        lambda x: x.split('@', 1) if isinstance(x, str) and '@' in x else [x, None])
    final_df['Option1 Value'] = option1_value_split.apply(lambda x: x[0])
    final_df['Variant Image'] = option1_value_split.apply(lambda x: x[1])

    buf = io.BytesIO()
    final_df.to_excel(buf, index=False)
    buf.seek(0)

    # ============ 拆分 Option1 Value 列 ============
    SOURCE_COL = 11
    TOTAL_COLS = 7
    SPLIT_HEADERS = [
        "Option1 Value", "Option2 Name", "Option2 Value",
        "Option3 Name", "Option3 Value", "Variant Price", "Variant Compare At Price",
    ]

    def split_prices(text):
        parts = text.rsplit("$", 2)
        if len(parts) == 3:
            return parts[0], parts[1], parts[2]
        return text, "", ""

    def split_option_value(value):
        if value is None:
            return [""] * TOTAL_COLS
        text = str(value)
        option_text, price1, price2 = split_prices(text)
        option_parts = option_text.split("&") if option_text else []
        result = [""] * TOTAL_COLS
        for idx, part in enumerate(option_parts[:5]):
            result[idx] = part
        result[-2] = float(price1) if price1 else ""
        result[-1] = float(price2) if price2 else ""
        return result

    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    insert_at = SOURCE_COL + 1
    ws.insert_cols(insert_at, TOTAL_COLS - 1)
    for offset, header in enumerate(SPLIT_HEADERS):
        ws.cell(row=1, column=SOURCE_COL + offset, value=header)
    for row in range(2, ws.max_row + 1):
        source_value = ws.cell(row=row, column=SOURCE_COL).value
        values = split_option_value(source_value)
        for offset, val in enumerate(values):
            ws.cell(row=row, column=SOURCE_COL + offset, value=val)

    duplicate_option_start_col = SOURCE_COL + len(SPLIT_HEADERS)
    ws.delete_cols(duplicate_option_start_col, 4)

    split_variant_price_col = SOURCE_COL + 5
    split_compare_price_col = SOURCE_COL + 6
    target_variant_price_col = SOURCE_COL + len(SPLIT_HEADERS) + 6
    target_compare_price_col = target_variant_price_col + 1
    for row in range(2, ws.max_row + 1):
        variant_price = ws.cell(row=row, column=split_variant_price_col).value
        compare_price = ws.cell(row=row, column=split_compare_price_col).value
        if compare_price in (None, "", 0, 0.0, "0", "0.0"):
            compare_price = variant_price
        ws.cell(row=row, column=target_variant_price_col, value=variant_price)
        ws.cell(row=row, column=target_compare_price_col, value=compare_price)
    ws.delete_cols(split_variant_price_col, 2)

    COL_J = 10
    COL_L = 12
    COL_N = 14
    for row in range(2, ws.max_row + 1):
        if not ws.cell(row=row, column=COL_J).value:
            ws.cell(row=row, column=COL_L, value="")
            ws.cell(row=row, column=COL_N, value="")

    out_buf = io.BytesIO()
    wb.save(out_buf)
    out_buf.seek(0)
    pd.read_excel(out_buf).to_csv(output_file_name, index=False, encoding='utf-8-sig')


print('\n' + '#' * 110)
print('第四步：转 Shopify（直接输出 UTF-8 CSV）')
print('#' * 110)
for src in (raw_path, output_path):
    shopify_path = src.with_name(f"{src.stem}_Shopify.csv")
    convert_to_shopify(str(src), str(shopify_path))
    print(f"已保存: {shopify_path}")


# ==================================================================================
# 第五步：Shopify 转 WP（逻辑复制自 03switch/03shopify转WP.py）
# ==================================================================================
def WPTitle(row):
    title = ''
    if row['Option1 Value'] != "":
        title += "-" + row['Option1 Value']
        if row['Option2 Value'] != "":
            title += "," + row['Option2 Value']
            if row['Option3 Value'] != "":
                title += "," + row['Option3 Value']
    return title


def WPAdd(value, obj):
    if value != "" and value not in obj.split(','):
        obj += "," + value
    return obj


def build_variation(row, items):
    """构造 WP 子变体行。items 是该 Handle 的实体列表，items[0] 为父级。

    注意：Position 取 append 前的 len(items)，与原逻辑一致。
    """
    v = {}
    v["Type"] = "variation"
    v["SKU"] = row['Variant SKU']
    v["Name"] = items[0]["Name"] + WPTitle(row)
    v["Published"] = "1"
    v["Visibility in catalog"] = "visible"
    v["Description"] = ""
    v["In stock?"] = "1"
    v["Stock"] = "99999"
    v["Sale price"] = row['Variant Price']
    v["Regular price"] = (row['Variant Compare At Price']
                          if row['Variant Compare At Price'] != '' else row['Variant Price'])
    v["Categories"] = items[0]["Categories"]
    v["Images"] = row['Variant Image']
    v["Parent"] = items[0]["SKU"]
    v["Position"] = len(items)
    v["Attribute 1 name"] = items[0]["Attribute 1 name"]
    v["Attribute 1 value(s)"] = row['Option1 Value']
    v["Attribute 1 visible"] = ""
    v["Attribute 1 global"] = "1" if row['Option1 Value'] != "" else ""
    v["Attribute 2 name"] = items[0]["Attribute 2 name"]
    v["Attribute 2 value(s)"] = row['Option2 Value']
    v["Attribute 2 visible"] = ""
    v["Attribute 2 global"] = "1" if row['Option2 Value'] != "" else ""
    v["Attribute 3 name"] = items[0]["Attribute 3 name"]
    v["Attribute 3 value(s)"] = row['Option3 Value']
    v["Attribute 3 visible"] = ""
    v["Attribute 3 global"] = "1" if row['Option3 Value'] != "" else ""
    return v


def generate_parent_sku(handle="", counter=0, used_skus=None):
    """基于商品 Handle 生成确定性 parent SKU，同一 Handle 始终返回相同 SKU。"""
    while True:
        hash_input = f"{handle}_{counter}" if handle else str(counter)
        h = hashlib.md5(hash_input.encode()).hexdigest()
        digits = "".join(c for c in h if c.isdigit())
        if len(digits) < 8:
            digits = (digits + "0" * 8)[:8]
        sku = f"SKU{digits[:8]}"
        if used_skus is None or sku not in used_skus:
            if used_skus is not None:
                used_skus.add(sku)
            return sku
        counter += 1


def ShopifyWP(rows, folder_path, filename):
    list_items = {}
    used_parent_skus = set()

    for row in rows:
        entity = {}

        if row['Handle'] not in list_items:
            list_items[row['Handle']] = []

            if row['Option1 Value'] != "" and row['Option1 Value'] != "Default Title":
                entity["Type"] = "variable"
                entity["SKU"] = generate_parent_sku(row['Handle'], used_skus=used_parent_skus)
            else:
                entity["Type"] = "simple"
                entity["SKU"] = row['Variant SKU']

            entity["Name"] = row['Title']
            entity["Published"] = "1"
            entity["Visibility in catalog"] = "visible"
            entity["Description"] = row['Body (HTML)']
            entity["In stock?"] = "1"
            entity["Stock"] = "99999"

            entity["Sale price"] = "" if entity["Type"] == "variable" else row['Variant Price']
            entity["Regular price"] = "" if entity["Type"] == "variable" else (
                row['Variant Compare At Price'] if row['Variant Compare At Price'] != '' else row['Variant Price'])
            entity["Categories"] = row['Collection'] if row['Collection'] != '' else row['Type']
            entity["Images"] = row['Image Src']
            entity["Parent"] = ""
            entity["Position"] = "0"
            raw_opt1_name = row['Option1 Name']
            entity["Attribute 1 name"] = "" if (raw_opt1_name == 'Title' or raw_opt1_name.startswith('$')) else raw_opt1_name
            entity["Attribute 1 value(s)"] = "" if row['Option1 Value'] == 'Default Title' else row['Option1 Value']
            entity["Attribute 1 visible"] = "1" if entity["Attribute 1 name"] != "" else ""
            entity["Attribute 1 global"] = "1" if entity["Attribute 1 name"] != "" else ""
            entity["Attribute 2 name"] = row['Option2 Name']
            entity["Attribute 2 value(s)"] = row['Option2 Value']
            entity["Attribute 2 visible"] = "1" if row['Option2 Name'] != "" else ""
            entity["Attribute 2 global"] = "1" if row['Option2 Name'] != "" else ""
            entity["Attribute 3 name"] = row['Option3 Name']
            entity["Attribute 3 value(s)"] = row['Option3 Value']
            entity["Attribute 3 visible"] = "1" if row['Option3 Name'] != "" else ""
            entity["Attribute 3 global"] = "1" if row['Option3 Name'] != "" else ""

            list_items[row['Handle']].append(entity)

            if entity["Type"] == "variable":
                items = list_items[row['Handle']]
                items.append(build_variation(row, items))
        else:
            if row['Option1 Value'] != "" and row['Option1 Value'] != "Default Title":
                items = list_items[row['Handle']]
                items.append(build_variation(row, items))

                items[0]["Images"] = WPAdd(row['Image Src'], items[0]["Images"])
                items[0]["Attribute 1 value(s)"] = WPAdd(row['Option1 Value'], items[0]["Attribute 1 value(s)"])
                items[0]["Attribute 2 value(s)"] = WPAdd(row['Option2 Value'], items[0]["Attribute 2 value(s)"])
                items[0]["Attribute 3 value(s)"] = WPAdd(row['Option3 Value'], items[0]["Attribute 3 value(s)"])
            else:
                list_items[row['Handle']][0]["Images"] = WPAdd(row['Image Src'], list_items[row['Handle']][0]["Images"])

    if list_items:
        folder_path = Path(folder_path)
        csv_fname = folder_path / f"wp-{filename}"

        csv_head = ["Type", "SKU", "Name", "Published", "Visibility in catalog", "Description", "In stock?", "Stock",
                    "Sale price", "Regular price", "Categories", "Images", "Parent", "Position",
                    "Attribute 1 name", "Attribute 1 value(s)", "Attribute 1 visible", "Attribute 1 global",
                    "Attribute 2 name", "Attribute 2 value(s)", "Attribute 2 visible", "Attribute 2 global",
                    "Attribute 3 name", "Attribute 3 value(s)", "Attribute 3 visible", "Attribute 3 global"]

        with open(csv_fname, 'w', newline='', encoding='utf-8-sig') as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=csv_head)
            writer.writeheader()

            for handle, items in list_items.items():
                for item in items:
                    writer.writerow(item)

        print(f'WP文件 {csv_fname} 已生成！')


def convert_to_wp(csv_input):
    rows = []
    with Path(csv_input).open(newline='', encoding='utf-8-sig') as csvfile:
        reader = csv.DictReader(csvfile)
        for row in reader:
            rows.append(dict(row))
    ShopifyWP(rows, Path(csv_input).parent, Path(csv_input).name)


print('\n' + '#' * 110)
print('第五步：Shopify 转 WP')
print('#' * 110)
for src in (raw_path.with_name(f"{input_path.stem}_原价_Shopify.csv"),
            output_path.with_name(f"{input_path.stem}_价格匹配_Shopify.csv")):
    convert_to_wp(src)
