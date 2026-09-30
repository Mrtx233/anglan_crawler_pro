import csv
import glob
import hashlib
import os
import random
import shutil
import string
import sys
import tempfile
import threading
import time
from collections import Counter
from pathlib import Path

import tkinter as tk
from tkinter import BooleanVar, IntVar, Tk, StringVar, filedialog, messagebox
from tkinter import ttk
from tkinter.scrolledtext import ScrolledText

import pandas as pd
import openpyxl

# ── 图片转换依赖：Pillow（缺失时自动安装） ──
try:
    from PIL import Image
except ImportError:
    print("正在安装 Pillow...")
    os.system(f'"{sys.executable}" -m pip install Pillow -q')
    from PIL import Image

# ── 常量 ──
PRICE_LIBRARY = [
    8.99, 9.01, 9.31, 9.33, 9.37, 9.63, 9.71, 9.72, 9.84, 9.86, 9.95, 9.99,
    10.03, 10.05, 10.11, 10.13, 10.15, 10.35, 10.36, 10.38, 10.54, 10.63, 10.64, 10.66,
    10.71, 10.78, 10.83, 10.94, 10.95, 10.99, 11.22, 11.36, 11.48, 11.65, 11.87, 11.95,
    11.99, 12.22, 12.25, 12.36, 12.41, 12.87, 12.92, 12.95, 12.99, 13.02, 13.13, 13.14,
    13.47, 13.69, 13.95, 13.99, 14.25, 14.32, 14.62, 14.83, 14.91, 14.95, 14.99, 15.31,
    15.52, 15.78, 15.95, 15.99, 16.34, 16.53, 16.66, 16.77, 16.95, 16.99, 17.11, 17.34,
    17.57, 17.95, 17.99, 18.62, 18.81, 18.84, 18.95, 18.99, 19.04, 19.25, 19.95, 19.99,
    20.14, 20.19, 20.34, 20.68, 20.74, 20.95, 20.99, 21.3, 21.47, 21.74, 21.95, 21.99,
    22.31, 22.34, 22.81, 22.95, 22.99, 23.14, 23.32, 23.95, 23.99, 24.22, 24.46, 24.63,
    24.95, 24.99, 25.37, 25.95, 25.99, 26.18, 26.34, 26.95, 26.99, 27.31, 27.54, 27.95,
    27.99, 28.11, 28.95, 28.99, 29, 29.31, 29.95, 29.99, 30.54, 30.95, 30.99, 31.26,
    31.52, 31.95, 31.99, 32.14, 32.24, 32.95, 32.99, 33.62, 33.84, 33.95, 33.99, 34.18,
    34.62, 34.95, 34.99, 35, 35.14, 35.33, 35.95, 35.99, 36.47, 36.95, 36.99, 37.21,
    37.95, 37.99, 38.14, 38.52, 38.95, 38.99, 39, 39.54, 39.95, 39.99, 40.02, 40.15,
    40.95, 40.99, 41.32, 41.95, 41.99, 42.51, 42.95, 42.99, 43.16, 43.95, 43.99, 44.82,
    44.95, 44.99, 45.17, 45.95, 45.99, 46.57, 47.82, 48.36, 49, 49.31, 50.03, 50.13,
    50.81, 55, 59, 69, 79, 85, 89, 99, 109, 119, 129, 189, 219, 259, 299, 319,
    349, 369, 399, 459, 499, 519, 599,
]

STYLES_HEADERS = [
    "title", "name", "price1", "price2",
    "styles1", "styles2", "styles3",
    "src_links", "link-href", "details",
]

# ── Shopify 价格修正 ──
DEFAULT_SUFFIX = "_价格修正"
MARKUP = 1.2
SALE_COLUMN = "Sale price"
REGULAR_COLUMN = "Regular price"

# ── Excel 按 title 拆分 ──
DEFAULT_TITLE_COL = "title"
DEFAULT_PRICE_COL = "price1"
DEFAULT_PARTS = 2

# ── 图片批量转 JPG ──
HEIF_EXTS = {".heic", ".heif"}
IMAGE_CONVERT_EXTS = {".webp", ".bmp", ".tiff", ".tif", ".gif", ".avif"}

# ── styles 文本转义保留字符 ──
STYLE_RESERVED = r"\&#@"

# ── src_links 精简 ──
SRC_ONLY_KEEP = 4
SLIM_SUFFIX = "_图片精简"


# ══════════════════════════════════════════
# 通用工具
# ══════════════════════════════════════════

def clean_cell(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def strip_quotes(path):
    """去除路径两端可能包裹的引号（用户粘贴路径时可能带有引号）。"""
    path = str(path).strip()
    if len(path) >= 2 and (
            (path[0] == '"' and path[-1] == '"') or (path[0] == "'" and path[-1] == "'")
    ):
        path = path[1:-1].strip()
    return path


def detect_encoding(path):
    for enc in ("utf-8-sig", "gbk", "latin-1"):
        try:
            with open(path, encoding=enc, newline="") as f:
                f.read(2048)
            return enc
        except UnicodeDecodeError:
            continue
    return "utf-8-sig"


def read_table(path):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(path)
    return pd.read_csv(path, encoding=detect_encoding(path))


def output_with_suffix(input_path, suffix, extension=None):
    input_path = Path(input_path)
    ext = extension or input_path.suffix
    return input_path.with_name(f"{input_path.stem}{suffix}{ext}")


def parse_price(value):
    """把价格（字符串或数字）转成 float，无法解析返回 None。"""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "").replace("$", "").replace("￥", "").replace("¥", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def price_sort_key(value):
    """价格升序，无法解析的值排到末尾。"""
    price = parse_price(value)
    return (price is None, price if price is not None else 0.0)


def read_csv_rows(path):
    encoding = detect_encoding(path)
    with open(path, encoding=encoding, newline="") as f:
        return list(csv.DictReader(f)), encoding


# ══════════════════════════════════════════
# SKU 生成
# ══════════════════════════════════════════

def generate_variant_sku(used_skus=None):
    while True:
        first_digit = random.choice("123456789")
        digits1 = "".join(random.choices(string.digits, k=2))
        letters1 = "".join(random.choices(string.ascii_uppercase, k=2))
        letter1 = random.choice(string.ascii_uppercase)
        digits2 = "".join(random.choices(string.digits, k=2))
        letters2 = "".join(random.choices(string.ascii_uppercase, k=2))
        digits3 = "".join(random.choices(string.digits, k=3))
        sku = f"{first_digit}{digits1}{letters1}-{letter1}{digits2}{letters2}-{digits3}{letters1}"
        if used_skus is None or sku not in used_skus:
            if used_skus is not None:
                used_skus.add(sku)
            return sku


def generate_parent_sku(handle="", counter=0, used_skus=None):
    """基于商品 Handle 生成确定性 parent SKU，同一 Handle 始终返回相同 SKU。"""
    while True:
        hash_input = f"{handle}_{counter}" if handle else str(counter)
        h = hashlib.md5(hash_input.encode()).hexdigest()
        digits = "".join(c for c in h if c.isdigit())
        if len(digits) < 8:
            digits = (digits + "0" * 8)[:8]
        sku = f"SKU{digits[:8]}"
        if used_skus is None or sku not in used_skus:
            if used_skus is not None:
                used_skus.add(sku)
            return sku
        counter += 1


# ══════════════════════════════════════════
# styles 转义
# ══════════════════════════════════════════

def style_escape(value):
    """转义选项值/选项名里的保留分隔符（\\ & # @）；不含则原样返回。"""
    if value is None:
        return ""
    text = str(value)
    if not any(ch in text for ch in STYLE_RESERVED):
        return text
    return (text.replace("\\", "\\\\")
                .replace("&", "\\&")
                .replace("#", "\\#")
                .replace("@", "\\@"))


def split_unescaped(text, sep):
    """按未转义的 sep 切分（\\& \\# \\@ \\\\ 视为字面量，不参与切分）。"""
    parts = []
    buf = []
    escaped = False
    for ch in text:
        if escaped:
            buf.append(ch)
            escaped = False
        elif ch == "\\":
            buf.append(ch)
            escaped = True
        elif ch == sep:
            parts.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    parts.append("".join(buf))
    return parts


def last_unescaped_index(text, sep):
    """返回最后一个未转义 sep 的下标，没有则返回 -1。"""
    found = -1
    escaped = False
    for i, ch in enumerate(text):
        if escaped:
            escaped = False
        elif ch == "\\":
            escaped = True
        elif ch == sep:
            found = i
    return found


def format_price(value):
    """styles 用：数值去掉多余的 .0，非数值原样返回，空值返回 '0'。"""
    text = clean_cell(value)
    if not text:
        return "0"
    try:
        number = float(text)
        return str(int(number)) if number == int(number) else str(number)
    except ValueError:
        return text


# ══════════════════════════════════════════
# 功能 1：WP 价格检查
# ══════════════════════════════════════════

def price_check(input_file, log):
    input_path = Path(strip_quotes(input_file))
    df = read_table(input_path)

    sale_col = None
    regular_col = None
    for col in df.columns:
        col_lower = str(col).lower()
        if "sale" in col_lower and "price" in col_lower:
            sale_col = col
        if "regular" in col_lower and "price" in col_lower:
            regular_col = col

    if not sale_col or not regular_col:
        raise ValueError(f"未找到 Sale price 或 Regular price 列。当前列: {df.columns.tolist()}")

    extracted_prices = []
    for _, row in df.iterrows():
        sale_price = pd.to_numeric(row[sale_col], errors="coerce")
        regular_price = pd.to_numeric(row[regular_col], errors="coerce")
        if pd.notna(sale_price):
            extracted_prices.append(round(float(sale_price), 2))
        elif pd.notna(regular_price):
            extracted_prices.append(round(float(regular_price), 2))

    missing_prices = sorted([price for price in set(PRICE_LIBRARY) if price not in set(extracted_prices)])

    log(f"读取: {input_path}")
    log(f"使用价格列: {sale_col} / {regular_col}")
    for price in missing_prices:
        log(f"{price:.2f}")
    log("-" * 50)
    log(f"{len(missing_prices)} 个")
    return None


# ══════════════════════════════════════════
# 功能 2：WP 还原 styles
# ══════════════════════════════════════════

def parse_wp_products(input_path):
    rows, _ = read_csv_rows(input_path)
    products = []
    sku_to_product = {}
    for row in rows:
        row_type = row.get("Type", "").strip()
        if row_type in ("variable", "simple"):
            product_data = {"parent": row, "variations": []}
            products.append(product_data)
            sku_to_product[row.get("SKU", "")] = product_data
        elif row_type == "variation":
            product_data = sku_to_product.get(row.get("Parent", "").strip())
            if product_data:
                product_data["variations"].append(row)
    return products


def build_styles_segment(variation):
    parts = [style_escape(variation.get("Attribute 1 value(s)", "") or "")]
    for name_key, value_key in (
            ("Attribute 2 name", "Attribute 2 value(s)"),
            ("Attribute 3 name", "Attribute 3 value(s)"),
    ):
        name = variation.get(name_key, "") or ""
        value = variation.get(value_key, "") or ""
        if name and value:
            parts.extend([style_escape(name), style_escape(value)])
    segment = (f"{'&'.join(parts)}"
               f"${format_price(variation.get('Sale price', ''))}"
               f"${format_price(variation.get('Regular price', ''))}")
    image = (variation.get("Images", "") or "").split(",")[0].strip()
    return f"{segment}@{image}" if image else segment


def reconstruct_styles_product(product_data):
    parent = product_data["parent"]
    variations = product_data["variations"]
    product_row = {header: "" for header in STYLES_HEADERS}
    product_row["name"] = parent.get("Name", "")
    product_row["details"] = parent.get("Description", "")
    product_row["title"] = parent.get("Categories", "")
    product_row["link-href"] = parent.get("Link-Href", "")

    if variations:
        product_row["price1"] = variations[0].get("Sale price", "")
        product_row["price2"] = variations[0].get("Regular price", "")
    else:
        product_row["price1"] = parent.get("Sale price", "")
        product_row["price2"] = parent.get("Regular price", "")

    option1_name = parent.get("Attribute 1 name", "")
    segments = [
        build_styles_segment(variation)
        for variation in variations
        if variation.get("Attribute 1 value(s)", "") not in ("", "Default Title")
    ]
    if option1_name and segments:
        product_row["styles1"] = f"{style_escape(option1_name)}#{'#'.join(segments)}"
    elif segments:
        product_row["styles1"] = "#".join(segments)

    images = [img.strip() for img in (parent.get("Images", "") or "").split(",") if img.strip()]
    product_row["src_links"] = "#".join(images)
    return product_row


def wp_to_styles(input_file, log):
    input_path = Path(strip_quotes(input_file))
    products = parse_wp_products(input_path)

    stem = input_path.stem
    if stem.startswith("wp-"):
        stem = stem[3:]
    output_path = input_path.parent / f"{stem}_styles.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = stem[:31]
    ws.append(STYLES_HEADERS)

    for product_data in products:
        product = reconstruct_styles_product(product_data)
        ws.append([product.get(header, "") for header in STYLES_HEADERS])

    wb.save(output_path)
    log(f"读取: {input_path}")
    log(f"解析到 {len(products)} 个商品")
    log(f"输出: {output_path} ({ws.max_row - 1} 行)")
    return output_path


# ══════════════════════════════════════════
# 功能 3：SKU 查重
# ══════════════════════════════════════════

def sku_check(input_file, log):
    input_path = Path(input_file)
    rows, encoding = read_csv_rows(input_path)
    sku_rows = {}
    for line_no, row in enumerate(rows, start=2):
        sku = row.get("SKU", "").strip()
        if sku:
            sku_rows.setdefault(sku, []).append(
                (line_no, row.get("Type", ""), row.get("Name", ""), row.get("Parent", "")))

    duplicates = {sku: entries for sku, entries in sku_rows.items() if len(entries) > 1}
    parent_count = sum(1 for row in rows if row.get("Type", "") in ("variable", "simple"))
    variation_count = sum(1 for row in rows if row.get("Type", "") == "variation")
    parent_skus = {row.get("SKU", "").strip() for row in rows if row.get("Type", "") in ("variable", "simple")}
    orphans = [
        (line_no, row)
        for line_no, row in enumerate(rows, start=2)
        if row.get("Type", "") == "variation" and row.get("Parent", "").strip() not in parent_skus
    ]

    log(f"读取编码: {encoding}")
    log(f"总行数: {len(rows)}")
    log(f"variable/simple: {parent_count}，variation: {variation_count}")
    log(f"SKU 总数: {sum(len(v) for v in sku_rows.values())}，唯一 SKU 数: {len(sku_rows)}")
    log(f"重复 SKU 数: {len(duplicates)}，孤立变体数: {len(orphans)}")

    if duplicates:
        log("\n重复 SKU 明细:")
        for sku, entries in sorted(duplicates.items()):
            log(f"  SKU {sku} 出现 {len(entries)} 次")
            for line_no, row_type, name, parent in entries[:10]:
                parent_text = f", Parent={parent}" if parent else ""
                log(f"    行 {line_no}: Type={row_type}, Name={name[:60]}{parent_text}")
    else:
        log("没有发现重复 SKU。")

    if orphans:
        log("\n孤立变体明细:")
        for line_no, row in orphans[:20]:
            log(f"  行 {line_no}: SKU={row.get('SKU', '')}, Parent={row.get('Parent', '')}, Name={row.get('Name', '')[:60]}")
        if len(orphans) > 20:
            log(f"  还有 {len(orphans) - 20} 条未显示。")
    return None


# ══════════════════════════════════════════
# 功能 4：SKU 更新
# ══════════════════════════════════════════

def sku_update(input_file, log):
    input_path = Path(input_file)
    rows, encoding = read_csv_rows(input_path)
    if not rows:
        raise ValueError("文件中没有数据行。")
    fieldnames = list(rows[0].keys())
    old_to_new_parent = {}
    used_skus = set()
    parent_count = 0
    variation_count = 0

    for row in rows:
        row_type = row.get("Type", "").strip()
        if row_type in ("variable", "simple"):
            old_sku = row.get("SKU", "")
            handle = row.get("Name", "").strip()
            new_sku = generate_parent_sku(handle, used_skus=used_skus)
            row["SKU"] = new_sku
            old_to_new_parent[old_sku] = new_sku
            parent_count += 1
        elif row_type == "variation":
            row["SKU"] = generate_variant_sku(used_skus)
            old_parent = row.get("Parent", "").strip()
            if old_parent in old_to_new_parent:
                row["Parent"] = old_to_new_parent[old_parent]
            variation_count += 1

    output_path = output_with_suffix(input_path, "_新SKU", ".csv")
    with open(output_path, "w", encoding=encoding, newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    log(f"已生成新 SKU 文件: {output_path}")
    log(f"更新父商品 SKU: {parent_count}，变体 SKU: {variation_count}")
    log(f"唯一 SKU 数: {len(used_skus)}")
    return output_path


# ══════════════════════════════════════════
# 功能 5：Shopify 价格修正
# ══════════════════════════════════════════

def collect_csv_files(target):
    """目标可以是单个 csv 文件，也可以是一个文件夹（递归收集其中的 csv）。"""
    path = Path(strip_quotes(target))
    if path.is_dir():
        return sorted(p for p in path.rglob("*.csv") if p.is_file())
    if path.is_file():
        return [path]
    return []


def write_csv(output_path, header, rows, overwrite_from=None):
    """写出 utf-8-sig（带 BOM）的 CSV；覆盖原文件时先写临时文件再替换，避免中途失败损坏数据。"""
    output_path = Path(output_path)

    def _dump(target):
        with open(target, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(header)
            writer.writerows(rows)

    if overwrite_from is not None:
        fd, tmp_name = tempfile.mkstemp(dir=str(output_path.parent), suffix=".tmp")
        os.close(fd)
        try:
            _dump(tmp_name)
            os.replace(tmp_name, output_path)
        except PermissionError:
            raise PermissionError(f"文件被其他程序占用（如 Excel），请关闭后重试：{output_path}") from None
        finally:
            if os.path.exists(tmp_name):
                os.remove(tmp_name)
    else:
        try:
            _dump(output_path)
        except PermissionError:
            raise PermissionError(f"文件被其他程序占用（如 Excel），请关闭后重试：{output_path}") from None


def fix_equal_prices(csv_path, overwrite=False, suffix=DEFAULT_SUFFIX, log=print):
    csv_path = Path(csv_path)
    encoding = detect_encoding(csv_path)

    with open(csv_path, encoding=encoding, newline="") as f:
        reader = csv.reader(f)
        try:
            header = next(reader)
        except StopIteration:
            log(f"[跳过] 空文件：{csv_path.name}")
            return None
        rows = list(reader)

    if SALE_COLUMN not in header or REGULAR_COLUMN not in header:
        log(f"[跳过] 缺少 {SALE_COLUMN} / {REGULAR_COLUMN} 列：{csv_path.name}")
        return None

    sale_idx = header.index(SALE_COLUMN)
    regular_idx = header.index(REGULAR_COLUMN)

    changed = 0
    for row in rows:
        if len(row) <= max(sale_idx, regular_idx):
            continue
        sale = parse_price(row[sale_idx])
        regular = parse_price(row[regular_idx])
        if sale is None or regular is None:
            continue
        if abs(sale - regular) > 1e-9:
            continue
        row[regular_idx] = f"{regular * MARKUP:.2f}"
        changed += 1

    if changed == 0:
        log(f"[完成] {csv_path.name}：没有需要修正的行（共 {len(rows)} 行）")
        return {"file": csv_path, "total": len(rows), "changed": 0, "output": csv_path}

    output_path = csv_path if overwrite else csv_path.with_name(f"{csv_path.stem}{suffix}{csv_path.suffix}")
    write_csv(output_path, header, rows, overwrite_from=csv_path if overwrite else None)

    log(f"[完成] {csv_path.name}：共 {len(rows)} 行，修正 {changed} 行 -> {output_path.name}")
    return {"file": csv_path, "total": len(rows), "changed": changed, "output": output_path}


def fix_equal_prices_target(target, log):
    """适配 GUI 的包装：target 可以是单个 csv 或一个文件夹。默认另存为新文件。"""
    files = collect_csv_files(target)
    if not files:
        raise ValueError(f"未找到 CSV 文件: {target}")

    log(f"开始处理 {len(files)} 个文件（另存为 {DEFAULT_SUFFIX} 后缀）")
    log("-" * 60)

    total_rows = total_changed = failed = 0
    for path in files:
        try:
            stat = fix_equal_prices(path, overwrite=False, suffix=DEFAULT_SUFFIX, log=log)
        except Exception as exc:
            failed += 1
            log(f"[失败] {path.name}：{exc}")
            continue
        if stat:
            total_rows += stat["total"]
            total_changed += stat["changed"]

    log("-" * 60)
    summary = f"处理完成：{len(files) - failed} 个文件，共 {total_rows} 行，修正 {total_changed} 行"
    if failed:
        summary += f"，失败 {failed} 个"
    log(summary)

    return Path(strip_quotes(target))


# ══════════════════════════════════════════
# 功能 6：Excel 按 title 拆分
# ══════════════════════════════════════════

def collect_xlsx_files(target):
    """目标可以是单个 xlsx 文件，也可以是一个文件夹（递归收集其中的 xlsx）。"""
    path = Path(strip_quotes(target))
    if path.is_dir():
        return sorted(
            p for p in path.rglob("*.xlsx")
            if p.is_file() and not p.name.startswith("~$")
        )
    if path.is_file() and path.suffix.lower() == ".xlsx":
        return [path]
    return []


def save_workbook(wb, path):
    """保存工作簿；被其他程序占用时抛出可读的错误。"""
    try:
        wb.save(path)
    except PermissionError:
        raise PermissionError(f"文件被其他程序占用（如 Excel），请关闭后重试：{path}") from None


def save_split_workbook(path, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    save_workbook(wb, path)


def split_by_title(xlsx_path, parts, title_col=DEFAULT_TITLE_COL, price_col=DEFAULT_PRICE_COL,
                   overwrite=True, log=print):
    """按 title_col 分组、按 price_col 排序后轮询分配到 parts 个 xlsx，返回统计信息 dict。"""
    xlsx_path = Path(xlsx_path)
    parts = max(1, int(parts))

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = list(next(rows_iter))
    except StopIteration:
        wb.close()
        log(f"[跳过] 空文件：{xlsx_path.name}")
        return None
    rows = [list(r) for r in rows_iter]
    wb.close()

    if title_col not in header:
        log(f"[跳过] 缺少「{title_col}」列：{xlsx_path.name}")
        return None

    ti = header.index(title_col)
    has_price = price_col in header
    if not has_price:
        log(f"[提示] 缺少「{price_col}」列，组内保持原顺序：{xlsx_path.name}")
    pi = header.index(price_col) if has_price else None

    groups = {}
    for r in rows:
        groups.setdefault(r[ti], []).append(r)
    if pi is not None:
        for g in groups.values():
            g.sort(key=lambda r: price_sort_key(r[pi]))

    buckets = [[] for _ in range(parts)]
    for g in groups.values():
        for i, r in enumerate(g):
            buckets[i % parts].append(r)

    stem = xlsx_path.stem
    outdir = xlsx_path.parent
    outputs = []
    for k, bucket in enumerate(buckets, 1):
        out_path = outdir / f"{stem}_{k}.xlsx"
        if not overwrite and out_path.exists():
            out_path = outdir / f"{stem}_{k}_新.xlsx"
        save_split_workbook(out_path, header, bucket)
        outputs.append(out_path)
        log(f"[完成] {out_path.name} -> {len(bucket)} 行")

    log(f"[完成] {xlsx_path.name}：共 {len(rows)} 行 / {len(groups)} 组，拆分 {parts} 份")
    return {
        "file": xlsx_path,
        "total": len(rows),
        "groups": len(groups),
        "parts": parts,
        "outputs": outputs,
    }


# ══════════════════════════════════════════
# 功能 7：图片批量转 JPG
# ══════════════════════════════════════════

def register_heif(log):
    """尝试注册 pillow-heif，成功返回 True。"""
    try:
        from pillow_heif import register_heif_opener
        register_heif_opener()
        return True
    except ImportError:
        log("[警告] 未安装 pillow-heif，.heic/.heif 无法转换（其他格式不受影响）。")
    except Exception as e:
        log(f"[警告] 无法加载 pillow-heif（.heic/.heif 将跳过）：{e}")
    return False


def collect_image_files(images_dir, exts):
    """收集目录下所有指定后缀的图片，Windows 大小写不敏感去重。"""
    raw_files = []
    for ext in exts:
        raw_files.extend(glob.glob(os.path.join(images_dir, f"*{ext}")))
        raw_files.extend(glob.glob(os.path.join(images_dir, f"*{ext.upper()}")))
    unique = {}
    for f in raw_files:
        unique[os.path.normcase(f)] = f
    return list(unique.values())


def update_images_csv(csv_file, current_exts, log):
    """把 CSV 中 Images 列里的图片后缀替换成 .jpg，并备份原文件。"""
    try:
        with open(csv_file, "r", encoding="utf-8") as f:
            rows = list(csv.reader(f))

        if not rows:
            log("CSV 文件为空。")
            return

        header = rows[0]
        img_col = None
        for idx, col_name in enumerate(header):
            if col_name.strip().lower() == "images":
                img_col = idx
                break

        if img_col is None:
            log("[警告] CSV 中未找到 'Images' 列，跳过 CSV 更新。")
            return

        updated = 0
        for row in rows[1:]:
            if img_col < len(row):
                original = row[img_col]
                new_val = original
                for ext in current_exts:
                    new_val = new_val.replace(ext, ".jpg").replace(ext.upper(), ".jpg")
                if new_val != original:
                    row[img_col] = new_val
                    updated += 1

        backup_path = csv_file + ".bak"
        if os.path.exists(backup_path):
            os.remove(backup_path)
        os.rename(csv_file, backup_path)

        with open(csv_file, "w", encoding="utf-8", newline="") as f:
            csv.writer(f).writerows(rows)

        log(f"CSV 已更新 {updated} 行，原文件已备份为: {os.path.basename(backup_path)}")

    except Exception as e:
        log(f"[CSV 错误] 更新 CSV 文件时出错: {e}")


# ══════════════════════════════════════════
# 功能 8：图片链接精简
# ══════════════════════════════════════════

def extract_style_images(styles_value):
    """提取 styles1 各变体段 @ 后的图片链接，按出现顺序去重返回。"""
    text = clean_cell(styles_value)
    if not text:
        return []
    images = []
    seen = set()
    for segment in split_unescaped(text, "#"):
        pos = last_unescaped_index(segment, "@")
        if pos < 0:
            continue
        url = segment[pos + 1:].strip()
        if url and url not in seen:
            seen.add(url)
            images.append(url)
    return images


def split_src_links(src_value):
    """src_links 按 # 拆分（图片链接不参与转义），保持原顺序。"""
    return [url.strip() for url in clean_cell(src_value).split("#") if url.strip()]


def slim_src_links(styles_value, src_value, keep=SRC_ONLY_KEEP):
    """精简 src_links：styles1 未引用的保留前 keep 张，被引用的全部保留。

    返回 (新 src_links, 原图片数, 精简后图片数)。
    """
    referenced = set(extract_style_images(styles_value))
    src_list = split_src_links(src_value)
    src_only = [url for url in src_list if url not in referenced]
    shared = [url for url in src_list if url in referenced]
    kept = src_only[:keep] + shared
    return "#".join(kept), len(src_list), len(kept)


def slim_src_links_xlsx(xlsx_path, keep=SRC_ONLY_KEEP, overwrite=True, log=print):
    """按 styles1 精简 xlsx 的 src_links 列；覆盖原文件时先备份为 .bak。"""
    xlsx_path = Path(xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    header = [clean_cell(cell.value) for cell in ws[1]]
    for column in ("styles1", "src_links"):
        if column not in header:
            wb.close()
            log(f"[跳过] 缺少「{column}」列：{xlsx_path.name}")
            return None

    styles_idx = header.index("styles1") + 1
    src_idx = header.index("src_links") + 1

    rows = changed = before_total = after_total = 0
    for row_no in range(2, ws.max_row + 1):
        styles_value = ws.cell(row=row_no, column=styles_idx).value
        src_cell = ws.cell(row=row_no, column=src_idx)
        new_links, before, after = slim_src_links(styles_value, src_cell.value, keep)
        rows += 1
        before_total += before
        after_total += after
        if new_links != clean_cell(src_cell.value):
            src_cell.value = new_links
            changed += 1

    if overwrite:
        output_path = xlsx_path
        backup_path = Path(f"{xlsx_path}.bak")
        try:
            shutil.copy2(xlsx_path, backup_path)
        except PermissionError:
            wb.close()
            raise PermissionError(f"备份失败，文件被其他程序占用（如 Excel）：{xlsx_path}") from None
        log(f"已备份原文件: {backup_path.name}")
    else:
        output_path = xlsx_path.with_name(f"{xlsx_path.stem}{SLIM_SUFFIX}{xlsx_path.suffix}")

    save_workbook(wb, output_path)
    wb.close()

    log(f"[完成] {xlsx_path.name}：{rows} 行，改写 {changed} 行，"
        f"图片 {before_total} -> {after_total} 张 -> {output_path.name}")
    return {
        "file": xlsx_path,
        "rows": rows,
        "changed": changed,
        "before": before_total,
        "after": after_total,
        "output": output_path,
    }


def slim_src_links_target(target, keep=SRC_ONLY_KEEP, overwrite=True, log=print):
    """适配 GUI 的包装：target 可以是单个 xlsx 或一个文件夹（递归收集）。"""
    files = collect_xlsx_files(target)
    if not files:
        raise ValueError(f"未找到 xlsx 文件: {target}")

    mode_text = "覆盖原文件，自动备份 .bak" if overwrite else f"另存为 {SLIM_SUFFIX} 后缀"
    log(f"开始处理 {len(files)} 个文件（{mode_text}）")
    log("-" * 60)

    total_rows = total_changed = total_before = total_after = failed = 0
    for path in files:
        try:
            stat = slim_src_links_xlsx(path, keep=keep, overwrite=overwrite, log=log)
        except Exception as exc:
            failed += 1
            log(f"[失败] {path.name}：{exc}")
            continue
        if stat:
            total_rows += stat["rows"]
            total_changed += stat["changed"]
            total_before += stat["before"]
            total_after += stat["after"]

    log("-" * 60)
    summary = (f"处理完成：{len(files) - failed} 个文件，共 {total_rows} 行，"
               f"改写 {total_changed} 行，图片 {total_before} -> {total_after} 张")
    if failed:
        summary += f"，失败 {failed} 个"
    log(summary)

    return Path(strip_quotes(target))


# ══════════════════════════════════════════
# GUI
# ══════════════════════════════════════════

class SwitchToolGui:
    def __init__(self, root):
        self.root = root
        self.root.title("商品数据转换工具")

        window_width = 1180
        window_height = 780
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        x = (screen_width - window_width) // 2
        y = (screen_height - window_height) // 2
        self.root.geometry(f"{window_width}x{window_height}+{x}+{y}")
        self.root.minsize(1050, 680)

        self.run_buttons = []
        self.tab_buttons = {}
        self.frames = {}
        self._build_ui()

    def _build_ui(self):
        style = ttk.Style()
        if sys.platform.startswith("win"):
            style.theme_use("vista")

        default_font = ("Microsoft YaHei UI", 10)
        style.configure(".", font=default_font)
        style.configure("TButton", padding=6)
        style.configure("TLabelframe.Label", font=("Microsoft YaHei UI", 11, "bold"), foreground="#0078D4")

        # ── 底部状态栏 ──
        bottom = tk.Frame(self.root, bg="#E8F4F8")
        bottom.pack(side="bottom", fill="x")

        self.status = StringVar(value="✨ 就绪")
        tk.Label(bottom, textvariable=self.status, font=("Microsoft YaHei UI", 10),
                 bg="#E8F4F8", fg="#333333").pack(side="left", padx=20, pady=8)
        ttk.Button(bottom, text="🗑️ 清空日志", command=self.clear_log).pack(side="right", padx=20, pady=6)

        outer = ttk.Frame(self.root)
        outer.pack(fill="both", expand=True)

        header_frame = tk.Frame(outer, bg="#0078D4")
        header_frame.pack(fill="x")
        tk.Label(header_frame, text="✨ 商品数据转换工具",
                 font=("Microsoft YaHei UI", 16, "bold"), bg="#0078D4", fg="#FFFFFF").pack(side="left", padx=20, pady=15)

        paned = ttk.Panedwindow(outer, orient="horizontal")
        paned.pack(fill="both", expand=True, padx=20, pady=20)

        # ── 左：功能面板 ──
        left_panel = ttk.Frame(paned)
        paned.add(left_panel, weight=14)

        self.tab_frame = tk.Frame(left_panel)
        self.tab_frame.pack(fill="x", pady=(0, 20))
        for i in range(4):
            self.tab_frame.columnconfigure(i, weight=1)

        self.cards_frame = ttk.Frame(left_panel)
        self.cards_frame.pack(fill="both", expand=True)

        # ── 右：日志 ──
        right_panel = ttk.Frame(paned)
        paned.add(right_panel, weight=9)

        log_frame = ttk.LabelFrame(right_panel, text="  📝 运行日志  ", padding=15)
        log_frame.pack(fill="both", expand=True, padx=(10, 0))

        self.log_box = ScrolledText(
            log_frame, wrap="word",
            font=("Consolas", 11),
            bg="#FFFFFF", fg="#333333",
            insertbackground="#000000",
            selectbackground="#ADD6FF",
            highlightcolor="#0078D4", highlightthickness=1,
            padx=12, pady=12, relief="flat",
        )
        self.log_box.pack(fill="both", expand=True)

        # ── 通用工具面板 ──
        operations = [
            ("wp价格检查", "WP 价格检查", "选择 WP CSV文件",
             price_check, [("CSV 文件", "*.csv")], False),
            ("WP 还原 styles", "WP 还原 styles", "选择 wp-*.csv 文件",
             wp_to_styles, [("CSV 文件", "*.csv")], False),
            ("SKU 查重", "SKU 查重", "选择 wp-*.csv 文件",
             sku_check, [("CSV 文件", "*.csv")], False),
            ("SKU 更新", "SKU 更新", "选择 wp-*.csv 文件",
             sku_update, [("CSV 文件", "*.csv")], False),
            ("Shopify 价格修正", "Shopify 价格修正", "选择 Shopify CSV 文件或文件夹",
             fix_equal_prices_target, [("CSV 文件", "*.csv")], True),
        ]
        for idx, (internal_name, btn_name, prompt, func, filetypes, allow_dir) in enumerate(operations):
            self._build_tool_panel(idx, internal_name, btn_name, prompt, func, filetypes, allow_dir)

        # ── 第 6 个：Excel 按 title 拆分 ──
        self._build_split_panel(5)

        # ── 第 7 个：图片批量转 JPG ──
        self._build_image_panel(6)

        # ── 第 8 个：图片链接精简 ──
        self._build_src_slim_panel(7)

        self.select_tab("WP 价格检查")

    # ---------- 通用工具面板 ----------

    def _build_tool_panel(self, idx, internal_name, btn_name, prompt, func, filetypes, allow_dir=False):
        f = ttk.Frame(self.cards_frame)
        self.frames[internal_name] = f

        content = ttk.Frame(f)
        content.pack(fill="x", expand=False)

        ttk.Label(content, text=prompt, font=("Microsoft YaHei UI", 12, "bold"),
                  foreground="#0056b3").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 15))

        var = StringVar()
        ttk.Entry(content, textvariable=var, font=("Microsoft YaHei UI", 11)).grid(
            row=1, column=0, sticky="ew", ipady=5)

        if allow_dir:
            browse_cmd = lambda: self._browse_dir(var)
        else:
            browse_cmd = lambda: self.browse_file(var, filetypes)
        ttk.Button(content, text="📂 浏览", command=browse_cmd).grid(row=1, column=1, padx=(15, 8))

        run_button = ttk.Button(content, text="▶ 运行",
                                command=lambda: self.run_operation(internal_name, func, var.get()))
        run_button.grid(row=1, column=2)
        self.run_buttons.append(run_button)

        content.columnconfigure(0, weight=1)
        self._create_nav_button(idx, internal_name, btn_name)

    # ---------- 拆分面板 ----------

    def _build_split_panel(self, idx):
        internal_name = "Excel 按 title 拆分"
        f = ttk.Frame(self.cards_frame)
        self.frames[internal_name] = f

        self.split_path_var = StringVar()
        self.split_title_var = StringVar(value=DEFAULT_TITLE_COL)
        self.split_price_var = StringVar(value=DEFAULT_PRICE_COL)
        self.split_parts_var = IntVar(value=DEFAULT_PARTS)
        self.split_overwrite_var = BooleanVar(value=True)

        content = ttk.Frame(f)
        content.pack(fill="x", expand=False)
        content.columnconfigure(1, weight=1)

        ttk.Label(
            content,
            text=("按「分组列」把行归组，组内按「价格列」升序排序后，\n"
                  "轮询分配（第 1 行 → 第 1 个文件，第 2 行 → 第 2 个文件…），\n"
                  "使同一款式的不同价格尽量分散到各文件。\n"
                  "结果保存到源文件同目录，命名为「源文件名_1.xlsx」「源文件名_2.xlsx」…"),
            font=("Microsoft YaHei UI", 10),
            foreground="#555555",
            justify="left",
        ).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 15))

        ttk.Label(content, text="目标:", font=("Microsoft YaHei UI", 11, "bold")).grid(
            row=1, column=0, sticky="e", padx=(0, 10), pady=6)
        ttk.Entry(content, textvariable=self.split_path_var, font=("Microsoft YaHei UI", 11)).grid(
            row=1, column=1, sticky="ew", ipady=5, pady=6)
        ttk.Button(content, text="📂 浏览",
                   command=lambda: self._browse_dir(self.split_path_var)).grid(
            row=1, column=2, padx=(15, 0), pady=6)

        param_row = ttk.Frame(content)
        param_row.grid(row=2, column=0, columnspan=3, sticky="w", pady=(6, 0))

        ttk.Label(param_row, text="分组列:", font=("Microsoft YaHei UI", 11, "bold")).pack(side="left")
        ttk.Entry(param_row, textvariable=self.split_title_var, width=12).pack(side="left", padx=(4, 15))

        ttk.Label(param_row, text="价格列:", font=("Microsoft YaHei UI", 11, "bold")).pack(side="left")
        ttk.Entry(param_row, textvariable=self.split_price_var, width=12).pack(side="left", padx=(4, 15))

        ttk.Label(param_row, text="拆分份数:", font=("Microsoft YaHei UI", 11, "bold")).pack(side="left")
        ttk.Spinbox(param_row, from_=1, to=99, width=5, textvariable=self.split_parts_var).pack(side="left", padx=4)

        ttk.Checkbutton(content, text="覆盖同名文件（不勾选则加「_新」另存）",
                        variable=self.split_overwrite_var).grid(
            row=3, column=0, columnspan=3, sticky="w", pady=(8, 0))

        run_button = ttk.Button(content, text="🚀 开始拆分", command=self._run_split)
        run_button.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(15, 0))
        self.run_buttons.append(run_button)

        self._create_nav_button(idx, internal_name, "Excel 按 title 拆分")

    def _run_split(self):
        input_path = strip_quotes(self.split_path_var.get())
        if not input_path:
            messagebox.showwarning("缺少文件", "请先选择 Excel 文件或文件夹。")
            return
        if not Path(input_path).exists():
            messagebox.showerror("路径不存在", input_path)
            return

        try:
            parts = int(self.split_parts_var.get())
        except (tk.TclError, ValueError):
            messagebox.showwarning("提示", "拆分份数必须是数字。")
            return
        if parts < 1:
            messagebox.showwarning("提示", "拆分份数至少为 1。")
            return

        title_col = self.split_title_var.get().strip()
        if not title_col:
            messagebox.showwarning("提示", "请填写分组列名。")
            return
        price_col = self.split_price_var.get().strip()
        overwrite = self.split_overwrite_var.get()

        self.set_running(True)
        self._log_separator("Excel 按 title 拆分")
        self.log(f"输入: {input_path}")
        self.log(f"分组列: {title_col}    价格列: {price_col or '（空，保持原序）'}    份数: {parts}")

        def worker():
            try:
                files = collect_xlsx_files(input_path)
                if not files:
                    raise ValueError(f"未找到 xlsx 文件: {input_path}")
                self.log(f"待处理 {len(files)} 个文件")
                self.log("-" * 60)

                total_rows = total_files = failed = 0
                for path in files:
                    try:
                        stat = split_by_title(path, parts, title_col, price_col,
                                              overwrite=overwrite, log=self.log)
                    except Exception as exc:
                        failed += 1
                        self.log(f"[失败] {path.name}：{exc}")
                        continue
                    if stat:
                        total_files += 1
                        total_rows += stat["total"]

                self.log("-" * 60)
                summary = f"处理完成：{total_files} 个文件，共 {total_rows} 行，各拆成 {parts} 份"
                if failed:
                    summary += f"，失败 {failed} 个"
                self.log(summary)
            except Exception as exc:
                self.log(f"失败: {exc}")
                self.root.after(0, lambda: messagebox.showerror("运行失败", str(exc)))
            finally:
                self.root.after(0, lambda: self.set_running(False))

        threading.Thread(target=worker, daemon=True).start()

    # ---------- 图片转 JPG 面板 ----------

    def _build_image_panel(self, idx):
        internal_name = "图片批量转 JPG"
        f = ttk.Frame(self.cards_frame)
        self.frames[internal_name] = f

        self.img_dir_var = StringVar()
        self.img_csv_var = StringVar()

        content = ttk.Frame(f)
        content.pack(fill="x", expand=False)
        content.columnconfigure(1, weight=1)

        ttk.Label(
            content,
            text=("把指定文件夹内的 WebP / BMP / TIFF / GIF / AVIF / HEIC 等图片\n"
                  "批量转为 JPG（白色背景合成透明通道），原图删除。\n"
                  "可选：同时把 CSV 的「Images」列后缀替换为 .jpg（原 CSV 会备份为 .bak）。"),
            font=("Microsoft YaHei UI", 10),
            foreground="#555555",
            justify="left",
        ).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 15))

        ttk.Label(content, text="图片文件夹:", font=("Microsoft YaHei UI", 11, "bold")).grid(
            row=1, column=0, sticky="e", padx=(0, 10), pady=6)
        ttk.Entry(content, textvariable=self.img_dir_var, font=("Microsoft YaHei UI", 11)).grid(
            row=1, column=1, sticky="ew", ipady=5, pady=6)
        ttk.Button(content, text="📂 浏览", command=self._browse_image_dir).grid(
            row=1, column=2, padx=(15, 0), pady=6)

        ttk.Label(content, text="CSV 文件(可选):", font=("Microsoft YaHei UI", 11, "bold")).grid(
            row=2, column=0, sticky="e", padx=(0, 10), pady=6)
        ttk.Entry(content, textvariable=self.img_csv_var, font=("Microsoft YaHei UI", 11)).grid(
            row=2, column=1, sticky="ew", ipady=5, pady=6)
        ttk.Button(content, text="📂 浏览", command=self._browse_image_csv).grid(
            row=2, column=2, padx=(15, 0), pady=6)

        self.img_progress = tk.DoubleVar()
        ttk.Progressbar(content, variable=self.img_progress, maximum=100).grid(
            row=3, column=0, columnspan=3, sticky="ew", pady=(12, 0))

        run_button = ttk.Button(content, text="🚀 开始转换", command=self._run_image_conversion)
        run_button.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(15, 0))
        self.run_buttons.append(run_button)

        self._create_nav_button(idx, internal_name, "图片批量转 JPG")

    def _browse_image_dir(self):
        folder = filedialog.askdirectory(title="选择图片文件夹")
        if folder:
            self.img_dir_var.set(strip_quotes(folder))

    def _browse_image_csv(self):
        file = filedialog.askopenfilename(
            title="选择 CSV 文件（可选）",
            filetypes=[("CSV 文件", "*.csv"), ("所有文件", "*.*")],
        )
        if file:
            self.img_csv_var.set(strip_quotes(file))

    def _run_image_conversion(self):
        images_dir = strip_quotes(self.img_dir_var.get())
        csv_file = strip_quotes(self.img_csv_var.get())

        if not images_dir or not os.path.isdir(images_dir):
            messagebox.showwarning("提示", "请选择有效的图片文件夹！")
            return

        self.set_running(True)
        self._log_separator("图片批量转 JPG")
        self.img_progress.set(0)

        def worker():
            try:
                self._process_images(images_dir, csv_file)
            except Exception as exc:
                self.log(f"[严重错误] {exc}")
                self.root.after(0, lambda: messagebox.showerror("运行失败", str(exc)))
            finally:
                self.root.after(0, lambda: self.set_running(False))

        threading.Thread(target=worker, daemon=True).start()

    def _process_images(self, images_dir, csv_file):
        log = self.log
        log("=" * 30)
        log(f"图片文件夹: {images_dir}")
        if csv_file:
            log(f"CSV 文件:   {csv_file}")
        log("=" * 30)

        # 1. HEIF 支持检测
        current_exts = set(IMAGE_CONVERT_EXTS)
        if any(glob.glob(os.path.join(images_dir, f"*{ext}")) for ext in HEIF_EXTS):
            if register_heif(log):
                current_exts |= HEIF_EXTS
            else:
                log("  [跳过] 检测到 .heic/.heif 文件，但 pillow-heif 不可用，将不处理。")

        # 2. 收集 + 去重
        files = collect_image_files(images_dir, current_exts)
        total = len(files)
        if total == 0:
            log(f"未找到需要转换的图片（支持: {', '.join(sorted(current_exts))}）")
            return

        ext_count = Counter(os.path.splitext(f)[1].lower() for f in files)
        log(f"实际需要转换 {total} 张图片 (已自动去重):")
        for ext, count in sorted(ext_count.items()):
            log(f"  {ext}: {count} 张")
        log("开始转换...")

        success = fail = 0

        # 3. 逐张转换
        for i, src_path in enumerate(files, 1):
            jpg_path = os.path.splitext(src_path)[0] + ".jpg"
            try:
                img = Image.open(src_path)
                if img.mode in ("RGBA", "LA", "P"):
                    background = Image.new("RGB", img.size, (255, 255, 255))
                    if img.mode == "P":
                        img = img.convert("RGBA")
                    background.paste(img, mask=img.split()[-1] if "A" in img.mode else None)
                    img = background
                elif img.mode != "RGB":
                    img = img.convert("RGB")

                img.save(jpg_path, "JPEG", quality=90)
                img.close()
                os.remove(src_path)
                success += 1
            except Exception as e:
                log(f"  [失败] {os.path.basename(src_path)}: {e}")
                fail += 1

            pct = (i / total) * 100
            self.root.after(0, lambda p=pct: self.img_progress.set(p))

            if i % 10 == 0 or i == total:
                log(f"  进度: {i}/{total}  (成功 {success}, 失败 {fail})")

        log("")
        log(f"图片转换完成！成功 {success}，失败 {fail}")

        # 4. CSV 更新
        if not csv_file:
            log("未提供 CSV 文件，跳过更新。")
        elif not os.path.isfile(csv_file):
            log(f"[警告] CSV 文件不存在: {csv_file}")
        else:
            log("正在更新 CSV 文件...")
            update_images_csv(csv_file, current_exts, log)

        log("✅ 所有的任务已完成！")

    # ---------- 图片链接精简面板 ----------

    def _build_src_slim_panel(self, idx):
        internal_name = "图片链接精简"
        f = ttk.Frame(self.cards_frame)
        self.frames[internal_name] = f

        self.slim_path_var = StringVar()
        self.slim_keep_var = IntVar(value=SRC_ONLY_KEEP)
        self.slim_overwrite_var = BooleanVar(value=True)

        content = ttk.Frame(f)
        content.pack(fill="x", expand=False)
        content.columnconfigure(1, weight=1)

        ttk.Label(
            content,
            text=("按 styles1 精简 src_links 列：\n"
                  "  · 已被 styles1 引用的图片 → 全部保留\n"
                  "  · styles1 未引用的图片 → 只保留开头若干张，其余删除\n"
                  "精简结果写回源文件的 src_links 列。"),
            font=("Microsoft YaHei UI", 10),
            foreground="#555555",
            justify="left",
        ).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 15))

        ttk.Label(content, text="目标:", font=("Microsoft YaHei UI", 11, "bold")).grid(
            row=1, column=0, sticky="e", padx=(0, 10), pady=6)
        ttk.Entry(content, textvariable=self.slim_path_var, font=("Microsoft YaHei UI", 11)).grid(
            row=1, column=1, sticky="ew", ipady=5, pady=6)
        ttk.Button(content, text="📂 浏览",
                   command=lambda: self._browse_dir(self.slim_path_var)).grid(
            row=1, column=2, padx=(15, 0), pady=6)

        param_row = ttk.Frame(content)
        param_row.grid(row=2, column=0, columnspan=3, sticky="w", pady=(8, 0))

        ttk.Label(param_row, text="未引用图片保留张数:", font=("Microsoft YaHei UI", 11, "bold")).pack(side="left")
        ttk.Spinbox(param_row, from_=0, to=99, width=5, textvariable=self.slim_keep_var).pack(
            side="left", padx=(6, 0))

        ttk.Checkbutton(content, text="覆盖原文件（自动备份为 .bak；不勾选则另存为「_图片精简」）",
                        variable=self.slim_overwrite_var).grid(
            row=3, column=0, columnspan=3, sticky="w", pady=(10, 0))

        run_button = ttk.Button(content, text="🚀 开始精简", command=self._run_src_slim)
        run_button.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(15, 0))
        self.run_buttons.append(run_button)

        self._create_nav_button(idx, internal_name, "图片链接精简")

    def _run_src_slim(self):
        input_path = strip_quotes(self.slim_path_var.get())
        if not input_path:
            messagebox.showwarning("缺少文件", "请先选择 Excel 文件或文件夹。")
            return
        if not Path(input_path).exists():
            messagebox.showerror("路径不存在", input_path)
            return

        try:
            keep = int(self.slim_keep_var.get())
        except (tk.TclError, ValueError):
            messagebox.showwarning("提示", "保留张数必须是数字。")
            return
        if keep < 0:
            messagebox.showwarning("提示", "保留张数不能为负数。")
            return

        overwrite = self.slim_overwrite_var.get()

        self.set_running(True)
        self._log_separator("图片链接精简")
        self.log(f"输入: {input_path}")
        self.log(f"未引用图片保留张数: {keep}    覆盖原文件: {'是（备份 .bak）' if overwrite else '否'}")

        def worker():
            try:
                slim_src_links_target(input_path, keep=keep, overwrite=overwrite, log=self.log)
            except Exception as exc:
                self.log(f"失败: {exc}")
                self.root.after(0, lambda: messagebox.showerror("运行失败", str(exc)))
            finally:
                self.root.after(0, lambda: self.set_running(False))

        threading.Thread(target=worker, daemon=True).start()

    # ---------- 公共 UI ----------

    def _create_nav_button(self, idx, internal_name, display_text):
        """将导航按钮自动分为每行四个的网格"""
        r, c = divmod(idx, 4)
        btn = tk.Button(
            self.tab_frame, text=display_text, font=("Microsoft YaHei UI", 11, "bold"),
            relief="flat", cursor="hand2", bd=0, padx=5, pady=12,
            command=lambda n=internal_name: self.select_tab(n),
        )
        btn.grid(row=r, column=c, sticky="ew", padx=3, pady=3)
        self.tab_buttons[internal_name] = btn

    def _log_separator(self, name):
        self._append_log(f"[{time.strftime('%H:%M:%S')}][{name}]======================")

    def select_tab(self, name):
        """切换选中功能面板，更新按钮高亮状态"""
        for n, btn in self.tab_buttons.items():
            if n == name:
                btn.config(bg="#0078D4", fg="#FFFFFF", activebackground="#005A9E", activeforeground="#FFFFFF")
                self.frames[n].pack(fill="both", expand=True)
            else:
                btn.config(bg="#E1E4E8", fg="#333333", activebackground="#D0D4D9", activeforeground="#333333")
                self.frames[n].pack_forget()
        self._log_separator(name)

    def browse_file(self, var, filetypes):
        selected = filedialog.askopenfilename(
            title="选择输入文件",
            filetypes=filetypes + [("所有文件", "*.*")],
        )
        if selected:
            var.set(strip_quotes(selected))

    def _browse_dir(self, var):
        """目录选择：既可以选文件夹，也可以选单个文件。"""
        choice = messagebox.askyesnocancel(
            "选择目标",
            "是 → 选择文件夹（递归处理其中的 CSV / XLSX）\n否 → 选择单个文件",
        )
        if choice is None:
            return
        if choice:
            selected = filedialog.askdirectory(title="选择文件夹")
        else:
            selected = filedialog.askopenfilename(
                title="选择文件",
                filetypes=[("所有支持", "*.csv *.xlsx"), ("CSV 文件", "*.csv"),
                           ("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
            )
        if selected:
            var.set(strip_quotes(selected))

    def log(self, text=""):
        self.root.after(0, self._append_log, str(text))

    def _append_log(self, text):
        self.log_box.insert("end", f"{text}\n")
        self.log_box.see("end")

    def clear_log(self):
        self.log_box.delete("1.0", "end")

    def set_running(self, running):
        state = "disabled" if running else "normal"
        for button in self.run_buttons:
            button.configure(state=state)
        self.status.set("⏳ 运行中..." if running else "✨ 就绪")

    def run_operation(self, name, func, input_file):
        input_file = strip_quotes(input_file)
        if not input_file:
            messagebox.showwarning("缺少文件", "请先选择输入文件或文件夹。")
            return
        if not Path(input_file).exists():
            messagebox.showerror("路径不存在", input_file)
            return

        self.set_running(True)
        self._log_separator(name)
        self.log(f"输入: {input_file}")

        def worker():
            try:
                output = func(input_file, self.log)
                self.log(f"完成: {name}")
                if output:
                    self.log(f"输出: {output}")
            except Exception as exc:
                self.log(f"失败: {exc}")
                self.root.after(0, lambda: messagebox.showerror("运行失败", str(exc)))
            finally:
                self.root.after(0, lambda: self.set_running(False))

        threading.Thread(target=worker, daemon=True).start()


def main():
    root = Tk()
    SwitchToolGui(root)
    root.mainloop()


if __name__ == "__main__":
    main()