"""
Spider 公共辅助函数。

各 spider 文件通过 import 复用，避免重复代码。
"""

from pathlib import Path

import openpyxl


SPIDERS_DIR = Path(__file__).resolve().parent.parent / "spiders"
OUTPUT_BASE = Path(__file__).resolve().parent.parent / "output_styles"


def get_spider_category(spider_file_path):
    """检测 spider 文件所在的分类子文件夹名。

    spiders/threadheads.py         → ""
    spiders/xxx/folkclothing.py    → "xxx"
    """
    parent = Path(spider_file_path).resolve().parent
    if parent == SPIDERS_DIR:
        return ""
    return parent.name


def load_targets(filepath):
    """从 Excel 文件读取 (title, url) 列表。"""
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    targets = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        title = str(row[0]).strip() if row and row[0] else ""
        url = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if url:
            targets.append((title, url))
    wb.close()
    return targets


def load_cached_links(filepath):
    """从已有的 详细链接.xlsx 读取缓存的链接。"""
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    targets = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        title = str(row[0]).strip() if row and row[0] else ""
        url = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if url:
            targets.append((title, url))
    wb.close()
    return targets


def load_done_urls(spider_name, category=""):
    """扫描 output_styles 目录，从最新的 styles.xlsx 中读取已采集的商品 URL。

    category: 分类子文件夹名（如 "midsummera.com男女装日常休闲款"），
              为空时按 output_styles/{spider_name}/ 查找。
    """
    if category:
        out_base = OUTPUT_BASE / category / spider_name
    else:
        out_base = OUTPUT_BASE / spider_name
    if not out_base.is_dir():
        return set()
    prev_files = sorted(out_base.glob(f"*/{spider_name}_styles.xlsx"), reverse=True)
    if not prev_files:
        return set()
    latest = prev_files[0]
    done = set()
    try:
        wb = openpyxl.load_workbook(latest, read_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        href_col = headers.index("link-href") if "link-href" in headers else -1
        if href_col >= 0:
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[href_col]:
                    done.add(str(row[href_col]).strip())
        wb.close()
        print(f"[resume] 已采集 URL: {len(done)} 条 ← {latest}")
    except Exception as e:
        print(f"[resume] 加载旧数据失败: {e}")
    return done
