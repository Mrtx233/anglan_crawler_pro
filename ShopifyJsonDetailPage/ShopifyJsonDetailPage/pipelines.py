from datetime import datetime
from pathlib import Path

import openpyxl
from itemadapter import ItemAdapter
from scrapy import signals


# Windows GBK 控制台安全输出
def _safe_print(text):
    try:
        print(text)
    except UnicodeEncodeError:
        print(str(text).encode("gbk", errors="replace").decode("gbk"))


# ── styles 转义感知拆分 ──────────────────────────────────────
# styles 用 & # @ 作分隔符；选项值里的 & # @ 会被反斜杠转义（如 "Tie \& Square"），
# 这里只按未转义的 # / @ 切分，避免把转义字符误当成分隔符。
_STYLE_RESERVED = r"\&#@"


def _style_split(text, sep):
    parts = []
    buf = []
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        if ch == "\\" and i + 1 < n and text[i + 1] in _STYLE_RESERVED:
            buf.append(text[i:i + 2])
            i += 2
            continue
        if ch == sep:
            parts.append("".join(buf))
            buf = []
            i += 1
            continue
        buf.append(ch)
        i += 1
    parts.append("".join(buf))
    return parts


def _style_split_first(text, sep):
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        if ch == "\\" and i + 1 < n and text[i + 1] in _STYLE_RESERVED:
            i += 2
            continue
        if ch == sep:
            return text[:i], text[i + 1:], True
        i += 1
    return text, "", False


def _extract_variant_images(styles_str):
    """从 styles 编码中提取 @image URL，保持顺序、去重。"""
    if not styles_str:
        return []
    imgs, seen = [], set()
    for seg in _style_split(styles_str, "#"):
        if "@" in seg:
            _head, url, has_image = _style_split_first(seg, "@")
            if has_image:
                url = url.strip()
                if url and url not in seen:
                    seen.add(url)
                    imgs.append(url)
    return imgs


# ── 字段列表 ────────────────────────────────────────────────

_FIELDS = [
    "title", "name", "price1", "price2",
    "styles1", "styles2", "styles3",
    "src_links", "link-href", "details",
]


class ShopifyJsonPipeline:

    def open_spider(self, spider):
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")

        # 从 spider 获取分类文件夹（由 get_spider_category 检测）
        category = getattr(spider, "_category", "")
        if category:
            group = str(Path(category) / spider.name)
        else:
            group = spider.name

        out_base = Path(__file__).resolve().parent / "output_styles" / group
        out_dir = out_base / ts
        out_dir.mkdir(parents=True, exist_ok=True)

        self.out_dir = out_dir
        self.xlsx_path = out_dir / f"{spider.name}_styles.xlsx"
        self.rows: list[dict] = []
        self.cursor = 0
        self.failures = 0

        # ── 断点恢复：加载上次已有的数据 ──
        done_urls: set[str] = set()
        if out_base.is_dir():
            prev_files = sorted(out_base.glob(f"*/{spider.name}_styles.xlsx"), reverse=True)
        else:
            prev_files = []
        if prev_files:
            latest = prev_files[0]
            try:
                wb = openpyxl.load_workbook(latest, read_only=True)
                ws = wb.active
                headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
                href_col = headers.index("link-href") if "link-href" in headers else -1
                for row in ws.iter_rows(min_row=2, values_only=True):
                    row_dict = {
                        h: (v or "") for h, v in zip(headers, row)
                    }
                    self.rows.append(row_dict)
                    if href_col >= 0 and row[href_col]:
                        done_urls.add(str(row[href_col]).strip())
                wb.close()
                print(f"[resume] 加载已有数据: {len(self.rows)} 条 ← {latest}")
            except Exception as e:
                print(f"[resume] 加载旧数据失败: {e}")

        self._done_urls = done_urls
        # 暴露给 spider，让 spider 跳过已采集的 URL
        spider._done_urls = done_urls

        spider.crawler.signals.connect(self._on_dropped, signal=signals.item_dropped)

    def _on_dropped(self, item, response, exception, spider):
        self.cursor += 1
        self.failures += 1
        total = getattr(spider, "total_items", "?")
        _safe_print(f"[{self.cursor}/{total}] 获取失败: {exception}")

    def process_item(self, item, spider):
        adapter = ItemAdapter(item)
        row = {f: adapter.get(f, "") for f in _FIELDS}

        # 合并商品图: 前5张 spider 图 + styles 中的 @image 变体图，去重
        base_imgs = [u.strip() for u in row.get("src_links", "").split("#") if u.strip()][:5]
        seen = set(base_imgs)
        for key in ("styles1", "styles2", "styles3"):
            for img in _extract_variant_images(row.get(key, "")):
                if img not in seen:
                    seen.add(img)
                    base_imgs.append(img)
        row["src_links"] = "#".join(base_imgs)

        self.rows.append(row)
        self.cursor += 1
        total = getattr(spider, "total_items", "?")
        _safe_print(f"[{self.cursor}/{total}] 获取成功")
        return item

    def close_spider(self, spider):
        _safe_print(f"完成: {len(self.rows)} 成功, {self.failures} 失败")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = spider.name
        ws.append(_FIELDS)
        for r in self.rows:
            ws.append([r.get(f, "") for f in _FIELDS])
        wb.save(self.xlsx_path)
        print(f"XLSX saved: {self.xlsx_path} ({len(self.rows)} rows)")
