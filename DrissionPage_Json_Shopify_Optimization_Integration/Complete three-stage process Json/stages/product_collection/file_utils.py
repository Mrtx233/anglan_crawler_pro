# ======================== 文件读写层 ========================
#
# 职责: 链接 JSON 读取、商品记录 JSON 保存/加载、安全输出路径构造、
#       恢复文件命名。（旧 XLSX 读写保留但已不再使用）
# 依赖: config（RECORD_FIELDS / RECOVERY 命名）、link_collection.file_utils
#       （链接 JSON 解析）、openpyxl、pathlib。
# 禁止: 导入 tkinter、DrissionPage。
#
# 链接表来源说明:
#   第一阶段现在按域名输出 {域名}.json（分类名 → 商品链接数组），
#   本层不再有 Excel 链接表读取逻辑；解析复用 link_collection 的
#   同一份实现，保证 GUI 分类列表与采集端分组键逐字符一致。
#
# 输出约定（第二阶段）:
#   {域名}商品.json，外层信封含域名，products 为商品记录数组：
#     {"domain": "...", "products": [{category, name, link_href, ...}, ...]}
#   字段定义见 config.RECORD_FIELDS 与 product_record 的模块说明。
#
# 安全约定（本版强化）:
#   1. build_output_file_path 保证输出路径绝不等于输入路径，
#      相同时抛 RuntimeError 阻止覆盖输入文件。
#   2. 已有商品记录 JSON 读取失败时，调用方必须改用
#      build_recovery_json_path 另存，绝不覆盖原文件。
#   3. save_records_json 采用「同目录临时文件 + os.replace」原子写入，
#      失败清理临时文件并抛出，原文件不受影响。
# ============================================================

import json
import os
import tempfile
import time
from pathlib import Path

import openpyxl

from stages.link_collection.file_utils import read_targets_from_domain_json
from stages.product_collection import config

# 恢复文件名中的标记
RECOVERY_MARK = "recovery"

# 第一阶段的链接表后缀
LINK_FILE_SUFFIX = ".json"

# 第二阶段商品记录的文件后缀
RECORD_FILE_SUFFIX = ".json"


def read_targets(filepath):
    """读取链接文件，返回 [(分类名, 链接), ...]。

    支持两种来源：
      - 第一阶段的 {域名}.json：分类名 → 商品链接数组；
      - 手工准备的 CSV：第 1 列分类名、第 2 列链接，跳过表头。
    """
    path = Path(filepath)
    if path.suffix.lower() == LINK_FILE_SUFFIX:
        return read_targets_from_domain_json(str(path))
    return _read_targets_from_csv(path)


def _read_targets_from_csv(path):
    """读取两列 CSV 链接表；空行与链接为空的行直接丢弃。"""
    import csv

    targets = []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        next(reader, None)  # 跳过表头
        for row in reader:
            title = str(row[0]).strip() if row else ""
            url = str(row[1]).strip() if len(row) > 1 else ""
            if url:
                targets.append((title, url))
    return targets


def read_categories(filepath):
    """按首次出现顺序返回该文件里会被采集的分类名（去重）。

    刻意复用 read_targets 而不是另写一遍读表逻辑：采集端用
    group_targets_by_json_url 得到的分组键就是这里的 title，GUI 的分类级
    图片设置也以它为键。同一个来源保证两边永远对得上——只要读表逻辑分家，
    GUI 里设了却匹配不上的静默失效就可能出现。

    链接为空的行本来就不会被采集，因此不计入分类列表。
    """
    categories = []
    seen = set()
    for title, _url in read_targets(filepath):
        if title not in seen:
            seen.add(title)
            categories.append(title)
    return categories


def _dump_record_envelope(domain, records):
    """紧凑 JSON：每条记录一个缩进块，内部子结构保持可读。

    ensure_ascii=False 保留中文；记录之间换行，避免整文件压成一行后
    无法用肉眼核对采集结果。
    """
    body = ",\n".join(
        "    " + json.dumps(record, ensure_ascii=False, indent=2).replace("\n", "\n    ")
        for record in records
    )
    if not body:
        body = ""
    return (
        "{\n"
        f"  \"domain\": {json.dumps(str(domain or ''), ensure_ascii=False)},\n"
        "  \"products\": [\n" + body + ("\n" if body else "") + "  ]\n"
        "}"
    )


def save_records_json(records, output_path, domain=None):
    """原子写入商品记录 JSON：先写同目录临时文件，成功后 os.replace 覆盖目标。

    domain 省略时从输出文件名推断（{域名}商品.json 的域名段）。
    失败时清理临时文件并抛出异常；原文件不会被破坏。
    """
    output_path = Path(output_path)
    if domain is None:
        stem = output_path.stem
        domain = stem[: -len("商品")] if stem.endswith("商品") else stem
    target_dir = output_path.parent if str(output_path.parent) else Path(".")
    target_dir.mkdir(parents=True, exist_ok=True)

    fd, tmp_path = tempfile.mkstemp(
        prefix=f".{output_path.stem}_tmp_",
        suffix=RECORD_FILE_SUFFIX,
        dir=str(target_dir),
    )
    os.close(fd)
    try:
        with open(tmp_path, "w", encoding="utf-8") as f:
            f.write(_dump_record_envelope(domain, records))
        os.replace(tmp_path, str(output_path))
    except Exception:
        try:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
        except OSError:
            pass
        raise


def load_records_json(filepath):
    """读取已有商品记录 JSON，返回记录列表（用于断点续采）。

    刻意区分两种失败：
      - 文件不存在：调用方应先判断，本函数直接抛 FileNotFoundError；
      - 文件存在但内容损坏：抛 ValueError，调用方必须另存 recovery，
        绝不覆盖原文件。
    因此本函数**不**把损坏吞成空列表——否则已有结果会被静默清空。
    """
    with open(filepath, "r", encoding="utf-8") as f:
        loaded = json.load(f)
    if isinstance(loaded, dict):
        products = loaded.get("products")
    elif isinstance(loaded, list):
        products = loaded  # 兼容裸数组写法
    else:
        raise ValueError(f"商品记录格式错误（应为对象或数组）: {filepath}")
    if not isinstance(products, list):
        raise ValueError(f"商品记录缺少 products 数组: {filepath}")
    return [item for item in products if isinstance(item, dict)]


# ── 以下 XLSX 读写为 JSON 迁移前的旧路径，已不再被 scraper_runner 调用。
#    保留是为了让本次改动可单独回退，也便于第三阶段继续读取历史中间表。
def save_xlsx(rows, output_path):
    """原子写入 XLSX：先写同目录临时文件，成功后 os.replace 覆盖目标。

    签名与原实现保持一致，所有调用方无需改动。
    失败时清理临时文件并抛出异常；原文件不会被破坏。
    """
    output_path = Path(output_path)
    target_dir = output_path.parent if str(output_path.parent) else Path(".")
    target_dir.mkdir(parents=True, exist_ok=True)

    fd, tmp_path = tempfile.mkstemp(
        prefix=f".{output_path.stem}_tmp_",
        suffix=".xlsx",
        dir=str(target_dir),
    )
    os.close(fd)
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(config.CSV_FIELDS)
        for row in rows:
            ws.append([row.get(f, "") for f in config.CSV_FIELDS])
        wb.save(tmp_path)
        wb.close()
        os.replace(tmp_path, str(output_path))
    except Exception:
        try:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
        except OSError:
            pass
        raise


def load_xlsx(filepath):
    """读取已有输出表，返回 (表头列表, 行字典列表)。

    旧 XLSX 路径，已不再被 scraper_runner 调用。
    用于断点续采：只关心每行的 link-href；空单元格统一归一成空串
    （注意 0 也会被转成空串，因为 `v or ""` 对 0 同样成立）。
    """
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(rows_iter)]
    results = []
    for row in rows_iter:
        results.append({h: (v or "") for h, v in zip(headers, row)})
    wb.close()
    return headers, results


def _replace_path_component(
    path,
    source_name=config.INPUT_DIRNAME,
    target_name=config.OUTPUT_DIRNAME,
):
    """仅替换完整目录名，避免字符串 replace 误改其它路径片段。"""
    path_obj = Path(path)
    parts = list(path_obj.parts)
    for index, part in enumerate(parts):
        if part.lower() == source_name.lower():
            parts[index] = target_name
            return Path(*parts), True
    return path_obj, False


def build_output_file_path(input_file):
    """构造安全输出路径，保证永远不会与输入文件为同一路径。

    输入是第一阶段的 01INPUT_XLSX/.../{域名}.json，输出是 02OUTPUT_XLSX/
    下同批次同任务的 {域名}商品.json：目录段按原约定整段替换，文件名换成
    {域名}商品.json——既保证与输入永不同名，也让产物一眼可辨是第二阶段的
    商品记录（与第一阶段的 {域名}.json 也不会同名）。
    """
    input_path = Path(input_file)
    replaced, found = _replace_path_component(input_path)
    output_dir = replaced.parent if found else input_path.parent / config.OUTPUT_DIRNAME
    output_path = output_dir / f"{input_path.stem}商品{RECORD_FILE_SUFFIX}"

    try:
        same_path = output_path.resolve(strict=False) == input_path.resolve(strict=False)
    except OSError:
        same_path = os.path.abspath(str(output_path)) == os.path.abspath(str(input_path))
    if same_path:
        raise RuntimeError(f"输出路径与输入文件相同，已阻止覆盖: {input_path}")
    return str(output_path)


def build_recovery_json_path(output_file):
    """已有商品记录 JSON 无法读取时，为本轮数据生成绝不覆盖旧文件的恢复名。"""
    output_path = Path(output_file)
    timestamp = time.strftime("%Y%m%d-%H%M%S")
    candidate = output_path.with_name(
        f"{output_path.stem}.{RECOVERY_MARK}-{timestamp}{output_path.suffix}"
    )
    counter = 1
    while candidate.exists():
        candidate = output_path.with_name(
            f"{output_path.stem}.{RECOVERY_MARK}-{timestamp}-{counter}{output_path.suffix}"
        )
        counter += 1
    return str(candidate)


def build_recovery_xlsx_path(output_file):
    """旧 XLSX 恢复文件名构造；JSON 迁移后仅保留给回退路径使用。"""
    output_path = Path(output_file)
    timestamp = time.strftime("%Y%m%d-%H%M%S")
    candidate = output_path.with_name(
        f"{output_path.stem}.{RECOVERY_MARK}-{timestamp}{output_path.suffix}"
    )
    counter = 1
    while candidate.exists():
        candidate = output_path.with_name(
            f"{output_path.stem}.{RECOVERY_MARK}-{timestamp}-{counter}{output_path.suffix}"
        )
        counter += 1
    return str(candidate)