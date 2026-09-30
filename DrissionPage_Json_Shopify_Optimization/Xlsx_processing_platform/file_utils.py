"""商品中间表读取与原子保存。"""

import os
import tempfile
from pathlib import Path

import openpyxl

import config


def save_xlsx(rows, output_path):
    """原子写入 XLSX：先写同目录临时文件，成功后 os.replace 覆盖目标。

    签名与原实现保持一致，所有调用方无需改动。
    失败时清理临时文件并抛出异常；原文件不会被破坏。
    """
    output_path = Path(output_path)
    target_dir = output_path.parent if str(output_path.parent) else Path(".")
    target_dir.mkdir(parents=True, exist_ok=True)

    fd, tmp_path = tempfile.mkstemp(
        prefix=f".{output_path.stem}_tmp_",
        suffix=".xlsx",
        dir=str(target_dir),
    )
    os.close(fd)
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(config.CSV_FIELDS)
        for row in rows:
            ws.append([row.get(f, "") for f in config.CSV_FIELDS])
            for field in ('price1', 'price2'):
                ws.cell(ws.max_row, config.CSV_FIELDS.index(field) + 1).number_format = '0.00'
        wb.save(tmp_path)
        wb.close()
        os.replace(tmp_path, str(output_path))
    except Exception:
        try:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
        except OSError:
            pass
        raise


def load_product_rows(filepath, *, with_locations=False):
    """按列名读取中间表，保留 0/False，跳过全空行。"""
    wb = openpyxl.load_workbook(filepath, read_only=True)
    try:
        rows = wb.active.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(rows, ())]
        required = [field for field in config.CSV_FIELDS if field != "link-href"]
        missing = [field for field in required if field not in headers]
        if missing:
            raise ValueError(f"{Path(filepath).name} 缺少字段: {', '.join(missing)}")
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise ValueError(f"{Path(filepath).name} 存在重复列名")
        result = []
        for row_number, values in enumerate(rows, 2):
            if all(value is None or value == "" for value in values):
                continue
            row = dict(zip(headers, values))
            item = {field: "" if row.get(field) is None else row[field]
                    for field in config.CSV_FIELDS}
            result.append((row_number, item) if with_locations else item)
        return result
    finally:
        wb.close()
